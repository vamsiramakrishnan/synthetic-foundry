#!/usr/bin/env python3
"""Ingest APQC Process Classification Framework workbooks into versioned data.

Input: a directory of the Excel files APQC publishes (one per framework: the
cross-industry PCF and the industry PCFs). Output: one gzipped JSON file per
framework under ``src/worldloom/_data/pcf/`` plus a provenance ledger.

What is taken from each workbook, and nothing else:

- the ``Combined`` sheet: every process element's stable ``PCF ID``, its
  ``Hierarchy ID`` (the human index, which changes between releases), its
  name, and whether APQC has benchmarking metrics for it;
- the element description, from the ``Element Description`` column where the
  workbook has one, else from the ``Glossary terms`` sheet by PCF ID;
- the ``Metrics`` sheet where present (the cross-industry file): metric id,
  name, formula, units and category per element;
- the copyright and attribution notice, verbatim, from the workbook's text
  boxes or front sheets (releases differ). APQC's licence permits copying,
  publishing and derivative works on the condition that every copy carries
  the notice, so every output file embeds it and a workbook whose notice
  cannot be found is refused.

The parent of an element is derived from its hierarchy id: ``1.1.1`` sits
under ``1.1``, ``1.1`` under ``1.0``. Levels are 1 (category) to 5 (task).
Rows are written in hierarchy order, keys sorted, so the same workbook yields
the same bytes. The raw workbooks are inputs, not repository assets.

Usage::

    python tools/ingest_apqc.py --input ./apqc_pcf --out src/worldloom/_data/pcf
"""

from __future__ import annotations

import argparse
import gzip
import hashlib
import html
import json
import re
import sys
import zipfile
from pathlib import Path
from typing import Any

SCHEMA = "worldloom.pcf/v1"

_VERSION = re.compile(r"VERSION NUMBER\s*([0-9][0-9.]*)", re.IGNORECASE)
_GENERATED = re.compile(r"GENERATED ON\s*([0-9/]+)", re.IGNORECASE)
_TITLE = re.compile(r"^\s*(.*?)PROCESS CLASSIFICATION FRAMEWORK", re.IGNORECASE | re.DOTALL)
_HEADING = re.compile(r"COPYRIGHT AND ATTRIBUTION", re.IGNORECASE)


def _text_boxes(path: Path) -> list[str]:
    """Every text box in the workbook's drawings, as plain text."""
    out: list[str] = []
    with zipfile.ZipFile(path) as archive:
        for name in sorted(archive.namelist()):
            if name.startswith("xl/drawings/") and name.endswith(".xml"):
                xml = archive.read(name).decode("utf-8")
                text = "".join(html.unescape(t) for t in re.findall(r"<a:t>([^<]*)</a:t>", xml))
                if text.strip():
                    out.append(text)
    return out


def _front_matter(path: Path, workbook: Any) -> list[str]:
    """The workbook's prose: its text boxes, then the cells of its front sheets.

    Releases differ in where they keep the title and the licence: 7.4 and 7.2.1
    put them in drawing text boxes, 7.2.2 in the cells of the Introduction,
    About and Copyright sheets. Both are read; each block is one string.
    """
    import openpyxl

    blocks = _text_boxes(path)
    # The read-only reader reports these sheets as empty (they declare no
    # dimensions), so they are read through the ordinary loader; the data
    # sheets stay on the streaming reader.
    front = openpyxl.load_workbook(path, read_only=False, data_only=True)
    for sheet in front.sheetnames:
        if sheet.lower().startswith(("introduction", "about", "copyright")):
            cells = [str(cell.value) for row in front[sheet].iter_rows() for cell in row if cell.value]
            if cells:
                blocks.append("\n".join(cells))
    # Last resort: the shared-string table holds every cell string, so a
    # licence kept on a sheet the loop above did not name is still found.
    with zipfile.ZipFile(path) as archive:
        if "xl/sharedStrings.xml" in archive.namelist():
            table = archive.read("xl/sharedStrings.xml").decode("utf-8")
            for text in re.findall(r"<t[^>]*>([^<]*)</t>", table):
                if "royalty-free" in text.lower():
                    blocks.append(html.unescape(text))
    return blocks


def _notice(blocks: list[str]) -> str:
    """The copyright and attribution notice, verbatim from its first character.

    Taken from the licence sentence's own block: from the "©" (or, in the
    APQC and IBM industry frameworks, the sentence that introduces it) to the
    end of that block. Whitespace is normalised; nothing else is changed.
    """
    for block in blocks:
        if "royalty-free" not in block.lower():
            continue
        text = block
        heading = _HEADING.search(text)
        if heading:
            text = text[heading.end():]
        start = text.find("This Industry Process Classification Framework")
        if start < 0:
            start = text.find("©")
        if start < 0:
            continue
        return " ".join(text[start:].split())
    return ""


def _framework_identity(path: Path, boxes: list[str]) -> tuple[str, str, str]:
    """(framework slug, version, generated-on) from the title box, else the filename."""
    title_box = next((b for b in boxes if "PROCESS CLASSIFICATION FRAMEWORK" in b.upper() and "VERSION NUMBER" in b.upper()), "")
    version = ""
    generated = ""
    name = ""
    if title_box:
        m = _VERSION.search(title_box)
        version = m.group(1) if m else ""
        g = _GENERATED.search(title_box)
        generated = g.group(1) if g else ""
        t = _TITLE.search(title_box)
        name = t.group(1).strip() if t else ""
    if not name:
        stem = path.stem
        stem = re.sub(r"^K\d+_?\s*", "", stem)
        stem = re.split(r"_v\d|_vs|\s*-\s*Excel", stem, maxsplit=1)[0]
        name = stem.replace("APQC Process Classification Framework (PCF) - ", "").replace("_", " ").strip()
    if not version:
        m = re.search(r"[_ ]v(\d)(\d)(\d)?\b", path.stem)
        if m:
            version = ".".join(part for part in m.groups() if part)
        else:
            m2 = re.search(r"Version (\d[\d.]*)", path.stem)
            version = m2.group(1) if m2 else "unknown"
    slug = re.sub(r"[^a-z0-9]+", "_", name.lower()).strip("_") or "framework"
    return slug, version, generated


def _level(hierarchy: str) -> int:
    return 1 if hierarchy.endswith(".0") else hierarchy.count(".") + 1


def _parent(hierarchy: str) -> str | None:
    if hierarchy.endswith(".0"):
        return None
    head, _, _ = hierarchy.rpartition(".")
    return f"{head}.0" if "." not in head else head


def _sort_key(hierarchy: str) -> tuple[int, ...]:
    return tuple(int(part) for part in hierarchy.split("."))


def _cell(value: Any) -> str:
    return "" if value is None else str(value).strip()


def ingest(path: Path) -> dict[str, Any]:
    import openpyxl

    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    blocks = _front_matter(path, workbook)
    slug, version, generated = _framework_identity(path, blocks)
    notice = _notice(blocks)

    glossary: dict[str, str] = {}
    if "Glossary terms" in workbook.sheetnames:
        rows = workbook["Glossary terms"].iter_rows(values_only=True)
        next(rows, None)
        for row in rows:
            if row and _cell(row[0]):
                glossary[_cell(row[0])] = _cell(row[3]) if len(row) > 3 else ""

    combined = workbook["Combined"].iter_rows(values_only=True)
    header = [_cell(c) for c in next(combined)]
    column = {name: index for index, name in enumerate(header)}
    described = "Element Description" in column
    elements: list[dict[str, Any]] = []
    seen: set[str] = set()
    for row in combined:
        if not row or not _cell(row[column["PCF ID"]]):
            continue
        pcf_id = _cell(row[column["PCF ID"]])
        hierarchy = _cell(row[column["Hierarchy ID"]])
        if not hierarchy or pcf_id in seen:
            continue
        seen.add(pcf_id)
        description = _cell(row[column["Element Description"]]) if described else glossary.get(pcf_id, "")
        elements.append({
            "pcf_id": pcf_id,
            "hierarchy_id": hierarchy,
            "level": _level(hierarchy),
            "parent_hierarchy_id": _parent(hierarchy),
            "name": _cell(row[column["Name"]]),
            "description": description,
            "metrics_available": _cell(row[column["Metrics available?"]]).upper() == "Y" if "Metrics available?" in column else False,
        })
    elements.sort(key=lambda e: _sort_key(e["hierarchy_id"]))
    by_hierarchy = {e["hierarchy_id"]: e["pcf_id"] for e in elements}
    for element in elements:
        parent = element.pop("parent_hierarchy_id")
        element["parent_pcf_id"] = by_hierarchy.get(parent) if parent else None

    metrics: list[dict[str, Any]] = []
    if "Metrics" in workbook.sheetnames:
        rows = workbook["Metrics"].iter_rows(values_only=True)
        next(rows, None)
        for row in rows:
            if not row or not _cell(row[0]):
                continue
            metrics.append({
                "element_pcf_id": _cell(row[0]),
                "metric_id": _cell(row[4]),
                "name": _cell(row[5]),
                "formula": _cell(row[6]),
                "units": _cell(row[7]).lower(),
                "category": _cell(row[3]),
            })
        metrics.sort(key=lambda m: (_sort_key(next((e["hierarchy_id"] for e in elements if e["pcf_id"] == m["element_pcf_id"]), "0.0")), m["metric_id"], m["name"]))

    digest = hashlib.sha256(path.read_bytes()).hexdigest()
    return {
        "schema": SCHEMA,
        "framework": slug,
        "version": version,
        "source": {"file": path.name, "sha256": digest, "bytes": path.stat().st_size, "generated_on": generated},
        "notice": notice,
        "levels": {"1": "category", "2": "process_group", "3": "process", "4": "activity", "5": "task"},
        "elements": elements,
        "metrics": metrics,
    }


def write(document: dict[str, Any], out: Path) -> Path:
    out.mkdir(parents=True, exist_ok=True)
    target = out / f"{document['framework']}@{document['version']}.json.gz"
    payload = json.dumps(document, sort_keys=True, ensure_ascii=False, separators=(",", ":")).encode("utf-8")
    with gzip.GzipFile(target, "wb", mtime=0) as handle:
        handle.write(payload)
    return target


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__.split("\n\n")[0])
    parser.add_argument("--input", required=True, type=Path, help="Directory of APQC PCF .xlsx files.")
    parser.add_argument("--out", required=True, type=Path, help="Output directory (src/worldloom/_data/pcf).")
    args = parser.parse_args(argv)

    ledger: list[dict[str, Any]] = []
    skipped: list[str] = []
    for path in sorted(args.input.glob("*.xls*")):
        if path.suffix.lower() != ".xlsx":
            skipped.append(f"{path.name}: not an xlsx workbook")
            continue
        try:
            document = ingest(path)
        except (KeyError, zipfile.BadZipFile, ValueError) as error:
            skipped.append(f"{path.name}: {error}")
            continue
        if not document["notice"]:
            skipped.append(f"{path.name}: no APQC copyright notice found; refusing to write a copy without it")
            continue
        target = write(document, args.out)
        ledger.append({
            "framework": document["framework"], "version": document["version"], "file": target.name,
            "source_file": document["source"]["file"], "source_sha256": document["source"]["sha256"],
            "generated_on": document["source"]["generated_on"],
            "elements": len(document["elements"]), "metrics": len(document["metrics"]),
            "described": sum(1 for e in document["elements"] if e["description"]),
        })
        print(f"{target.name:48} elements={len(document['elements']):5} described={ledger[-1]['described']:5} metrics={len(document['metrics'])}")
    ledger.sort(key=lambda row: row["file"])
    (args.out / "provenance.json").write_text(
        json.dumps({"schema": "worldloom.pcf-provenance/v1", "licence": "APQC PCF notice embedded in every file; see `notice`.",
                    "frameworks": ledger, "skipped": skipped}, indent=2, sort_keys=True) + "\n", encoding="utf-8")
    for line in skipped:
        print("skipped:", line, file=sys.stderr)
    return 0 if ledger else 1


if __name__ == "__main__":
    raise SystemExit(main())
