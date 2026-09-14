"""A document type declares that it is a deck, and the deck renderer obeys.

`render.pptx.HANDLES` was a frozenset holding one name. A pack could author a
board pack with an outline, a filing rule and a size, and the only format it
could never reach was the one a board pack is read in. `deck: true` on the
authored type is the door: `doctypes.install` registers it, `describe` reads
it back, `registries.scoped` puts the set back afterwards, and an unset flag
never touches the wire of a type authored before decks were declarable.
"""

from __future__ import annotations

import io
import json
import pathlib

import pptx as python_pptx
import pytest

from worldloom import MonthEndClose, RetailWorld, World, doctypes, packs, registries
from worldloom.narrative import DeterministicProvider
from worldloom.render import pptx as pptx_render

EXAMPLES = pathlib.Path(__file__).resolve().parents[1] / "examples" / "artifact-types"
AUTHORED = EXAMPLES / "franchise-network.json"
STATEMENT = "franchisee_trading_statement"
PERIOD = "2026-03"


@pytest.fixture(autouse=True)
def _restore_the_registries():
    with registries.scoped():
        yield


def _pack(deck: bool) -> dict:
    document = json.loads(AUTHORED.read_text(encoding="utf-8"))
    if deck:
        document["artifact_types"][0]["deck"] = True
    return document


def _world(document: dict) -> World:
    return RetailWorld.from_pack(packs.load(document), seed=8128).build().run(
        MonthEndClose(period=PERIOD, include_operational_incident=True)
    )


def test_an_unset_deck_flag_stays_off_the_wire() -> None:
    types = doctypes.load(_pack(deck=False)["artifact_types"])
    dumped = doctypes.to_document(types)["artifact_types"][0]
    assert "deck" not in dumped and dumped["word"] is True
    assert types[0].deck is False

    declared = doctypes.load(_pack(deck=True)["artifact_types"])
    assert doctypes.to_document(declared)["artifact_types"][0]["deck"] is True


def test_the_core_port_marks_the_executive_summary_as_the_deck_it_always_was() -> None:
    ported = {spec.key: spec for spec in doctypes.load(EXAMPLES / "core.json")}
    assert ported["executive_summary"].deck is True
    assert doctypes.describe("executive_summary").deck is True
    assert all(not spec.deck for key, spec in ported.items() if key != "executive_summary")


def test_an_authored_deck_type_is_registered_installed_and_restored() -> None:
    before = set(pptx_render.HANDLES)
    with registries.scoped():
        doctypes.install(doctypes.load(_pack(deck=True)["artifact_types"]))
        assert STATEMENT in pptx_render.HANDLES
        assert doctypes.describe(STATEMENT).deck is True
    assert set(pptx_render.HANDLES) == before


def test_an_authored_deck_renders_as_a_deck_and_a_document() -> None:
    """The same outline reaches Word, PDF, Markdown and now PowerPoint."""
    with registries.scoped():
        world = _world(_pack(deck=True)).narrate(DeterministicProvider()).render("pptx", "docx", "markdown")
    intent = next(i for i in world.artifact_intents if i.artifact_type == STATEMENT)
    by_suffix = {
        suffix: {r.artifact_id: r.payload for r in world._rendered if r.path.endswith(suffix)}
        for suffix in (".pptx", ".docx", ".md")
    }
    assert intent.id in by_suffix[".pptx"], "the statement declared itself a deck"
    assert intent.id in by_suffix[".docx"] and intent.id in by_suffix[".md"]

    deck = python_pptx.Presentation(io.BytesIO(by_suffix[".pptx"][intent.id]))
    headings = [
        shape.text_frame.text for slide in deck.slides for shape in slide.shapes
        if shape.has_text_frame
    ]
    ir = next(ir for ir in world.artifact_irs if ir.intent_id == intent.id)
    for section in ir.sections:
        if not section.hidden:
            assert section.heading in headings, section.heading

    # A world whose type did not declare a deck renders none for it — the
    # flag, not the outline, is what decides.
    # Its own scope: the same key installed with a different flag would be
    # refused as a redefinition, which is `install`'s rule, not a defect.
    with registries.scoped():
        plain = _world(_pack(deck=False)).narrate(DeterministicProvider()).render("pptx")
    produced = {r.artifact_id for r in plain._rendered if r.path.endswith(".pptx")}
    assert next(i for i in plain.artifact_intents if i.artifact_type == STATEMENT).id not in produced
