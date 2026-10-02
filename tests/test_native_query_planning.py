from __future__ import annotations

import base64
import hashlib
import re
from dataclasses import replace
from datetime import UTC, datetime
from io import BytesIO

import pytest

from worldloom.models import (
    ArtifactIR,
    ArtifactSection,
    Authority,
    CanonicalFact,
    Cell,
    Column,
    Company,
    FormulaKind,
    Quantity,
    Row,
    Table,
)
from worldloom.native_artifacts import inspect_artifact
from worldloom.native_corpus import (
    NativeContent,
    NativeCorpusPlan,
    render_native_corpus,
)
from worldloom.native_eval_bridge import native_task_cases
from worldloom.native_query_planning import NativeWorkloadPlan, plan_native_workload
from worldloom.native_reference import qualify_native_task, reference_submission
from worldloom.native_tasks import NativeCitation, grade_native_task, public_contract
from worldloom.studio import NativeSuiteRequest, ProjectSpec, UseCase, preset
from worldloom.studio.native_suite import propose
from worldloom.world import World


def _world(surface="legacy", period="2026-01") -> tuple:
    facts = tuple(CanonicalFact(id=f"FACT-{kind}-{subject}".upper(), kind="financial.revenue." + kind,
        subject="store:" + subject, period=period, value=Quantity(amount=value, unit="AUD"),
        valid_from=datetime(2026, 1, 1, tzinfo=UTC), authority=Authority.SYSTEM_OF_RECORD)
        for kind, subject, value in (("actual", "east", 125), ("actual", "west", 70),
                                     ("budget", "east", 100), ("budget", "west", 60)))
    sections = [ArtifactSection(heading=f"{f.subject} {f.kind.rsplit('.', 1)[1]} revenue",
        body=f"Revenue evidence for {f.subject}: {{{{fact:{f.id}}}}}.", fact_ids=[f.id]) for f in facts]
    world = World(seed=8128, company=Company(id="CO-1", name="Northstar Retail", industry="retail",
        headquarters="Sydney", fiscal_year_start_month=7, employees_total=100), _facts=facts,
        _artifact_irs=(ArtifactIR(id="ART-SOURCE", intent_id="INTENT-1", title="Revenue close", sections=sections),))
    plans = tuple(NativeCorpusPlan(artifact_id=f"{group}-{format}", format=format, title=f"{group} revenue close",
        surface=surface, minimum_units=2, contents=tuple(NativeContent(source_artifact_id="ART-SOURCE", section_index=index,
            placement="notes" if format == "pptx" else "body") for index in indexes))
        for group, indexes in (("actual", (0, 1)), ("budget", (2, 3))) for format in ("docx", "pptx", "xlsx"))
    return world, {plan.artifact_id: render_native_corpus(world, plan) for plan in plans}


def _plan(**overrides) -> NativeWorkloadPlan:
    return NativeWorkloadPlan(use_case_id="revenue-review", objective="Reconcile the reporting period's revenue evidence.",
                              **overrides)


def test_business_workload_is_executable_deterministic_and_hides_source_oracles() -> None:
    world, rendered = _world()
    workload = plan_native_workload(world, rendered, _plan(max_tasks=24))
    assert workload == plan_native_workload(world, dict(reversed(list(rendered.items()))), _plan(max_tasks=24))
    assert set(workload.operation_counts) == {"read", "analyze", "update", "create"}
    assert workload.reference_qualified == len(workload.tasks) == 24
    assert workload.capability_coverage["cross_artifact"] > 0
    assert workload.capability_coverage["actual_budget_variance"] > 0
    assert workload.capability_coverage["formula_output"] > 0
    assert workload.capability_coverage["content_update"] > 0
    assert not any(f.code == "reference_qualification_failed" for f in workload.findings)
    payloads = {key: value.payload for key, value in rendered.items()}
    for task in workload.tasks:
        assert qualify_native_task(task, payloads).passed
        public = public_contract(task)
        assert "expected" not in public
        for assertion in task.assertions:
            refs = assertion.calculation.operands if assertion.calculation else (assertion.target,)
            for ref in refs:
                assert ref is not None and ref.locator not in task.prompt
        assert "125" not in task.prompt and "100" not in task.prompt


def test_cross_file_variance_requires_both_actual_sources_and_rejects_fabricated_values() -> None:
    world, rendered = _world()
    workload = plan_native_workload(world, rendered, _plan(operations=("analyze",), discovery_scope="cross_artifact"))
    task = next(t for t in workload.tasks if "financial.revenue.actual" in t.prompt and "financial.revenue.budget" in t.prompt)
    assert len(task.inputs) == 2
    operands = task.assertions[0].calculation.operands
    assert operands[0].artifact_id != operands[1].artifact_id
    payloads = {key: value.payload for key, value in rendered.items()}
    reference = reference_submission(task, payloads)
    assert grade_native_task(task, payloads, reference).passed
    answer = reference.answers[0]
    wrong_value = reference.model_copy(update={"answers": (answer.model_copy(update={"value": "999999"}),)})
    assert not grade_native_task(task, payloads, wrong_value).passed
    wrong_citation = NativeCitation(artifact_id=operands[1].artifact_id, locator=operands[0].locator)
    same_value_wrong_source = reference.model_copy(update={"answers": (answer.model_copy(update={"citations": (wrong_citation, operands[1])}),)})
    assert not grade_native_task(task, payloads, same_value_wrong_source).passed


def test_content_updates_preserve_every_unrelated_unit_and_pin_the_source_version() -> None:
    from docx import Document

    world, rendered = _world()
    workload = plan_native_workload(world, rendered, _plan(formats=("docx",), operations=("update",), discovery_scope="cross_artifact"))
    task = workload.tasks[0]
    source_id = task.output.source_artifact_id
    source_ref = task.output.assertions[0].target
    assert source_ref is not None and source_ref.locator.startswith("paragraph:")
    assert source_ref.locator != "paragraph:1"
    payloads = {key: value.payload for key, value in rendered.items()}
    reference = reference_submission(task, payloads)
    assert grade_native_task(task, payloads, reference).passed
    submitted = reference.files[0]
    document = Document(BytesIO(base64.b64decode(submitted.content_base64)))
    document.paragraphs[0].text = "Unapproved change to another heading"
    out = BytesIO()
    document.save(out)
    corrupted = reference.model_copy(update={"files": (submitted.model_copy(update={"content_base64": base64.b64encode(out.getvalue()).decode()}),)})
    assert "update_unaffected_content_changed" in grade_native_task(task, payloads, corrupted).findings
    stale = reference.model_copy(update={"files": (submitted.model_copy(update={"source_sha256": "0" * 64}),)})
    assert "update_source_version_mismatch" in grade_native_task(task, payloads, stale).findings
    assert source_id in payloads


def test_replicated_formats_do_not_purchase_independent_cross_file_coverage() -> None:
    world, rendered = _world()
    replicas = {key: value for key, value in rendered.items() if key.startswith("actual-")}
    workload = plan_native_workload(world, replicas, _plan(discovery_scope="cross_artifact"))
    assert not workload.tasks
    assert any(f.code == "independent_cross_artifact_evidence_missing" for f in workload.findings)
    assert {f.operation for f in workload.findings if f.code == "operation_unavailable"} == set(_plan().operations)


def test_manifest_tampering_and_bytes_from_another_company_are_refused() -> None:
    world, rendered = _world()
    key, original = next(iter(rendered.items()))
    stale = replace(original, payload=original.payload + b"tampered")
    with pytest.raises(ValueError, match="source bytes changed"):
        plan_native_workload(world, {key: stale}, _plan())
    bad_entry = original.manifest.evidence[0].model_copy(update={"text_sha256": "0" * 64})
    bad_manifest = original.manifest.model_copy(update={"evidence": (bad_entry, *original.manifest.evidence[1:])})
    with pytest.raises(ValueError, match="provenance changed"):
        plan_native_workload(world, {key: replace(original, manifest=bad_manifest)}, _plan())
    with pytest.raises(ValueError, match="unknown canonical facts"):
        plan_native_workload(replace(world, _facts=()), {key: original}, _plan())


@pytest.mark.parametrize("mutation", ["section", "facts", "heading", "number", "unit"])
def test_valid_bytes_do_not_authorize_forged_provenance_bindings(mutation: str) -> None:
    world, rendered = _world()
    key = "actual-xlsx" if mutation in ("number", "unit") else "actual-docx"
    original = rendered[key]
    entries = list(original.manifest.evidence)
    if mutation in ("number", "unit"):
        index = next(i for i, entry in enumerate(entries) if entry.value is not None)
        change = {"value": 999} if mutation == "number" else {"unit": "USD"}
        entries[index] = entries[index].model_copy(update=change)
    else:
        change = {"section_index": 1} if mutation == "section" else {"fact_ids": (world.facts[1].id,)}
        if mutation == "heading":
            ir = world.artifact_irs[0]
            section = ir.sections[0].model_copy(update={"heading": ir.sections[1].heading})
            world = replace(world, _artifact_irs=(ir.model_copy(update={"sections": [section, *ir.sections[1:]]}),))
        else:
            entries[0] = entries[0].model_copy(update=change)
    manifest = original.manifest.model_copy(update={"evidence": tuple(entries)})
    with pytest.raises(ValueError, match=r"provenance disagrees|heading is not bound|numeric evidence disagrees"):
        plan_native_workload(world, {key: replace(original, manifest=manifest)}, _plan())


def test_ambiguous_business_evidence_is_reported_without_a_locator_fallback() -> None:
    world, _ = _world()
    sections = [s.model_copy(update={"heading": "Revenue", "body": "Revenue: {{fact:" + f.id + "}}"})
                for s, f in zip(world.artifact_irs[0].sections, world.facts, strict=True)]
    world = replace(world, _artifact_irs=(world.artifact_irs[0].model_copy(update={"sections": sections}),))
    corpus = NativeCorpusPlan(artifact_id="ambiguous", format="docx", title="Revenue", minimum_units=4,
        contents=tuple(NativeContent(source_artifact_id="ART-SOURCE", section_index=i) for i in range(4)))
    workload = plan_native_workload(world, {"ambiguous": render_native_corpus(world, corpus)}, _plan(formats=("docx",), operations=("read",)))
    assert not workload.tasks
    assert sum(f.code == "ambiguous_evidence_selector" for f in workload.findings) == 4


def test_small_budgets_reach_multiple_real_capabilities() -> None:
    world, rendered = _world()
    workload = plan_native_workload(world, rendered, _plan(max_tasks=8))
    assert len(workload.tasks) == 8 and workload.truncated
    assert set(workload.operation_counts) == {"read", "analyze", "update", "create"}
    assert workload.capability_coverage["cross_artifact"] > 0


def test_cross_period_duplicates_without_visible_period_metadata_remain_a_gap() -> None:
    world, _ = _world()
    old_fact = world._facts[0]
    later = old_fact.model_copy(update={"id": "FACT-LATER", "period": "2026-02", "value": Quantity(amount=130, unit="AUD")})
    ir = world.artifact_irs[0].model_copy(update={"sections": [
        world.artifact_irs[0].sections[0], ArtifactSection(heading="Later revenue", body="Later revenue: {{fact:FACT-LATER}}", fact_ids=[later.id])
    ]})
    world = replace(world, _facts=(old_fact, later), _artifact_irs=(ir,))
    corpus = NativeCorpusPlan(artifact_id="periods", format="xlsx", title="Revenue", minimum_units=2,
        contents=tuple(NativeContent(source_artifact_id="ART-SOURCE", section_index=i) for i in range(2)))
    result = render_native_corpus(world, corpus)
    workload = plan_native_workload(world, {"periods": result}, _plan(formats=("xlsx",), operations=("analyze",)))
    assert not workload.tasks
    assert any(f.code == "ambiguous_numeric_selector" for f in workload.findings)


def test_studio_discovery_compiles_independent_business_work_without_locator_prompts() -> None:
    world, _ = _world()
    spec = preset().model_copy(update={"use_cases": (UseCase(id="revenue-review", title="Revenue review",
        objective="Reconcile the reporting period's revenue evidence."),)})
    request = NativeSuiteRequest(use_case_id="revenue-review", query_style="discovery", minimum_units=2,
        max_cases=2, max_tasks_per_case=16)
    result = propose(world, spec, request)
    assert result == propose(world, spec, request)
    proposal = ProjectSpec.model_validate(result["spec"])
    assert result["summary"]["prepared_cases"] == 1
    assert result["summary"]["reference_qualified"] == 16
    assert result["summary"]["capability_coverage"]["cross_artifact"] > 0
    assert {t.operation for t in proposal.native_tasks} == {"read", "analyze", "update", "create"}
    payloads = {p.artifact_id: render_native_corpus(world, p).payload for p in proposal.native_corpus}
    for task in proposal.native_tasks:
        assert qualify_native_task(task, payloads).passed
        for assertion in task.assertions:
            refs = assertion.calculation.operands if assertion.calculation else (assertion.target,)
            assert all(ref is not None and ref.locator not in task.prompt for ref in refs)


def test_studio_discovery_options_leave_default_request_serialization_unchanged() -> None:
    request = NativeSuiteRequest(use_case_id="review")
    assert request.model_dump(mode="json") == {"use_case_id": "review", "formats": ["docx", "pptx", "xlsx"],
        "operations": ["read", "analyze", "update", "create"], "minimum_units": 2, "max_cases": 12,
        "source_artifact_ids": []}
    assert NativeSuiteRequest.model_validate(request.model_dump(mode="json")) == request


def test_business_scalar_surfaces_qualify_cross_file_queries_in_every_native_format() -> None:
    world, rendered = _world("business")
    workload = plan_native_workload(world, rendered, _plan(max_tasks=32))
    assert workload.reference_qualified == len(workload.tasks) == 32
    assert set(workload.operation_counts) == {"read", "analyze", "update", "create"}
    assert workload.capability_coverage["actual_budget_variance"] > 0
    assert not any(f.code == "reference_qualification_failed" for f in workload.findings)
    assert all("FACT-" not in unit.text for result in rendered.values()
        for unit in inspect_artifact(result.payload, result.manifest.format).units)
    for format in ("docx", "pptx", "xlsx"):
        scoped = plan_native_workload(world, rendered, _plan(formats=(format,), operations=("analyze",), discovery_scope="cross_artifact"))
        assert scoped.tasks and all(len(task.inputs) == 2 for task in scoped.tasks)
        assert all(qualify_native_task(task, {key: value.payload for key, value in rendered.items()}).passed for task in scoped.tasks)


def test_unperiodized_business_scalars_accept_only_blank_period_fields_in_all_formats() -> None:
    world, rendered = _world("business", period=None)
    workload = plan_native_workload(world, rendered, _plan(operations=("read",), max_tasks=24))
    assert workload.reference_qualified == len(workload.tasks) > 0
    assert {item.format for task in workload.tasks for item in task.inputs} == {"docx", "pptx", "xlsx"}
    assert not any(f.code == "reference_qualification_failed" for f in workload.findings)
    changed = replace(world, _facts=tuple(fact.model_copy(update={"period": "2026-02"}) for fact in world.facts))
    for format in ("docx", "pptx", "xlsx"):
        with pytest.raises(ValueError, match="reporting period disagrees"):
            plan_native_workload(changed, {key: value for key, value in rendered.items() if key.endswith(format)},
                _plan(formats=(format,), operations=("read",)))


def _table_world(format="xlsx") -> tuple:
    world, _ = _world()
    first, _, budget, _ = world.facts
    table = Table(key="revenue", title="Store revenue reconciliation", columns=[
        Column(key="actual", label="Actual"), Column(key="budget", label="Budget"), Column(key="variance", label="Variance")],
        rows=[Row(key="east", label="East store", cells={
            "actual": Cell(value=first.value.amount, fact_id=first.id),
            "budget": Cell(value=budget.value.amount, fact_id=budget.id),
            "variance": Cell(value=25, formula=FormulaKind.DIFFERENCE, operands=["actual", "budget"])})])
    section = ArtifactSection(heading="Revenue close", table=table)
    world = replace(world, _facts=(first, budget), _artifact_irs=(world.artifact_irs[0].model_copy(update={"sections": [section]}),))
    plan = NativeCorpusPlan(artifact_id="reconciliation", format=format, title="Store revenue", surface="business",
        contents=(NativeContent(source_artifact_id="ART-SOURCE", section_index=0),))
    return world, render_native_corpus(world, plan)


@pytest.mark.parametrize("format", ["docx", "pptx", "xlsx"])
def test_table_only_business_sources_bind_values_and_formula_dependencies(format: str) -> None:
    world, result = _table_world(format)
    workload = plan_native_workload(world, {"reconciliation": result}, _plan(formats=(format,), max_tasks=32))
    assert workload.tasks and workload.capability_coverage["table_discovery"] == 3
    assert all(qualify_native_task(task, {"reconciliation": result.payload}).passed for task in workload.tasks)
    if format == "xlsx":
        assert workload.capability_coverage["native_formula_inspection"] == 1
    else:
        assert "native_formula_inspection" not in workload.capability_coverage


@pytest.mark.parametrize("mutation", ["column", "facts", "dependencies", "selector"])
def test_table_provenance_cannot_rebind_valid_bytes_to_other_semantics(mutation: str) -> None:
    world, result = _table_world()
    entries = list(result.manifest.evidence)
    index = next(i for i, entry in enumerate(entries) if entry.kind == "formula")
    original = entries[index]
    mutations = {"column": {"column_key": "actual"}, "facts": {"fact_ids": (world.facts[0].id,)},
        "dependencies": {"dependency_locators": tuple(reversed(original.dependency_locators))},
        "selector": {"metadata_locators": {**original.metadata_locators, "column_label": original.metadata_locators["row_label"]}}}
    entries[index] = original.model_copy(update=mutations[mutation])
    manifest = result.manifest.model_copy(update={"evidence": tuple(entries)})
    with pytest.raises(ValueError, match="native table"):
        plan_native_workload(world, {"reconciliation": replace(result, manifest=manifest)}, _plan())


@pytest.mark.parametrize("format", ["docx", "pptx", "xlsx"])
def test_sdk_intake_rejects_derived_truth_forged_in_ir_manifest_and_native_bytes(format: str) -> None:
    world, result = _table_world(format)
    original_tasks = plan_native_workload(world, {"reconciliation": result}, _plan(formats=(format,))).tasks
    entry = next(entry for entry in result.manifest.evidence if entry.kind == "formula")
    payload = result.payload
    if format == "docx":
        from docx import Document
        document = Document(BytesIO(payload))
        position = re.fullmatch(r"table:(\d+)/row:(\d+)/cell:(\d+)", entry.locator)
        assert position is not None
        document.tables[int(position[1]) - 1].cell(int(position[2]) - 1, int(position[3]) - 1).text = "999"
        stream = BytesIO()
        document.save(stream)
        payload = stream.getvalue()
    elif format == "pptx":
        from pptx import Presentation
        presentation = Presentation(BytesIO(payload))
        position = re.fullmatch(r"slide:(\d+)/shape:(\d+)/row:(\d+)/cell:(\d+)", entry.locator)
        assert position is not None
        shape = presentation.slides[int(position[1]) - 1].shapes[int(position[2]) - 1]
        shape.table.cell(int(position[3]) - 1, int(position[4]) - 1).text = "999"
        stream = BytesIO()
        presentation.save(stream)
        payload = stream.getvalue()
    snapshot = inspect_artifact(payload, format)
    units = {unit.locator: unit.text for unit in snapshot.units}
    forged_entry = entry.model_copy(update={"value": 999, "text_sha256": hashlib.sha256(units[entry.locator].encode()).hexdigest()})
    manifest = result.manifest.model_copy(update={"sha256": snapshot.sha256, "file_size_bytes": len(payload),
        "native_metrics": snapshot.metrics, "evidence": tuple(forged_entry if item == entry else item for item in result.manifest.evidence)})
    forged = replace(result, payload=payload, manifest=manifest)
    ir = world.artifact_irs[0]
    table = ir.sections[0].table
    row = table.rows[0]
    table = table.model_copy(update={"rows": [row.model_copy(update={"cells": {**row.cells,
        "variance": row.cells["variance"].model_copy(update={"value": 999})}})]})
    world = replace(world, _artifact_irs=(ir.model_copy(update={"sections": [ir.sections[0].model_copy(update={"table": table})]}),))
    tasks = tuple(task.model_copy(update={"inputs": tuple(item.model_copy(update={"sha256": snapshot.sha256})
        for item in task.inputs)}) for task in original_tasks)
    with pytest.raises(ValueError, match="formula literal disagrees with its operands"):
        plan_native_workload(world, {"reconciliation": forged}, _plan(formats=(format,)))
    with pytest.raises(ValueError, match="formula literal disagrees with its operands"):
        native_task_cases(tasks, {"reconciliation": forged}, namespace="northstar/seed-8128", world=world)


@pytest.mark.parametrize("format", ["docx", "pptx", "xlsx"])
def test_sdk_intake_rejects_relabeling_a_derived_table_unit(format: str) -> None:
    world, result = _table_world(format)
    entries = tuple(entry.model_copy(update={"unit": "USD"}) if entry.kind == "formula" else entry
        for entry in result.manifest.evidence)
    with pytest.raises(ValueError, match="unit disagrees with its canonical computation"):
        plan_native_workload(world, {"reconciliation": replace(result, manifest=result.manifest.model_copy(
            update={"evidence": entries}))}, _plan(formats=(format,)))


def test_studio_discovery_accepts_structured_sources_without_manufacturing_prose() -> None:
    world, _ = _table_world()
    spec = preset().model_copy(update={"use_cases": (UseCase(id="revenue-review", title="Revenue review",
        objective="Inspect the reconciliation model."),)})
    request = NativeSuiteRequest(use_case_id="revenue-review", formats=("xlsx",), operations=("read",),
        query_style="discovery", discovery_scope="artifact", minimum_units=1, max_cases=1)
    result = propose(world, spec, request)
    assert result["summary"]["source_sections"] == 1
    assert result["summary"]["capability_coverage"]["native_formula_inspection"] == 1
    assert result["summary"]["reference_qualified"] == 3
    assert not result["summary"]["unsupported"]


def test_studio_discovery_reports_unrenderable_source_graph_without_false_tasks() -> None:
    world, _ = _table_world()
    ir = world.artifact_irs[0]
    table = ir.sections[0].table
    row = table.rows[0]
    invalid = row.cells["variance"].model_copy(update={"operands": ["actual", "missing"]})
    table = table.model_copy(update={"rows": [row.model_copy(update={"cells": {**row.cells, "variance": invalid}})]})
    world = replace(world, _artifact_irs=(ir.model_copy(update={"sections": [ir.sections[0].model_copy(update={"table": table})]}),))
    spec = preset().model_copy(update={"use_cases": (UseCase(id="revenue-review", title="Revenue review",
        objective="Inspect the reconciliation model."),)})
    request = NativeSuiteRequest(use_case_id="revenue-review", formats=("xlsx",), operations=("read",),
        query_style="discovery", discovery_scope="artifact", minimum_units=1, max_cases=1)
    result = propose(world, spec, request)
    assert not result["summary"]["tasks"]
    assert result["spec"] == spec.model_dump(mode="json")
    assert any("native_source_graph_unavailable" in finding["reason"] for finding in result["summary"]["unsupported"])


def test_arithmetic_keeps_within_source_values_in_every_actual_native_format() -> None:
    world, rendered = _world("business")
    workload = plan_native_workload(world, rendered, _plan(operations=("analyze",), discovery_scope="artifact"))
    assert {(task.inputs[0].format, task.assertions[0].calculation.operation) for task in workload.tasks} == {
        (format, operation) for format in ("docx", "pptx", "xlsx") for operation in ("difference", "ratio")}
    assert {task.inputs[0].artifact_id for task in workload.tasks} == set(rendered)
    assert all(len(task.inputs) == 1 for task in workload.tasks)
    payloads = {identifier: result.payload for identifier, result in rendered.items()}
    for task in workload.tasks:
        assert qualify_native_task(task, payloads).passed
        assert all(ref.locator not in task.prompt for ref in task.assertions[0].calculation.operands)


@pytest.mark.parametrize("format", ["docx", "pptx"])
@pytest.mark.parametrize("duplicate_record", [False, True])
def test_record_column_scope_separates_table_copies_but_rejects_duplicate_records(format: str, duplicate_record: bool) -> None:
    world, _ = _table_world(format)
    actual = world.facts[0]
    ir = world.artifact_irs[0]
    sections = [*ir.sections, ArtifactSection(heading="Actual register", body="Actual revenue: {{fact:" + actual.id + "}}",
        fact_ids=[actual.id])]
    if duplicate_record:
        sections.append(ArtifactSection(heading="Repeated actual", body="Reviewed actual revenue: {{fact:" + actual.id + "}}",
            fact_ids=[actual.id]))
    world = replace(world, _artifact_irs=(ir.model_copy(update={"sections": sections}),))
    plan = NativeCorpusPlan(artifact_id="mixed-evidence", format=format, title="Revenue", surface="business",
        contents=tuple(NativeContent(source_artifact_id=ir.id, section_index=index) for index in range(len(sections))))
    result = render_native_corpus(world, plan)
    workload = plan_native_workload(world, {plan.artifact_id: result}, _plan(formats=(format,), operations=("read",)))
    ambiguous = [finding for finding in workload.findings if finding.code == "ambiguous_numeric_selector"]
    assert len(ambiguous) == (2 if duplicate_record else 0)
    records = [task for task in workload.tasks if "source-record table" in task.prompt]
    assert len(records) == (0 if duplicate_record else 1)
    assert all(qualify_native_task(task, {plan.artifact_id: result.payload}).passed for task in workload.tasks)


@pytest.mark.parametrize("format", ["docx", "pptx", "xlsx"])
def test_authored_sum_reconciliation_uses_native_constituents_not_inferred_populations(format: str) -> None:
    world, _ = _world("business")
    actual, west, budget, west_budget = world.facts
    table = Table(key="totals", title="Store revenue totals", columns=[Column(key="actual", label="Actual"),
        Column(key="budget", label="Budget")], rows=[
            Row(key="east", label="East", cells={"actual": Cell(value=125, fact_id=actual.id), "budget": Cell(value=100, fact_id=budget.id)}),
            Row(key="west", label="West", cells={"actual": Cell(value=70, fact_id=west.id), "budget": Cell(value=60, fact_id=west_budget.id)}),
            Row(key="total", label="Total", cells={"actual": Cell(value=195, formula=FormulaKind.SUM, operands=["east", "west"]),
                "budget": Cell(value=160, formula=FormulaKind.SUM, operands=["east", "west"])})])
    ir = world.artifact_irs[0].model_copy(update={"sections": [ArtifactSection(heading="Revenue", table=table)]})
    world = replace(world, _artifact_irs=(ir,))
    plan = NativeCorpusPlan(artifact_id="store-totals", format=format, title="Revenue totals", surface="business",
        contents=(NativeContent(source_artifact_id=ir.id, section_index=0),))
    result = render_native_corpus(world, plan)
    workload = plan_native_workload(world, {plan.artifact_id: result}, _plan(formats=(format,), operations=("analyze",)))
    sums = [task for task in workload.tasks if task.assertions[0].calculation.operation == "sum"]
    assert len(sums) == workload.capability_coverage["authored_sum_reconciliation"] == 2
    payloads = {plan.artifact_id: result.payload}
    assert {reference_submission(task, payloads).answers[0].value for task in sums} == {"195", "160"}
    for task in sums:
        assert qualify_native_task(task, payloads).passed
        assert len(task.assertions[0].calculation.operands) == 2
        assert all(ref.locator not in task.prompt for ref in task.assertions[0].calculation.operands)
    scalar_world, scalar_rendered = _world("business")
    scalar = plan_native_workload(scalar_world, scalar_rendered, _plan(operations=("analyze",)))
    assert all(task.assertions[0].calculation.operation != "sum" for task in scalar.tasks)


def test_explicit_calculation_format_cells_spend_fixed_budget_and_report_source_deficits() -> None:
    from worldloom.native_requirements import BenchmarkRequirements, CoverageRequirement

    world, rendered = _world("business")
    requirements = BenchmarkRequirements(cells=tuple(CoverageRequirement(name=format + "-ratios", operation="analyze",
        format=format, calculation="ratio", scope="artifact", min_independent_units=2) for format in ("docx", "pptx", "xlsx")))
    plan = _plan(max_tasks=6, operations=("read", "analyze"), requirements=requirements)
    workload = plan_native_workload(world, rendered, plan)
    assert len(workload.tasks) == 6
    assert all(task.operation == "analyze" and task.assertions[0].calculation.operation == "ratio" for task in workload.tasks)
    assert {format: sum(task.inputs[0].format == format for task in workload.tasks) for format in ("docx", "pptx", "xlsx")} == {
        "docx": 2, "pptx": 2, "xlsx": 2}
    assert not any(finding.code == "requirement_coverage_deficit" for finding in workload.findings)
    short = plan_native_workload(world, rendered, plan.model_copy(update={"max_tasks": 2}))
    assert len(short.tasks) == 2
    assert len([finding for finding in short.findings if finding.code == "requirement_coverage_deficit"]) == 3


def test_derived_and_superseded_source_lineage_cannot_purchase_independent_arithmetic() -> None:
    world, rendered = _world("business")
    facts = list(world.facts)
    facts[2] = facts[2].model_copy(update={"derived_from": [facts[0].id]})
    world = replace(world, _facts=tuple(facts))
    workload = plan_native_workload(world, rendered, _plan(operations=("analyze",), discovery_scope="cross_artifact"))
    assert not workload.tasks
    assert any(finding.code == "independent_cross_artifact_evidence_missing" for finding in workload.findings)


def test_creation_format_requirement_selects_deliverable_format() -> None:
    from worldloom.native_requirements import BenchmarkRequirements, CoverageRequirement

    world, rendered = _world("business")
    rendered = {key: value for key, value in rendered.items() if key.endswith("docx")}
    requirements = BenchmarkRequirements(cells=(CoverageRequirement(name="deck", operation="create", format="pptx"),))
    workload = plan_native_workload(world, rendered, _plan(max_tasks=1, operations=("create",),
        formats=("docx", "pptx"), requirements=requirements))
    assert len(workload.tasks) == 1
    task = workload.tasks[0]
    assert {item.format for item in task.inputs} == {"docx"}
    assert task.output.format == "pptx"
    assert not any(finding.code == "requirement_coverage_deficit" for finding in workload.findings)


def test_period_and_budget_families_survive_format_copies_with_visible_periods() -> None:
    world, _ = _world("business")
    actual, _, budget, _ = world.facts
    later = actual.model_copy(update={"id": "FACT-LATER", "period": "2026-02", "value": Quantity(amount=130, unit="AUD")})
    facts = (actual, later, budget)
    ir = world.artifact_irs[0].model_copy(update={"sections": [ArtifactSection(heading=f"Revenue {index}",
        body="Revenue: {{fact:" + fact.id + "}}", fact_ids=[fact.id]) for index, fact in enumerate(facts)]})
    world = replace(world, _facts=facts, _artifact_irs=(ir,))
    plans = [NativeCorpusPlan(artifact_id=format + "-periods", format=format, title="Revenue periods", surface="business",
        contents=tuple(NativeContent(source_artifact_id=ir.id, section_index=index) for index in range(3)))
        for format in ("docx", "pptx", "xlsx")]
    rendered = {plan.artifact_id: render_native_corpus(world, plan) for plan in plans}
    workload = plan_native_workload(world, rendered, _plan(operations=("analyze",), discovery_scope="artifact"))
    for capability in ("actual_budget_variance", "actual_budget_ratio", "period_change", "period_ratio"):
        assert workload.capability_coverage[capability] == 3
    assert all(qualify_native_task(task, {key: result.payload for key, result in rendered.items()}).passed for task in workload.tasks)


@pytest.mark.parametrize("mismatch", ["period", "ancestor", "rounding"])
def test_authored_sum_does_not_overstate_incompatible_or_rounded_reconciliation(mismatch: str) -> None:
    world, _ = _world("business")
    first = world.facts[0].model_copy(update={"value": Quantity(amount=2.4, unit="AUD")})
    second = world.facts[1].model_copy(update={"value": Quantity(amount=1.4, unit="AUD")})
    if mismatch == "period":
        second = second.model_copy(update={"period": "2026-02"})
    if mismatch == "ancestor":
        second = second.model_copy(update={"derived_from": [first.id]})
    table = Table(key="total", title="Revenue", columns=[Column(key="value", label="Value")], rows=[
        Row(key="first", label="East", cells={"value": Cell(value=2.4, fact_id=first.id)}),
        Row(key="second", label="West", cells={"value": Cell(value=1.4, fact_id=second.id)}),
        Row(key="total", label="Total", cells={"value": Cell(value=4 if mismatch == "rounding" else 3.8,
            formula=FormulaKind.SUM, operands=["first", "second"], formula_decimal_places=0 if mismatch == "rounding" else None)})])
    ir = world.artifact_irs[0].model_copy(update={"sections": [ArtifactSection(heading="Revenue", table=table)]})
    world = replace(world, _facts=(first, second), _artifact_irs=(ir,))
    plan = NativeCorpusPlan(artifact_id="total", format="xlsx", title="Revenue", surface="business",
        contents=(NativeContent(source_artifact_id=ir.id, section_index=0),))
    workload = plan_native_workload(world, {"total": render_native_corpus(world, plan)},
        _plan(formats=("xlsx",), operations=("analyze",)))
    assert "authored_sum_reconciliation" not in workload.capability_coverage


def test_numeric_pair_probes_are_linear_and_result_budget_is_hard(monkeypatch: pytest.MonkeyPatch) -> None:
    from decimal import Decimal

    import worldloom.native_query_planning as planning
    from worldloom.native_tasks import NativeInput

    world, _ = _world()
    template = world.facts[0]
    facts = tuple(template.model_copy(update={"id": f"FACT-{role}-{index}", "subject": f"store:{index}",
        "kind": "financial.revenue." + role}) for role in ("actual", "budget") for index in range(1000))
    numeric = []
    for role in ("actual", "budget"):
        members = [fact for fact in facts if fact.kind.endswith(role)]
        for format in ("docx", "pptx", "xlsx"):
            identifier = role + "-" + format
            source = planning._Source(NativeInput(artifact_id=identifier, format=format, path=identifier + "." + format),
                {}, frozenset(fact.id for fact in members))
            numeric.extend(planning._Numeric(source, NativeCitation(artifact_id=identifier, locator="value:" + fact.id),
                fact, Decimal(125), fact.subject, fact.kind, fact.subject) for fact in members)
    original = planning._compatible_pair
    probes = 0

    def counted(*args, **kwargs):
        nonlocal probes
        probes += 1
        return original(*args, **kwargs)

    monkeypatch.setattr(planning, "_compatible_pair", counted)
    pairs = planning._pairs(numeric, budget=30)
    assert len(pairs) == 30
    assert probes < len(numeric) * 12
    assert {pair.first.source.input.format for pair in pairs} == {"docx", "pptx", "xlsx"}
    assert any(not pair.cross_artifact for pair in pairs)
    assert pairs == planning._pairs(list(reversed(numeric)), budget=30)
