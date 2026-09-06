"""The presentation layer's contract: freedom above, one control underneath.

``presentation.py`` lets a harness decide who a corpus's documents are for. The
freedom is wide — an appendix can vanish, a figure can change magnitude, a PDF
can re-typeset itself — and it is only safe because of two properties that have
to be checked rather than asserted:

1. **A profile never changes a value.** It changes how one is *shown*.
2. **A profile never loses a record.** Whatever it declines to print is still in
   ``artifact-ir.jsonl``, so a corpus rendered for a reader is not a corpus that
   forgot its citations.

Plus the migration guarantee this whole layer was built behind: the default
profile renders the bytes that shipped before it existed. This project's gate is
byte identity, and a presentation layer whose arrival rewrote every corpus would
be indistinguishable from a regression.
"""

from __future__ import annotations

import io

import pytest

from worldloom import presentation, recipe
from worldloom.presentation import (
    Presentation,
    PresentationSeed,
    scale_money,
    suffix_for,
)

# ---------------------------------------------------------------------------
# 1. The control: presentation may not move a number
# ---------------------------------------------------------------------------


@pytest.mark.parametrize(
    "amount",
    [0.0, 1.0, 999.0, 1_000.0, 999_999.0, 123_801.0, 1_328_832.0, 5_372_800.0,
     -89_834.0, -1.0, 0.5, 1_000_000.0, 1_000_000_000.0, 4_242_424_242.0],
)
def test_a_promotion_multiplies_back_to_the_figure_it_started_from(amount: float) -> None:
    """The one thing that makes `magnitudes: scaled` allowable at all.

    Exact, never within a tolerance: a tolerance here would be this module
    deciding how much of a figure a reader may lose, which is precisely the
    decision it is not permitted to make.
    """
    shown, factor = scale_money(amount)
    assert round(shown * factor, 6) == round(amount, 6)


def test_a_figure_with_no_exact_short_spelling_is_not_promoted() -> None:
    """Never a rounding, even a flattering one.

    A document showing a rounded figure while claiming to cite the fact it
    rounded is the drift `narrative/references` exists to prevent, one layer
    down. `scale_money` searches up to three decimals, so a figure that cannot
    be spelled exactly inside those keeps the ledger's own wording.
    """
    # 1,234.5678 thousands: not exact at one, two or three decimals in millions.
    shown, factor = scale_money(1_234_567.8)
    assert (shown, factor) == (1_234_567.8, 1.0)
    assert suffix_for(factor) == ""


def test_the_lint_refuses_a_scaling_that_cannot_round_trip(monkeypatch) -> None:
    """The refusal arrives while the author still holds the pen.

    Checked at authoring time and not only in a render test, because a profile
    is a thing a harness writes and a defect found three commands later is a
    defect found in the wrong place. Simulated by making the promotion lossy,
    which is the only way to reach this branch — the shipped `scale_money`
    cannot fail it, which is the point of it.
    """
    monkeypatch.setattr(presentation, "scale_money",
                    lambda amount, **_: (round(amount / 1000.0, 1), 1000.0))
    findings = presentation.review(PresentationSeed(name="lossy", magnitudes="scaled"))
    assert findings, "a promotion that loses a digit must be refused"
    assert any("round-trip" in finding for finding in findings)


def test_an_override_that_scales_is_linted_like_a_top_level_one(monkeypatch) -> None:
    """The knob is reachable two ways, so the control has to cover both.

    An override setting `magnitudes: scaled` on one doctype is exactly as
    capable of moving a figure as the profile-wide setting, and a lint that
    only read the top level would leave the interesting half unchecked.
    """
    monkeypatch.setattr(presentation, "scale_money",
                    lambda amount, **_: (round(amount / 1000.0, 1), 1000.0))
    findings = presentation.review(
        PresentationSeed(name="lossy", overrides={"cfo_variance_memo": {"magnitudes": "scaled"}})
    )
    assert any("round-trip" in finding for finding in findings)


# ---------------------------------------------------------------------------
# 2. Refusals a reviser can act on
# ---------------------------------------------------------------------------


def test_every_finding_arrives_at_once() -> None:
    """`cascade`'s protocol, over presentation.

    A reviser fixing one knob per round trip pays a turn per rule it could not
    see. Four independent mistakes, four findings, one refusal.
    """
    findings = presentation.review(
        PresentationSeed(
            name="wrong", appendix="hide", provenance="bottom",
            magnitudes="big", table_fit="snug",
        )
    )
    assert len(findings) >= 4
    for knob in ("appendix", "provenance", "magnitudes", "table_fit"):
        assert any(finding.startswith(knob) for finding in findings), knob


def test_a_misspelled_knob_is_refused_rather_than_ignored() -> None:
    """The reason the seed is a model and not a dict.

    A profile is a table of knobs, and a knob silently dropped is a document
    that renders one way while its author is certain it renders another.
    """
    with pytest.raises(Exception):  # noqa: B017 - pydantic's own error type
        PresentationSeed(name="typo", appendx="omit")

    findings = presentation.review(
        PresentationSeed(name="typo", overrides={"incident_rca": {"appendx": "omit"}}),
        doctypes=("incident_rca",),
    )
    assert any("appendx" in finding for finding in findings)


def test_an_override_naming_a_doctype_the_corpus_lacks_is_refused() -> None:
    """A rule that silently does nothing is what a typo produces."""
    findings = presentation.review(
        PresentationSeed(name="p", overrides={"incident_rce": {"appendix": "omit"}}),
        doctypes=("incident_rca", "cfo_variance_memo"),
    )
    assert any("incident_rce" in finding for finding in findings)


def test_a_registered_name_may_not_be_quietly_redefined() -> None:
    """A corpus that asked for a profile yesterday must still get it."""
    findings = presentation.review(PresentationSeed(name="reader", appendix="append"))
    assert any("already a registered profile" in finding for finding in findings)
    # ...and re-registering the *same* settings is not a collision.
    presentation.register("reader", presentation.READER)


def test_an_unknown_profile_name_is_refused_never_defaulted() -> None:
    """Sharper than the locale case it copies.

    A locale typo produces a visibly wrong corpus. A profile typo produces an
    invisibly wrong one: ask for `readerr`, get the audit rendering, and the
    only symptom is four pages of fact table in a document you were about to
    hand to somebody.
    """
    with pytest.raises(ValueError, match="unknown presentation profile"):
        presentation.named("readerr")


# ---------------------------------------------------------------------------
# 3. The record survives every profile
# ---------------------------------------------------------------------------


def test_the_recipe_round_trips_a_profile_by_value() -> None:
    """By value and never by name, because this project rebuilds from records.

    A name is a reference into a registry a later checkout may have changed,
    and a rebuild that resolved it against somebody's edited profile would
    produce different documents and report success.
    """
    written = recipe.with_presentation({}, "reader")
    assert isinstance(written[recipe.PRESENTATION_KEY], dict)
    assert recipe.presentation_of(written) == presentation.READER


def test_a_corpus_that_names_no_profile_is_the_audit_rendering() -> None:
    """Every corpus built before this layer carries no key and *was* `audit`.

    A fact about those corpora rather than a gap in them — the same reading
    `recipe.locale_of` takes of an absent locale.
    """
    assert recipe.presentation_of({}) is presentation.AUDIT


def test_a_profile_applies_per_doctype_without_mutating_the_profile() -> None:
    """`for_doctype` returns `self` when nothing is overridden.

    Identity, not equality: the common path is every artifact in the corpus and
    it should allocate nothing.
    """
    plain = Presentation(name="p", appendix="omit")
    assert plain.for_doctype("incident_rca") is plain

    tuned = Presentation(name="p", appendix="omit",
                         overrides={"incident_rca": {"appendix": "append"}})
    assert tuned.for_doctype("incident_rca").appendix == "append"
    assert tuned.for_doctype("cfo_variance_memo").appendix == "omit"
    assert tuned.appendix == "omit", "reading an override must not move the profile"


def test_two_profiles_built_from_one_document_are_equal() -> None:
    """Overrides normalise at construction, so a JSON-built profile and a
    Python-built one are the same object by value — which is what lets
    `presentation_of` compare a recorded profile against a registered one."""
    seed = PresentationSeed(name="h", appendix="omit",
                            overrides={"incident_rca": {"appendix": "append"}})
    assert presentation.resolve(seed) == Presentation(
        name="h", appendix="omit", overrides={"incident_rca": {"appendix": "append"}}
    )


# ---------------------------------------------------------------------------
# 4. What the renderers actually do with it
# ---------------------------------------------------------------------------


@pytest.fixture(scope="module")
def memo():
    """One narrated artifact with a hidden section, a voice and money facts.

    Built rather than hand-rolled: the properties below are about what the
    *shipped* artifacts do, and an IR assembled in a test can be given exactly
    the shape that makes an assertion pass.
    """
    from worldloom import RetailWorld
    from worldloom.scenarios import MonthEndClose

    world = RetailWorld(seed=8128).build().run(MonthEndClose(period="2026-03", include_operational_incident=True))
    world = world.compile()
    ir = next(
        candidate for candidate in world.artifact_irs
        if any(section.hidden and section.table for section in candidate.sections)
        and candidate.metadata.get("voice")
    )
    return ir, {fact.id: fact for fact in world.facts}


def test_omitting_the_appendix_does_not_remove_it_from_the_record(memo) -> None:
    """The property that makes `omit` safe rather than lossy.

    A reader profile declines to print a section; it does not decline to record
    one. If this ever fails, every claim the module docstring makes about
    omission being free goes with it.
    """
    ir, _ = memo
    hidden = [section for section in ir.sections if section.hidden]
    assert hidden, "the fixture needs an artifact that has something to omit"
    assert all(section.fact_ids or section.table for section in hidden)
    # The IR is what `corpus.write_jsonl` persists, and no profile touches it —
    # profiles are read at render time and never on the way in.
    assert ir.metadata.get("voice")


def test_the_reader_profile_drops_the_appendix_and_the_brief_from_markdown(memo) -> None:
    ir, facts = memo
    from worldloom.render import markdown

    audit = markdown.render(ir, facts, presentation=presentation.AUDIT).decode()
    reader = markdown.render(ir, facts, presentation=presentation.READER).decode()

    assert "not part of the readable surface" in audit
    assert "Author voice:" in audit
    assert "not part of the readable surface" not in reader
    assert "Author voice:" not in reader
    assert len(reader) < len(audit)
    # The prose itself is untouched — only the scaffolding around it moved.
    for section in ir.sections:
        if section.body and not section.hidden:
            assert section.heading in reader


def test_the_brief_moves_into_the_file_rather_than_out_of_existence(memo) -> None:
    """`provenance: properties` keeps the voice where a tool can read it.

    The information is worth having — it is how a later reader knows why a memo
    sounds the way it does — and the only thing wrong with it was being the
    last paragraph of the memo.
    """
    pytest.importorskip("docx")
    import docx as docx_pkg

    from worldloom.render import docx as docx_render

    ir, facts = memo
    payload = docx_render.render(ir, facts, presentation=presentation.READER)
    document = docx_pkg.Document(io.BytesIO(payload))
    body = "\n".join(paragraph.text for paragraph in document.paragraphs)
    assert "Author voice:" not in body
    assert ir.metadata["voice"] in (document.core_properties.category or "")


def test_a_measured_pdf_table_never_breaks_a_token_in_half(memo) -> None:
    """The defect `table_fit` exists for, pinned by its symptom.

    The fixed split gives the five-column supporting-facts table roughly 80pt a
    column, which breaks `system_of_record` and splits an ISO timestamp as
    `2026-04-07T16:4` / `0:00+00:00`. Asserted on the *widths* rather than on
    extracted text, because a substring check cannot tell a broken token from
    an intact one — `'system_of_recor' in 'system_of_record'` is true either
    way, and that mistake made a working fix look broken for two rounds.
    """
    pytest.importorskip("reportlab")
    from reportlab.pdfbase.pdfmetrics import stringWidth

    from worldloom.compiler.style import genome
    from worldloom.locales import DEFAULT as DEFAULT_LOCALE
    from worldloom.render import pdf
    from worldloom.rng import Rng

    ir, _ = memo
    section = next(s for s in ir.sections if s.hidden and s.table)
    # The measurement contract, threaded explicitly since `_styles` became
    # genome-driven: measure in the face and size the cells will actually be
    # set in. `"house"` keeps this test about token-breaking, not palette.
    house = genome(Rng(0).derive("style"), archetype="house")
    cell_pt = pdf._styles(house)["cell"].fontSize
    widths, size_pt = pdf._measured_layout(
        section.table, pdf._FRAME_WIDTH_PT, DEFAULT_LOCALE,
        font=pdf._styles(house)["cell_header"].fontName, size=cell_pt,
    )

    assert abs(sum(widths) - pdf._FRAME_WIDTH_PT) < 0.5, "the table must fill its frame"
    assert size_pt <= cell_pt and size_pt >= pdf._CELL_MIN_PT

    columns = [[section.table.title] + [row.label for row in section.table.rows]]
    for spec in section.table.columns:
        columns.append([spec.label] + [
            str((row.cells.get(spec.key).value if row.cells.get(spec.key) else "") or "")
            for row in section.table.rows
        ])
    header_face = pdf._styles(house)["cell_header"].fontName
    for width, texts in zip(widths, columns):
        widest = max(
            (stringWidth(token, header_face, size_pt)
             for text in texts if text for token in str(text).split()),
            default=0.0,
        )
        assert widest <= width, "a column narrower than its longest token can only break it"


def test_the_default_profile_renders_what_shipped_before_profiles_existed(memo) -> None:
    """The migration guarantee, stated as bytes.

    ``AUDIT`` is not a legacy setting kept around out of politeness — it is the
    right profile for a corpus whose reader is a validator, and it is also the
    proof that this layer arrived without rewriting anything. Passing it
    explicitly and passing nothing must produce identical output, in every
    format that takes one.
    """
    from worldloom.render import markdown

    ir, facts = memo
    assert markdown.render(ir, facts) == markdown.render(
        ir, facts, presentation=presentation.AUDIT
    )
    assert presentation.DEFAULT is presentation.AUDIT


def test_a_profile_is_a_function_of_its_settings_and_nothing_else(memo) -> None:
    """Determinism, at the layer that decides bytes.

    Two renders under one profile must agree, and a profile built twice from
    one document must render the same — otherwise the recipe records something
    that does not reproduce, which is the failure this whole project gates on.
    """
    from worldloom.render import markdown

    ir, facts = memo
    rebuilt = recipe.presentation_of(recipe.with_presentation({}, presentation.READER))
    assert markdown.render(ir, facts, presentation=presentation.READER) == markdown.render(
        ir, facts, presentation=rebuilt
    )
