"""A reader-grade Word document is checked by reading its structure.

There is no office suite to render with in CI (LibreOffice will not load a
file in this container), so fidelity is verified on the package itself with
python-docx and the XML under it: the styles headings use, keep-with-next on
headings and captions, table header rows that repeat across pages and rows
that do not split, the custom properties that carry provenance off the page,
native review comments, the section break between front matter and body
(numbered from one), and the fields in the running header and footer.
"""

from __future__ import annotations

from io import BytesIO

import pytest

from worldloom import realism_profiles
from worldloom.narrative import ComposedProvider
from worldloom.render.ooxml import custom_properties
from worldloom.retail import RetailWorld
from worldloom.scenarios import MonthEndClose

W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"


@pytest.fixture(scope="module")
def rendered():  # type: ignore[no-untyped-def]
    world = RetailWorld(seed=8128).build().run(MonthEndClose(period="2026-03", include_operational_incident=True))
    planned = world.extend(recipe=realism_profiles.with_realism(world.recipe, "enterprise")).compile()
    return planned.narrate(ComposedProvider.for_world(planned)).render("docx")


def _payloads(world, artifact_type: str) -> list[tuple[str, bytes]]:  # type: ignore[no-untyped-def]
    ids = {ir.id for ir in world.artifact_irs
           if world.artifact_intents.by_id(ir.intent_id).artifact_type == artifact_type}
    return [(item.path, bytes(item.payload)) for item in world._rendered
            if item.artifact_id in ids and item.path.endswith(".docx")]


@pytest.fixture(scope="module")
def memo(rendered):  # type: ignore[no-untyped-def]
    from docx import Document

    _path, payload = next((p, b) for p, b in _payloads(rendered, "cfo_variance_memo") if "/revisions/" not in p)
    return Document(BytesIO(payload)), payload


def test_headings_use_word_styles_and_keep_with_what_follows(memo) -> None:  # type: ignore[no-untyped-def]
    document, _ = memo
    styles = {s.name for s in document.styles}
    assert {"Title", "Heading 1", "Heading 2", "Normal"} <= styles
    headings = [p for p in document.paragraphs if p.style.name.startswith("Heading")]
    assert len(headings) >= 4
    # Word keeps a heading with its next paragraph through the style (the
    # built-in heading styles set it) or on the paragraph itself.
    for heading in headings:
        style_keep = heading.style.paragraph_format.keep_with_next
        assert heading.paragraph_format.keep_with_next or style_keep is not False, heading.text
    captions = [p for p in document.paragraphs if p.paragraph_format.keep_with_next and not p.style.name.startswith("Heading")]
    assert captions, "no caption or table lead-in is kept with its table"


def test_table_header_rows_repeat_and_rows_do_not_split(memo) -> None:  # type: ignore[no-untyped-def]
    document, _ = memo
    multi = [t for t in document.tables if len(t.rows) > 2]
    assert multi
    for table in multi:
        first = table.rows[0]._tr
        assert first.find(f"{W}trPr/{W}tblHeader") is not None
        assert all(row._tr.find(f"{W}trPr/{W}cantSplit") is not None for row in table.rows)


def test_provenance_rides_in_the_custom_properties(memo) -> None:  # type: ignore[no-untyped-def]
    _, payload = memo
    properties = custom_properties(payload)
    assert properties["WorldloomProfile"] == "reader"
    assert properties["WorldloomFacts"].startswith("FACT-")


def test_front_matter_and_body_are_two_sections_numbered_from_one(memo) -> None:  # type: ignore[no-untyped-def]
    document, _ = memo
    sections = document.sections
    assert len(sections) == 2
    assert sections[0].different_first_page_header_footer
    body = sections[1]
    start = body._sectPr.find(f"{W}pgNumType")
    assert start is not None and start.get(f"{W}start") == "1"
    assert body.header.is_linked_to_previous and body.footer.is_linked_to_previous


def test_running_heads_carry_fields(memo) -> None:  # type: ignore[no-untyped-def]
    document, _ = memo
    head = document.sections[0].header._element.xml
    foot = document.sections[0].footer._element.xml
    assert "STYLEREF" in head
    assert " PAGE " in foot and " NUMPAGES " in foot
    contents = "".join(p._p.xml for p in document.paragraphs[:80])
    assert " TOC " in contents


def test_a_reviewed_revision_carries_native_comments(rendered) -> None:  # type: ignore[no-untyped-def]
    from zipfile import ZipFile

    commented = 0
    for _path, payload in _payloads(rendered, "cfo_variance_memo") + _payloads(rendered, "incident_rca"):
        with ZipFile(BytesIO(payload)) as archive:
            if "word/comments.xml" in archive.namelist():
                commented += "<w:comment " in archive.read("word/comments.xml").decode("utf-8")
    assert commented


def test_the_plain_layout_keeps_one_section() -> None:
    """The section break is the designed layout's; an audit rendering keeps its bytes."""
    from docx import Document

    world = RetailWorld(seed=8128).build().run(MonthEndClose(period="2026-03", include_operational_incident=True))
    planned = world.extend(recipe=realism_profiles.with_realism(world.recipe, "enterprise/v1")).compile()
    done = planned.narrate(ComposedProvider()).render("docx")
    _path, payload = next((p, b) for p, b in _payloads(done, "cfo_variance_memo") if "/revisions/" not in p)
    assert len(Document(BytesIO(payload)).sections) == 1
