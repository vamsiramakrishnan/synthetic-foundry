"""Native DAG receipts must describe actual file transitions and ordered calls."""
from __future__ import annotations

import base64
import hashlib
import json
import re
import sys
from collections.abc import Mapping
from datetime import UTC, datetime
from io import BytesIO
from pathlib import Path
from typing import Any

import pytest

from worldloom.benchmarks.core import NativeBenchmark
from worldloom.benchmarks.runner import CallableHarness, CommandHarness
from worldloom.benchmarks.workflows import (
    NativeWorkflowInput,
    NativeWorkflowPlan,
    NativeWorkflowStep,
    qualify_workflow,
    run_workflow,
    validate_workflow,
    workflow_capabilities,
)
from worldloom.corpus import write_json
from worldloom.models import (
    ArtifactIR,
    ArtifactSection,
    Authority,
    CanonicalFact,
    Company,
    Quantity,
)
from worldloom.native_corpus import (
    NativeContent,
    NativeCorpusPlan,
    render_native_corpus,
)
from worldloom.native_query_planning import NativeWorkloadPlan
from worldloom.native_tasks import (
    NativeAnswer,
    NativeCitation,
    NativeFile,
    NativeSubmission,
)
from worldloom.providers import digest
from worldloom.render.ooxml import normalise
from worldloom.world import World


@pytest.fixture(scope="module")
def benchmark() -> NativeBenchmark:
    pytest.importorskip("docx")
    pytest.importorskip("openpyxl")
    facts = tuple(CanonicalFact(id=f"FACT-{kind}-{subject}".upper(), kind="financial.revenue." + kind,
        subject="store:" + subject, period="2026-01", value=Quantity(amount=value, unit="AUD"),
        valid_from=datetime(2026, 1, 1, tzinfo=UTC), authority=Authority.SYSTEM_OF_RECORD)
        for kind, subject, value in (("actual", "east", 125), ("actual", "west", 70),
                                     ("budget", "east", 100), ("budget", "west", 60)))
    sections = [ArtifactSection(heading=f"{fact.subject} {fact.kind.rsplit('.', 1)[1]} revenue",
        body=f"Revenue evidence for {fact.subject}: {{{{fact:{fact.id}}}}}.", fact_ids=[fact.id]) for fact in facts]
    world = World(seed=8128, company=Company(id="CO-1", name="Northstar Retail", industry="retail",
        headquarters="Sydney", fiscal_year_start_month=7, employees_total=100), _facts=facts,
        _artifact_irs=(ArtifactIR(id="ART-SOURCE", intent_id="INTENT-1", title="Revenue close", sections=sections),))
    plans = tuple(NativeCorpusPlan(artifact_id=f"{group}-{format}", format=format, title=f"{group} revenue close",
        minimum_units=2, contents=tuple(NativeContent(source_artifact_id="ART-SOURCE", section_index=index) for index in indexes))
        for group, indexes in (("actual", (0, 1)), ("budget", (2, 3))) for format in ("docx", "xlsx"))
    rendered = {plan.artifact_id: render_native_corpus(world, plan) for plan in plans}
    return NativeBenchmark.from_rendered(world, rendered, NativeWorkloadPlan(use_case_id="workflow",
        # Fund the explicit DOCX/XLSX workflow alongside the planner's ratio
        # and sum families; a small general sampling budget is not a fixture.
        objective="Reconcile revenue evidence.", formats=("docx", "xlsx"), max_tasks=24))


@pytest.fixture
def plan(benchmark: NativeBenchmark) -> NativeWorkflowPlan:
    tasks = benchmark.workload.tasks
    read_doc = next(task for task in tasks if task.operation == "read" and len(task.inputs) == 1
        and task.inputs[0].artifact_id == "budget-docx" and "store:east" in task.prompt)
    analyze = next(task for task in tasks if task.operation == "analyze"
        and task.assertions[0].calculation.operation == "difference")
    update_doc = next(task for task in tasks if task.operation == "update" and task.output.source_artifact_id == "budget-docx")
    update_sheet = next(task for task in tasks if task.operation == "update" and len(task.inputs) == 1
        and task.output.source_artifact_id == "actual-xlsx")
    read_sheet = next(task for task in tasks if task.operation == "read" and len(task.inputs) == 1
        and task.inputs[0].artifact_id == "actual-xlsx" and "store:east" in task.prompt)
    create = next(task for task in tasks if task.operation == "create" and task.output.format == "docx")
    return NativeWorkflowPlan(id="revenue-close", benchmark_digest=benchmark.digest, steps=(
        NativeWorkflowStep(id="read", task_id=read_doc.id),
        NativeWorkflowStep(id="analyze", task_id=analyze.id, depends_on=("read",)),
        NativeWorkflowStep(id="update-doc", task_id=update_doc.id, depends_on=("analyze",)),
        NativeWorkflowStep(id="verify-doc", task_id=read_doc.id, depends_on=("update-doc",),
            input_bindings=(NativeWorkflowInput(input_artifact_id="budget-docx", parent_step_id="update-doc"),)),
        NativeWorkflowStep(id="update-sheet", task_id=update_sheet.id, depends_on=("analyze",)),
        NativeWorkflowStep(id="verify-sheet", task_id=read_sheet.id, depends_on=("update-sheet",),
            input_bindings=(NativeWorkflowInput(input_artifact_id="actual-xlsx", parent_step_id="update-sheet"),)),
        NativeWorkflowStep(id="create", task_id=create.id, depends_on=("verify-doc", "verify-sheet")),
    ))


def _target(request: dict[str, Any], inputs: Mapping[str, bytes]) -> NativeSubmission:
    """A public-only document worker: parse instructions and inspect actual files."""
    from docx import Document
    from openpyxl import load_workbook

    assert "assertions" not in request and "expected" not in request
    assert set(inputs) == {item["artifact_id"] for item in request["inputs"]}
    assert all(set(parent) == {"step_id", "execution_id", "submission_digest", "answers"}
        for parent in request["workflow"]["parents"])
    documents: dict[str, Any] = {}
    evidence: dict[str, tuple[str, str, str]] = {}
    for item in request["inputs"]:
        key = item["artifact_id"]
        assert hashlib.sha256(inputs[key]).hexdigest() == item["sha256"]
        if item["format"] == "docx":
            document = Document(BytesIO(inputs[key]))
            documents[key] = document
            for index, paragraph in enumerate(document.paragraphs[:-1]):
                if paragraph.text.startswith("store:"):
                    evidence[paragraph.text] = (key, f"paragraph:{index + 2}", document.paragraphs[index + 1].text)
        else:
            document = load_workbook(BytesIO(inputs[key]))
            documents[key] = document
            for row in document["Evidence"].iter_rows(min_row=2):
                evidence[str(row[0].value)] = (key, f"sheet:Evidence/cell:{row[1].coordinate}", str(row[1].value))
    if request["operation"] == "analyze":
        selectors = re.findall(r"measure '([^']+)' for subject '([^']+)' in reporting period '([^']+)' in '([^']+)'", request["prompt"])
        values, citations = [], []
        for measure, subject, _period, key in selectors:
            workbook = documents[key]
            row = next(row for row in workbook["Facts"].iter_rows(min_row=2) if row[3].value == measure and row[4].value == subject)
            values.append(row[1].value)
            citations.append(NativeCitation(artifact_id=key, locator=f"sheet:Facts/cell:{row[1].coordinate}"))
        return NativeSubmission(answers=(NativeAnswer(assertion_id="difference", value=str(values[0] - values[1]), citations=tuple(citations)),))
    headings = re.findall(r"section headed '([^']+)'", request["prompt"])
    found = [evidence[heading] for heading in headings]
    answers = tuple(NativeAnswer(assertion_id=shape["assertion_id"], value=value,
        citations=(NativeCitation(artifact_id=key, locator=locator),))
        for shape, (key, locator, value) in zip(request["answers"], found, strict=True))
    if request["output"] is None:
        return NativeSubmission(answers=answers)
    output = request["output"]
    source_hash = None
    if request["operation"] == "update":
        source_id = output["source_artifact_id"]
        document = documents[source_id]
        updated = found[0][2] + "\nRelated evidence: " + found[1][2]
        locator = found[0][1]
        if output["format"] == "docx":
            document.paragraphs[int(locator.split(":")[1]) - 1].text = updated
        else:
            sheet, cell = re.fullmatch(r"sheet:([^/]+)/cell:(.+)", locator).groups()
            document[sheet][cell] = updated
        source_hash = hashlib.sha256(inputs[source_id]).hexdigest()
    else:
        document = Document()
        for heading, (_, _, value) in zip(headings, found, strict=True):
            document.add_paragraph(heading)
            document.add_paragraph(value)
    stream = BytesIO()
    document.save(stream)
    payload = normalise(stream.getvalue(), created="1980-01-01T00:00:00Z")
    return NativeSubmission(answers=answers, files=(NativeFile(artifact_id=output["artifact_id"], format=output["format"],
        content_base64=base64.b64encode(payload).decode("ascii"), source_sha256=source_hash),))


def _receipt(directory: Path, step: str) -> Path:
    return directory / "receipts" / (digest([step]) + ".json")


def _rewrite(path: Path, document: dict[str, Any]) -> None:
    document["digest"] = digest({key: value for key, value in document.items() if key != "digest"})
    write_json(path, document)


def test_real_docx_and_xlsx_dag_passes_new_bytes_and_exact_resume_is_identical(
    benchmark: NativeBenchmark, plan: NativeWorkflowPlan, tmp_path: Path,
) -> None:
    calls: list[dict[str, Any]] = []
    def target(request: dict[str, Any], inputs: Mapping[str, bytes]) -> NativeSubmission:
        calls.append(request)
        return _target(request, inputs)
    harness = CallableHarness(target, {"worker": "public-native/v1"})
    assert qualify_workflow(benchmark, plan).passed
    assert qualify_workflow(benchmark, plan) == qualify_workflow(benchmark, plan)
    output = tmp_path / "run"
    report = run_workflow(benchmark, plan, harness, directory=output)
    assert report.passed and report.passed_count == report.total == 7
    assert [call["workflow"]["step_id"] for call in calls] == [step.id for step in plan.steps]
    for step_id in ("verify-doc", "verify-sheet"):
        step = next(item for item in report.steps if item.step_id == step_id)
        lineage = step.lineage[0]
        assert lineage.sha256 != lineage.source_input_sha256
        parent = json.loads(_receipt(output, lineage.parent_step_id).read_text())
        child = json.loads(_receipt(output, step_id).read_text())
        payload = base64.b64decode(parent["submission"]["files"][0]["content_base64"])
        assert lineage.sha256 == hashlib.sha256(payload).hexdigest()
        assert "Related evidence:" in child["submission"]["answers"][0]["value"]
    before = {path.relative_to(output): path.read_bytes() for path in output.rglob("*") if path.is_file()}
    assert run_workflow(benchmark, plan, harness, directory=output, resume=True) == report
    assert len(calls) == 7
    assert before == {path.relative_to(output): path.read_bytes() for path in output.rglob("*") if path.is_file()}
    assert report.observation == "runner_scheduled_calls_and_staged_bytes"


def test_interrupt_reissues_only_uncommitted_step_and_preserves_execution_identity(
    benchmark: NativeBenchmark, plan: NativeWorkflowPlan, tmp_path: Path,
) -> None:
    calls: list[str] = []
    def target(request: dict[str, Any], inputs: Mapping[str, bytes]) -> NativeSubmission:
        calls.append(request["execution_id"])
        if len(calls) == 4:
            raise KeyboardInterrupt()
        return _target(request, inputs)
    harness = CallableHarness(target, {"worker": "interrupt/v1"})
    output = tmp_path / "run"
    with pytest.raises(KeyboardInterrupt):
        run_workflow(benchmark, plan, harness, directory=output)
    assert len(list((output / "receipts").glob("*.json"))) == 3
    assert run_workflow(benchmark, plan, harness, directory=output, resume=True).passed
    assert len(calls) == 8 and calls[3] == calls[4]


def test_existing_command_protocol_stages_derived_bytes_and_preserves_workflow_context(
    benchmark: NativeBenchmark, plan: NativeWorkflowPlan, tmp_path: Path,
) -> None:
    command_plan = plan.model_copy(update={"steps": (
        plan.steps[0], plan.steps[2].model_copy(update={"depends_on": ("read",)}), plan.steps[3])})
    script = tmp_path / "native-worker.py"
    script.write_text("""import json, sys
from pathlib import Path
sys.path.insert(0, sys.argv[1])
from test_native_workflows import _target
payload = json.load(sys.stdin)
assert payload['schema'] == 'worldloom.native-harness-request/v1'
task = payload['task']
root = Path(payload['input_root'])
assert root == Path.cwd()
assert not list(root.rglob('manifest.json'))
inputs = {item['artifact_id']: (root / item['path']).read_bytes() for item in task['inputs']}
submission = _target(task, inputs)
print(json.dumps({'schema': 'worldloom.native-harness-response/v1', 'task_id': task['id'],
    'execution_id': task['execution_id'], 'submission': submission.model_dump(mode='json')}))
""", encoding="utf-8")
    harness = CommandHarness((sys.executable, str(script), str(Path(__file__).parent)), identity_files=(Path(__file__),))
    report = run_workflow(benchmark, command_plan, harness, directory=tmp_path / "command-run")
    assert report.passed and report.total == 3
    assert report.steps[-1].lineage[0].sha256 != report.steps[-1].lineage[0].source_input_sha256


@pytest.mark.parametrize("tamper", ["answer", "lineage", "order", "parent", "missing"])
def test_rehashed_receipt_tampering_is_rejected_before_any_calls(
    benchmark: NativeBenchmark, plan: NativeWorkflowPlan, tmp_path: Path, tamper: str,
) -> None:
    calls: list[str] = []
    def target(request: dict[str, Any], inputs: Mapping[str, bytes]) -> NativeSubmission:
        calls.append(request["execution_id"])
        return _target(request, inputs)
    harness = CallableHarness(target, {"worker": "tamper/v1"})
    output = tmp_path / "run"
    run_workflow(benchmark, plan, harness, directory=output)
    path = _receipt(output, "verify-doc")
    saved = json.loads(path.read_text())
    if tamper == "answer":
        saved["submission"]["answers"][0]["value"] = "invented success"
    elif tamper == "lineage":
        saved["result"]["lineage"][0]["sha256"] = "0" * 64
    elif tamper == "order":
        saved["pins"]["index"] = 0
    elif tamper == "parent":
        saved["pins"]["parents"] = {"read": "0" * 64}
    elif tamper == "missing":
        _receipt(output, "update-doc").unlink()
    _rewrite(path, saved)
    with pytest.raises(ValueError, match=r"grading|lineage|order|predecessor"):
        run_workflow(benchmark, plan, harness, directory=output, resume=True)
    assert len(calls) == 7


def test_mutated_parent_file_fails_and_blocks_descendants_while_independent_branch_runs(
    benchmark: NativeBenchmark, plan: NativeWorkflowPlan, tmp_path: Path,
) -> None:
    from docx import Document

    calls: list[str] = []
    def target(request: dict[str, Any], inputs: Mapping[str, bytes]) -> NativeSubmission:
        step = request["workflow"]["step_id"]
        calls.append(step)
        submission = _target(request, inputs)
        if step == "update-doc":
            file = submission.files[0]
            document = Document(BytesIO(base64.b64decode(file.content_base64)))
            document.paragraphs[0].text = "An unauthorized heading change"
            stream = BytesIO()
            document.save(stream)
            submission = submission.model_copy(update={"files": (file.model_copy(update={
                "content_base64": base64.b64encode(stream.getvalue()).decode("ascii")}),)})
        return submission
    harness = CallableHarness(target, {"worker": "mutate/v1"})
    report = run_workflow(benchmark, plan, harness, directory=tmp_path / "run")
    assert not report.passed and report.blocked_count == 2
    assert calls == ["read", "analyze", "update-doc", "update-sheet", "verify-sheet"]
    assert "update_unaffected_content_changed" in report.steps[2].grade.findings
    assert report.steps[3].status == report.steps[6].status == "blocked"
    assert not report.steps[3].lineage
    assert run_workflow(benchmark, plan, harness, directory=tmp_path / "run", resume=True) == report


def test_wrong_intermediate_answer_blocks_required_steps_even_if_writes_would_pass(
    benchmark: NativeBenchmark, plan: NativeWorkflowPlan, tmp_path: Path,
) -> None:
    def target(request: dict[str, Any], inputs: Mapping[str, bytes]) -> NativeSubmission:
        submission = _target(request, inputs)
        if request["operation"] == "analyze":
            submission = submission.model_copy(update={"answers": (submission.answers[0].model_copy(update={"value": "999"}),)})
        return submission
    report = run_workflow(benchmark, plan, CallableHarness(target, {"worker": "wrong-analysis/v1"}), directory=tmp_path / "run")
    assert report.passed_count == 1 and report.blocked_count == 5
    assert report.steps[1].grade.findings == ("answer_failed:difference",)


def test_schema_and_source_bound_task_refs_refuse_fabricated_graphs(
    benchmark: NativeBenchmark, plan: NativeWorkflowPlan,
) -> None:
    with pytest.raises(ValueError, match="dependency order"):
        NativeWorkflowPlan(id="wrong-order", benchmark_digest=benchmark.digest, steps=tuple(reversed(plan.steps)))
    with pytest.raises(ValueError, match="declared parent"):
        NativeWorkflowStep(id="missing-parent", task_id=plan.steps[0].task_id,
            input_bindings=(NativeWorkflowInput(input_artifact_id="budget-docx", parent_step_id="fake"),))
    with pytest.raises(ValueError, match="unknown benchmark task"):
        validate_workflow(benchmark, NativeWorkflowPlan(id="invented", benchmark_digest=benchmark.digest,
            steps=(NativeWorkflowStep(id="fake", task_id="free-oracle-task"),)))
    with pytest.raises(ValueError, match="identity changed"):
        validate_workflow(benchmark, plan.model_copy(update={"benchmark_digest": "0" * 64}))
    with pytest.raises(ValueError, match=r"extra_forbidden|Extra inputs"):
        NativeWorkflowPlan.model_validate({**plan.model_dump(mode="json"), "expected": "invented oracle"})


def test_same_format_wrong_source_and_write_rebinding_are_refused(
    benchmark: NativeBenchmark, plan: NativeWorkflowPlan,
) -> None:
    update_doc = plan.steps[2]
    actual_read = next(task for task in benchmark.workload.tasks if task.operation == "read"
        and len(task.inputs) == 1 and task.inputs[0].artifact_id == "actual-docx")
    wrong_source = NativeWorkflowPlan(id="wrong-source", benchmark_digest=benchmark.digest, steps=(
        update_doc.model_copy(update={"depends_on": ()}),
        NativeWorkflowStep(id="wrong", task_id=actual_read.id, depends_on=("update-doc",),
            input_bindings=(NativeWorkflowInput(input_artifact_id="actual-docx", parent_step_id="update-doc"),))))
    with pytest.raises(ValueError, match="same source artifact"):
        validate_workflow(benchmark, wrong_source)
    stale_oracle = NativeWorkflowPlan(id="stale-oracle", benchmark_digest=benchmark.digest, steps=(
        update_doc.model_copy(update={"depends_on": ()}),
        update_doc.model_copy(update={"id": "again", "depends_on": ("update-doc",), "input_bindings": (
            NativeWorkflowInput(input_artifact_id="budget-docx", parent_step_id="update-doc"),)})))
    with pytest.raises(ValueError, match="derived output contract"):
        validate_workflow(benchmark, stale_oracle)


def test_resume_pins_harness_plan_and_policy_before_calls(
    benchmark: NativeBenchmark, plan: NativeWorkflowPlan, tmp_path: Path,
) -> None:
    from worldloom.evalrun.agents import ScriptedAgent
    from worldloom.evalrun.policy import AgentPolicy

    class PolicyAgent(ScriptedAgent):
        policy: Any = AgentPolicy(system="Read the supplied files.", skills={"inspect": "Read actual input bytes."})

    calls: list[str] = []
    def target(request: dict[str, Any], inputs: Mapping[str, bytes]) -> NativeSubmission:
        calls.append(request["execution_id"])
        assert request["agent"]["policy"]["skills"]["inspect"] == "Read actual input bytes."
        return _target(request, inputs)
    agent = PolicyAgent(())
    harness = CallableHarness(target, {"worker": "policy/v1"})
    output = tmp_path / "run"
    run_workflow(benchmark, plan, harness, directory=output, agent=agent)
    with pytest.raises(ValueError, match="configuration changed"):
        run_workflow(benchmark, plan.model_copy(update={"id": "different-plan"}), harness, directory=output, resume=True, agent=agent)
    with pytest.raises(ValueError, match="configuration changed"):
        run_workflow(benchmark, plan, CallableHarness(target, {"worker": "policy/v2"}), directory=output, resume=True, agent=agent)
    agent.policy = AgentPolicy(system="Read the supplied files.", skills={"inspect": "Skip verification."})
    with pytest.raises(ValueError, match="configuration changed"):
        run_workflow(benchmark, plan, harness, directory=output, resume=True, agent=agent)
    assert len(calls) == 7


def test_parity_metadata_names_unobserved_legacy_capabilities() -> None:
    capabilities = workflow_capabilities()
    assert "conditional_branches" in capabilities["unavailable"]
    assert "iteration" in capabilities["unavailable"]
    assert "autonomous_agent_plan_scoring" in capabilities["unavailable"]
    assert "graded_update_bytes_as_inputs" in capabilities["supported"]
