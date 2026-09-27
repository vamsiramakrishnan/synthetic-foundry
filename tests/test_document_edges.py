"""Four rough edges read off a live-narrated reader-grade corpus, each fixed at its root.

1. **The direction twice.** "missing revenue plan by AUD 10.2m adverse",
   "missing by only AUD 1.1m adverse": a clause that already says which way a
   figure went leaves the figure bare, read from a direction lexicon in pack
   text rather than a list of phrases; a clause that says the other way is
   refused.
2. **Titles.** A slide title is complete, within the profile's words and
   characters, and carries no dash: a long claim is shortened to a shorter
   clause, never truncated.
3. **A table's unit.** Every money table states the ledger unit its cells are
   held in, in DOCX, PDF, Markdown and HTML; an audit rendering is unchanged.
4. **The floor.** A section's moves state how many sentences each must say,
   the brief carries it, and a section below it is refused as
   ``section_floor`` unless it has fewer facts than moves.
"""

from __future__ import annotations

import html
import re
from datetime import UTC, datetime
from io import BytesIO

import pytest

from worldloom import figures, presentation, realism_profiles, recipe, rhetoric, titles
from worldloom.artifact_text import extract
from worldloom.locales import DEFAULT
from worldloom.models import (
    Authority,
    CanonicalFact,
    Cell,
    Column,
    Quantity,
    Row,
    Table,
)
from worldloom.narrative import ComposedProvider, handshake, prompts, references
from worldloom.narrative.claims import validate
from worldloom.narrative.requests import (
    GeneratedClaim,
    GeneratedNarrative,
    NarrativeRequest,
    RequestMove,
)
from worldloom.presentation import AUDIT, READER
from worldloom.render.enterprise import presenter
from worldloom.retail import RetailWorld
from worldloom.scenarios import MonthEndClose

_WHEN = datetime(2026, 3, 31, tzinfo=UTC)


def _fact(fact_id: str, amount: float | None = None, unit: str = "AUD_thousands", *,
          kind: str = "financial.revenue.variance", text: str | None = None) -> CanonicalFact:
    return CanonicalFact(id=fact_id, kind=kind, subject="COMP-0001", period="2026-03",
                         value=Quantity(amount=amount, unit=unit) if amount is not None else None,
                         text_value=text, valid_from=_WHEN, authority=Authority.SYSTEM_OF_RECORD)


FACTS = {
    "FACT-0001": _fact("FACT-0001", -10_200.0),
    "FACT-0002": _fact("FACT-0002", -1_100.0),
    "FACT-0003": _fact("FACT-0003", 393_300.0, kind="financial.revenue.actual"),
    "FACT-0004": _fact("FACT-0004", -2_862.0, kind="financial.gross_profit.variance"),
}


def _spell(text: str, profile: presentation.Presentation = READER) -> str:
    return references.substitute(text, FACTS, locale=DEFAULT, presentation=profile)


# -- 1. the direction once ----------------------------------------------------------------


@pytest.mark.parametrize(("text", "spelled"), [
    # The live deck's two sentences.
    ("Food carried the group's shortfall, missing revenue plan by {{fact:FACT-0001}}.",
     "Food carried the group's shortfall, missing revenue plan by AUD 10.2m."),
    ("General Merchandise held close to plan, missing by only {{fact:FACT-0002}} on revenue.",
     "General Merchandise held close to plan, missing by only AUD 1.1m on revenue."),
    # A live writer's longer clause: the verb is eight words back.
    ("Food missed budget on revenue by a variance of {{fact:FACT-0001}}.",
     "Food missed budget on revenue by a variance of AUD 10.2m."),
    # Verbs and nouns the phrase list never named, before and after the figure.
    ("Revenue fell {{fact:FACT-0001}} for the month.", "Revenue fell AUD 10.2m for the month."),
    ("Food overspent by {{fact:FACT-0001}}.", "Food overspent by AUD 10.2m."),
    ("Food came in {{fact:FACT-0001}} below budget.", "Food came in AUD 10.2m below budget."),
    ("Digital beat plan while Food missed by {{fact:FACT-0001}}.",
     "Digital beat plan while Food missed by AUD 10.2m."),
    # Words that say nothing about direction leave the figure to say it.
    ("The widest gap is Food, at {{fact:FACT-0001}}.", "The widest gap is Food, at AUD 10.2m adverse."),
    ("Gross profit was off by {{fact:FACT-0001}}.", "Gross profit was off by AUD 10.2m adverse."),
    # A direction word in an earlier clause is that clause's, not this figure's.
    ("Revenue fell short, and gross profit finished {{fact:FACT-0001}}.",
     "Revenue fell short, and gross profit finished AUD 10.2m adverse."),
])
def test_a_clause_that_says_the_direction_leaves_the_figure_bare(text: str, spelled: str) -> None:
    out = _spell(text)
    assert out == spelled
    assert figures.defects(out) == []


def test_another_figures_direction_word_is_not_this_ones() -> None:
    out = _spell("Food finished {{fact:FACT-0001}} and General Merchandise {{fact:FACT-0002}}.")
    assert out == "Food finished AUD 10.2m adverse and General Merchandise AUD 1.1m adverse."
    assert figures.defects(out) == []


def test_the_lexicon_is_pack_text() -> None:
    from worldloom import packkit

    adverse = packkit.template("render.figures.direction_words.adverse")
    favourable = packkit.template("render.figures.direction_words.favourable")
    for word in ("miss", "missing", "shortfall", "short of", "fell", "below", "over budget", "overspend"):
        assert word in adverse.split(" | "), word
    for word in ("ahead of plan", "beat"):
        assert word in favourable.split(" | "), word
    assert figures.direction("missing revenue plan by ") == "adverse"
    assert figures.direction("ahead of plan by ") == "favourable"
    assert figures.direction("revenue of ") is None


def test_the_exact_spelling_is_untouched() -> None:
    assert _spell("Missing plan by {{fact:FACT-0001}}.", AUDIT) == "Missing plan by AUD 10,200 thousands adverse."


def test_a_typed_direction_or_a_contradiction_is_refused_as_number_spelling() -> None:
    request = NarrativeRequest(artifact_id="ART-0001", artifact_type="cfo_variance_memo", section="Position",
                               persona_id="", voice="plain", audience="board", author_title="CFO",
                               allowed_fact_ids=list(FACTS))

    def verdict(text: str):  # type: ignore[no-untyped-def]
        narrative = GeneratedNarrative(text=text, claims=[GeneratedClaim(text=text, supporting_fact_ids=["FACT-0001"])])
        return validate(request, narrative, FACTS, presentation=READER, locale=DEFAULT)

    assert verdict("Food carried the miss, missing plan by {{fact:FACT-0001}}.").accepted
    typed = verdict("Food carried the miss, missing plan by {{fact:FACT-0001}} adverse.")
    assert not typed.accepted and "direction said twice" in typed.feedback
    backwards = verdict("Food finished ahead of plan by {{fact:FACT-0001}}.")
    assert not backwards.accepted and "direction contradicts" in backwards.feedback


# -- 2. titles ------------------------------------------------------------------------------


def test_the_reader_profile_names_a_title_style() -> None:
    assert READER.titles == "reader" and AUDIT.titles == "free"
    assert "titles" in presentation.KNOBS and presentation.TITLES == titles.styles()
    assert titles.rules_for(AUDIT) is None
    style = titles.rules_for(READER)
    assert style is not None and style.max_words == 12 and style.max_chars == 80
    # A recipe that names the audit profile keeps its bytes: the knob is
    # written only when it differs from its default.
    assert "titles" not in recipe.with_presentation({}, "audit")["presentation"]


@pytest.mark.parametrize(("claim", "fitted"), [
    # A dash between clauses: the first clause is the point.
    ("Revenue missed plan by AUD 13.3m for the period — a miss, not a rounding difference",
     "Revenue missed plan by AUD 13.3m for the period"),
    # A trailing phrase cut away, keeping the figure.
    ("Revenue came in at AUD 617.2m against a budget of AUD 630.5m, the weakest month of the year so far",
     "Revenue came in at AUD 617.2m against a budget of AUD 630.5m"),
    # A clause after ", which" has no subject of its own and never titles a
    # slide: the live deck read "Alone carried a margin impact of 114 bps".
    ("Group gross margin moved 120 bps adverse, and most of that is traceable to promotional depth in Food,"
     " which alone carried a margin impact of 114 bps adverse", "Group gross margin moved 120 bps adverse"),
    # Short enough: whole, with a dash typeset as a colon.
    ("Revenue missed plan — AUD 13.3m", "Revenue missed plan: AUD 13.3m"),
])
def test_a_long_claim_is_fitted_to_a_complete_shorter_clause(claim: str, fitted: str) -> None:
    style = titles.rules_for(READER)
    assert style is not None
    out = titles.fit(claim, style, lambda piece: "AUD" in piece or "bps" in piece)
    assert out == fitted
    assert titles.lint(out, style) == []


def test_nothing_complete_fits_so_the_lead_fact_titles_the_slide() -> None:
    style = titles.rules_for(READER)
    assert style is not None
    long = "The confirmed cause of the failure was a stale legacy to new product hierarchy mapping in the master"
    assert titles.fit(long, style, lambda piece: "stale legacy" in piece) is None
    raw = "The confirmed cause of the failure was {{fact:FACT-0009}}."
    facts = {"FACT-0009": _fact("FACT-0009", kind="ops.cause",
                                text="Stale legacy-to-new product hierarchy mapping in the merchandising master")}
    spelled = raw.replace("{{fact:FACT-0009}}", facts["FACT-0009"].text_value or "")
    title = presenter.takeaway(raw, spelled, "Cause", facts,
                               lambda fid: "Stale legacy-to-new product hierarchy mapping in the merchandising master",
                               set(), style, lambda piece: "stale legacy" in piece.lower())
    assert title == "Stale legacy-to-new product hierarchy mapping in the merchandising master"
    assert "…" not in title


def test_the_title_lint_names_length_dashes_and_open_clauses() -> None:
    style = titles.rules_for(READER)
    findings = presenter.lint_titles([
        ("Revenue missed plan by AUD 13.3m for the period — a miss, not a rounding difference",
         ["AUD 13.3m"]),
        ("The confirmed cause of the failure was Stale legacy-to-new mapping in the…", ["Stale legacy-to-new"]),
        ("Revenue finished AUD 13.3m adverse against plan", ["AUD 13.3m adverse"]),
    ], style)
    assert any("over the reader style's 12 words" in f for f in findings)
    assert any("does not print in a title" in f for f in findings)
    # Unchanged without a style: the shipped lint.
    assert presenter.lint_titles([("Revenue finished AUD 13.3m adverse against plan", ["AUD 13.3m adverse"])]) == []


# -- a rendered reader-grade corpus ------------------------------------------------------------


@pytest.fixture(scope="module")
def planned():  # type: ignore[no-untyped-def]
    world = RetailWorld(seed=8128).build().run(MonthEndClose(period="2026-03", include_operational_incident=True))
    return world.extend(recipe=realism_profiles.with_realism(world.recipe, "enterprise")).compile()


@pytest.fixture(scope="module")
def narrated(planned):  # type: ignore[no-untyped-def]
    return planned.narrate(ComposedProvider.for_world(planned))


@pytest.fixture(scope="module")
def rendered(narrated):  # type: ignore[no-untyped-def]
    return narrated.render("docx", "pdf", "pptx", "markdown", "html")


def _file(world, artifact_type: str, ext: str) -> tuple[str, bytes]:  # type: ignore[no-untyped-def]
    for ir in world.artifact_irs:
        if world.artifact_intents.by_id(ir.intent_id).artifact_type != artifact_type:
            continue
        for item in world._rendered:
            if item.artifact_id == ir.id and item.path.endswith(f".{ext}") and "/revisions/" not in item.path:
                return item.path, bytes(item.payload)
    raise AssertionError(f"no {ext} for {artifact_type}")


def test_every_deck_title_is_complete_bounded_and_dashless(rendered) -> None:  # type: ignore[no-untyped-def]
    from pptx import Presentation as Deck

    _, payload = _file(rendered, "executive_summary", "pptx")
    style = titles.rules_for(READER)
    assert style is not None
    for slide in Deck(BytesIO(payload)).slides:
        title = slide.shapes.title.text if slide.shapes.title is not None else ""
        assert titles.lint(title, style) == [], title


def test_the_prose_carries_no_direction_said_twice(rendered) -> None:  # type: ignore[no-untyped-def]
    _, payload = _file(rendered, "cfo_variance_memo", "md")
    text = payload.decode()
    body = text.split("## Appendix", 1)[0]
    assert not [d for d in figures.defects(body) if "direction" in d]


@pytest.mark.parametrize("ext", ["docx", "pdf", "md", "html"])
def test_every_money_schedule_states_its_unit(rendered, ext: str) -> None:  # type: ignore[no-untyped-def]
    path, payload = _file(rendered, "cfo_variance_memo", ext)
    if ext in ("md", "html"):
        text = html.unescape(payload.decode())
    else:
        read = extract(path, payload)
        assert read is not None
        # PDF text extraction breaks "P&L" around its ampersand.
        text = re.sub(r"\s*&\s*", "&", read.text)
    for schedule in ("Business Unit P&L", "Category P&L", "Corporate Cost Base", "Store Performance"):
        pattern = r"\s+".join(re.escape(word) for word in f"{schedule} (AUD thousands)".split())
        assert re.search(pattern, text), (ext, schedule)
    # A schedule with a unit column of its own states it per row, not twice.
    assert not re.search(r"Variance Drivers\s*\(", text)


def test_the_audit_rendering_prints_no_unit_caption(narrated) -> None:  # type: ignore[no-untyped-def]
    audited = narrated.extend(recipe=recipe.with_presentation(narrated.recipe, "audit")).render("markdown")
    _, payload = _file(audited, "cfo_variance_memo", "md")
    assert "(AUD thousands)" not in payload.decode()


def test_a_table_in_two_currencies_states_each_in_its_header() -> None:
    facts = {"FACT-0001": _fact("FACT-0001", 5.0), "FACT-0002": _fact("FACT-0002", 7.0, "NZD_millions")}
    table = Table(key="t", title="Two books", columns=[Column(key="a", label="Australia"), Column(key="b", label="New Zealand")],
                  rows=[Row(key="r", label="Revenue", cells={"a": Cell(value=5.0, fact_id="FACT-0001"),
                                                             "b": Cell(value=7.0, fact_id="FACT-0002")})])
    caption, out = figures.unit_caption("Two books", table, facts, READER)
    assert caption == "Two books"
    assert [c.label for c in out.columns] == ["Australia (AUD thousands)", "New Zealand (NZD millions)"]
    assert figures.unit_caption("Two books", table, facts, AUDIT) == ("Two books", table)


# -- 4. the floor ------------------------------------------------------------------------------


def test_the_floor_is_stated_per_move_and_a_thin_section_is_exempt(planned) -> None:  # type: ignore[no-untyped-def]
    narrated = planned
    requests = handshake.pending(narrated)
    floored = [r for r in requests if r.floor is not None and not r.floor.exempt]
    exempt = [r for r in requests if r.floor is not None and r.floor.exempt]
    assert floored and exempt
    for request in floored:
        assert all(move.sentences >= 1 for move in request.moves)
        assert request.floor is not None and request.floor.sentences == sum(m.sentences for m in request.moves)
    for request in exempt:
        assert len(request.allowed_fact_ids) < len(request.moves)
        assert request.floor is not None and "fact(s) for" in request.floor.exempt
    document = handshake.requests_document(narrated)
    payload = next(r for r in document["requests"] if "floor" in r and "sentences" in r["floor"])
    assert all("sentences" in move for move in payload["moves"])
    assert any("section_floor" in rule for rule in document["rules"])
    brief = prompts.for_world(narrated).render(floored[0], {f.id: f for f in narrated.facts})
    assert "At least" in brief and "shorter is refused as section_floor" in brief


def test_a_section_below_its_floor_is_refused_with_the_fix() -> None:
    moves = [RequestMove(name="headline", instruction="", fact_ids=["FACT-0001"], sentences=1),
             RequestMove(name="comparison", instruction="", fact_ids=["FACT-0003", "FACT-0004"], sentences=2),
             RequestMove(name="implication", instruction="", fact_ids=["FACT-0001"], derived=True, sentences=1)]
    floor = rhetoric.floor(moves, 3)
    assert (floor.sentences, floor.paragraphs, floor.exempt) == (4, 1, "")
    request = NarrativeRequest(artifact_id="ART-0003", artifact_type="cfo_variance_memo", section="Position",
                               persona_id="", voice="plain", audience="board", author_title="CFO",
                               allowed_fact_ids=["FACT-0001", "FACT-0003", "FACT-0004"], moves=moves, floor=floor)
    thin = "Revenue finished {{fact:FACT-0001}}.\n\nRevenue came in at {{fact:FACT-0003}}, gross profit {{fact:FACT-0004}}."
    claims = [GeneratedClaim(text=thin, supporting_fact_ids=["FACT-0001", "FACT-0003", "FACT-0004"])]
    verdict = validate(request, GeneratedNarrative(text=thin, claims=claims), FACTS, presentation=READER, locale=DEFAULT)
    assert not verdict.accepted and "section_floor" in verdict.feedback and "headline 1" in verdict.feedback
    full = thin + " Gross profit fell further than revenue did.\n\nThe month does not meet plan."
    verdict = validate(request, GeneratedNarrative(text=full, claims=claims), FACTS, presentation=READER, locale=DEFAULT)
    assert verdict.accepted, verdict.feedback
    # Fewer facts than moves: exempt, and the reason is recorded.
    exempt = rhetoric.floor(moves, 2)
    assert exempt.sentences == 0 and "2 fact(s) for 3 move(s)" in exempt.exempt
    request = request.model_copy(update={"floor": exempt})
    assert validate(request, GeneratedNarrative(text=thin, claims=claims), FACTS, presentation=READER,
                    locale=DEFAULT).accepted


def test_a_measure_is_one_sentence_however_many_facets_it_has() -> None:
    facts = [_fact("FACT-0003", 393_300.0, kind="financial.revenue.actual"),
             _fact("FACT-0005", 403_500.0, kind="financial.revenue.budget"),
             _fact("FACT-0001", -10_200.0, kind="financial.revenue.variance")]
    assert rhetoric._move_sentences("comparison", False, [f.id for f in facts], {f.id: f for f in facts}) == 1
