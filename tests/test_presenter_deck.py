"""The presenter deck argues from facts and talks like a person.

Read off a rendered enterprise/v2 deck before this: content slides titled
"What the committee needs to note" and "The group's position in one line",
an agenda that listed those titles, and notes that said "That is the message
of this slide" and "Which brings us to the next point" on slide after slide.
"""

from __future__ import annotations

from io import BytesIO

import pytest

from worldloom import figures, packkit, prose_quality, realism_profiles
from worldloom.narrative import ComposedProvider, references
from worldloom.presentation import READER
from worldloom.render.enterprise import presenter
from worldloom.retail import RetailWorld
from worldloom.scenarios import MonthEndClose

BOILERPLATE = ("That is the message of this slide", "Which brings us to the next point",
               "everything else on the slide supports that")


@pytest.fixture(scope="module")
def rendered():  # type: ignore[no-untyped-def]
    world = RetailWorld(seed=8128).build().run(MonthEndClose(period="2026-03", include_operational_incident=True))
    planned = world.extend(recipe=realism_profiles.with_realism(world.recipe, "enterprise")).compile()
    return planned.narrate(ComposedProvider.for_world(planned)).render("pptx")


@pytest.fixture(scope="module")
def deck(rendered):  # type: ignore[no-untyped-def]
    from pptx import Presentation as Deck

    for ir in rendered.artifact_irs:
        if rendered.artifact_intents.by_id(ir.intent_id).artifact_type != "executive_summary":
            continue
        for item in rendered._rendered:
            if item.artifact_id == ir.id and item.path.endswith(".pptx") and "/revisions/" not in item.path:
                return Deck(BytesIO(bytes(item.payload)))
    raise AssertionError("no executive summary deck")


def _content(deck):  # type: ignore[no-untyped-def]
    """Every slide that makes a point: after the agenda, before the close."""
    slides = list(deck.slides)
    closing = next(i for i, s in enumerate(slides) if s.shapes.title is not None
                   and s.shapes.title.text == packkit.text("render.deck.closing.title"))
    return slides[2:closing]


def _grounded(world) -> list[str]:  # type: ignore[no-untyped-def]
    from worldloom.render.values import corpus_locale

    locale = corpus_locale(world)
    values: list[str] = []
    for fact in world.facts:
        spelled = references.render_value(fact, locale=locale, presentation=READER)
        values.append(spelled if fact.value is not None else " ".join(spelled.split()[:4]))
        values.extend(figures.spellings_of(fact, locale=locale, presentation=READER))
    return [v for v in values if v] + list(world.entity_names().values())


def test_every_content_slide_is_titled_with_a_fact(rendered, deck) -> None:  # type: ignore[no-untyped-def]
    content = _content(deck)
    assert len(content) >= 6
    grounded = _grounded(rendered)
    findings = presenter.lint_titles([(s.shapes.title.text, grounded) for s in content])
    assert findings == []


def test_the_lint_refuses_a_lead_in_and_a_title_with_no_fact() -> None:
    findings = presenter.lint_titles([("What the committee needs to note", ["AUD 13.3m adverse"]),
                                      ("The picture this month", ["AUD 13.3m adverse"]),
                                      ("Revenue finished AUD 13.3m adverse against plan", ["AUD 13.3m adverse"])])
    assert len(findings) == 2 and "lead-in" in findings[0] and "none of the slide's facts" in findings[1]


def test_the_agenda_lists_the_argument_not_the_titles(deck) -> None:  # type: ignore[no-untyped-def]
    agenda = list(deck.slides)[1]
    lines = [p.text for p in agenda.placeholders[1].text_frame.paragraphs if p.text]
    titles = {s.shapes.title.text for s in _content(deck)}
    labels = set(packkit.texts("render.deck.agenda.move."))
    assert lines and not set(lines) & titles
    assert set(lines) <= labels


def test_notes_vary_and_say_why_the_next_slide_follows(deck) -> None:  # type: ignore[no-untyped-def]
    notes = [s.notes_slide.notes_text_frame.text for s in deck.slides]
    assert not any(phrase in note for note in notes for phrase in BOILERPLATE)
    assert prose_quality.notes_repetition(notes) <= prose_quality.NOTES_THRESHOLD
    relations = [piece.split("{")[0].strip() for text in packkit.texts("render.deck.notes.relation.")
                 for piece in text.split(" | ") if piece.split("{")[0].strip()]
    bridged = [note for note in notes if any(start in note for start in relations)]
    assert len(bridged) >= len(_content(deck)) // 2


def test_the_notes_measure_reads_a_template_as_repetition() -> None:
    template = [f"Point {i}. That is the message of this slide. Which brings us to the next point: {i}."
                for i in range(8)]
    assert prose_quality.notes_repetition(template) > prose_quality.NOTES_THRESHOLD
    assert prose_quality.notes_repetition(["One claim here.", "A different claim there."]) == 0.0
