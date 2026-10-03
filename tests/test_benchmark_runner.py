"""Execute an actual workbook reader, without giving the target an oracle."""
from __future__ import annotations

import json
import os
import sys
import venv
from dataclasses import replace
from datetime import UTC, datetime
from io import BytesIO
from pathlib import Path
from typing import Any

import pytest

from worldloom.benchmarks import (
    CallableHarness,
    CommandHarness,
    HarnessFailure,
    NativeBenchmark,
    NativeWorkloadPlan,
    protocol_manifest,
    run_benchmark,
)
from worldloom.benchmarks.runner import REQUEST_SCHEMA, RESPONSE_SCHEMA
from worldloom.corpus import write_json
from worldloom.evalrun.agents import ScriptedAgent
from worldloom.evalrun.policy import AgentPolicy
from worldloom.models import (
    ArtifactIR,
    ArtifactSection,
    Authority,
    CanonicalFact,
    Company,
)
from worldloom.native_corpus import (
    NativeContent,
    NativeCorpusPlan,
    render_native_corpus,
)
from worldloom.native_tasks import (
    NativeAnswer,
    NativeCitation,
    NativeSubmission,
    public_contract,
)
from worldloom.providers import digest
from worldloom.world import World


@pytest.fixture
def benchmark() -> NativeBenchmark:
    pytest.importorskip("openpyxl")
    facts = tuple(CanonicalFact(id=f"FACT-{name.upper()}", kind="financial.revenue.actual", subject=f"store:{name}",
        period="2026-01", text_value=str(amount), valid_from=datetime(2026, 1, 1, tzinfo=UTC),
        authority=Authority.SYSTEM_OF_RECORD) for name, amount in (("east", 125), ("west", 70)))
    world = World(seed=8128, company=Company(id="CO-1", name="Northstar Retail", industry="retail",
        headquarters="Sydney", fiscal_year_start_month=7, employees_total=100), _facts=facts,
        _artifact_irs=(ArtifactIR(id="ART-SOURCE", intent_id="INTENT-1", title="January close", sections=[
            ArtifactSection(heading=f"{fact.subject} revenue", body="{{fact:" + fact.id + "}}",
                fact_ids=[fact.id]) for fact in facts]),))
    rendered = {name: render_native_corpus(world, NativeCorpusPlan(artifact_id=name, format="xlsx",
        title=f"{name} store revenues", contents=(NativeContent(source_artifact_id="ART-SOURCE", section_index=index),)))
        for index, name in enumerate(("east", "west"))}
    return NativeBenchmark.from_rendered(world, rendered, NativeWorkloadPlan(use_case_id="close",
        objective="Read store revenue evidence.", formats=("xlsx",), operations=("read",), discovery_scope="artifact"))


_READER = '''import json, os, sys
from pathlib import Path
from openpyxl import load_workbook
request = json.load(sys.stdin)
assert request["schema"] == "worldloom.native-harness-request/v1"
assert set(request) <= {"schema", "task", "input_root", "agent"}
task = request["task"]
assert set(task) == {"id", "operation", "prompt", "inputs", "answers", "output", "submission_schema", "execution_id"}
assert "expected" not in json.dumps(task)
root = Path(request["input_root"])
assert root == Path.cwd()
assert len(list(root.rglob("*.xlsx"))) == len(task["inputs"]) == 1
assert not list(root.rglob("oracle.json"))
source = task["inputs"][0]
with (root / source["path"]).open("rb") as handle:
    workbook = load_workbook(handle, data_only=False)
    value = str(workbook["Evidence"]["B2"].value)
workbook.close()
if os.environ.get("CALL_LOG"):
    with open(os.environ["CALL_LOG"], "a") as output:
        output.write(task["execution_id"] + "\\n")
if os.environ.get("MODE") == "bad-west" and source["artifact_id"] == "west":
    print("not JSON")
    raise SystemExit()
if os.environ.get("MODE") == "wrong":
    value = "999999"
if os.environ.get("MODE") == "policy":
    assert request["agent"]["policy"]["skills"]["verify"] == "Read workbook bytes."
submission = {"answers": [{"assertion_id": task["answers"][0]["assertion_id"], "value": value,
    "citations": [{"artifact_id": source["artifact_id"], "locator": "sheet:Evidence/cell:B2"}]}], "files": []}
print(json.dumps({"schema": "worldloom.native-harness-response/v1", "task_id": task["id"],
    "execution_id": task["execution_id"], "submission": submission}))
'''


def reader(tmp_path: Path, **environment: str) -> CommandHarness:
    script = tmp_path / "workbook reader.py"
    script.write_text(_READER, encoding="utf-8")
    return CommandHarness((sys.executable, str(script)), environment=environment)


@pytest.mark.skipif(os.name == "nt", reason="POSIX virtual environments use interpreter symlinks")
def test_command_keeps_the_selected_virtual_environment_after_staging(tmp_path: Path) -> None:
    environment = tmp_path / "isolated-python"
    venv.EnvBuilder(with_pip=False, symlinks=True).create(environment)
    executable = environment / "bin" / "python"
    assert executable.is_symlink()
    script = tmp_path / "prefix-worker.py"
    script.write_text('''import json, sys
assert sys.prefix == sys.argv[1], "staged process lost its selected virtual environment"
request = json.load(sys.stdin)
task = request["task"]
print(json.dumps({"schema": "worldloom.native-harness-response/v1", "task_id": task["id"],
    "execution_id": task["execution_id"], "submission": {"answers": [], "files": []}}))
''', encoding="utf-8")
    harness = CommandHarness((str(executable), str(script), str(environment)))
    reply = harness.submit({"id": "venv-check", "execution_id": "venv-execution", "inputs": []}, {})
    assert reply == NativeSubmission()
    assert harness.identity["argv"][0] == str(executable)


def test_real_reader_gets_only_public_task_bytes_and_exact_resume_makes_no_calls(
    benchmark: NativeBenchmark, tmp_path: Path,
) -> None:
    log = tmp_path / "calls.txt"
    harness = reader(tmp_path, CALL_LOG=str(log))
    out = tmp_path / "run"
    report = run_benchmark(benchmark, harness, directory=out, repeats=2)
    assert report.passed and report.passed_count == report.total == 4
    calls = log.read_text().splitlines()
    assert len(set(calls)) == 4
    prior = {path.relative_to(out): path.read_bytes() for path in out.rglob("*") if path.is_file()}
    assert run_benchmark(benchmark, harness, directory=out, repeats=2, resume=True) == report
    assert log.read_text().splitlines() == calls
    assert {path.relative_to(out): path.read_bytes() for path in out.rglob("*") if path.is_file()} == prior
    assert all(trial.grade.metrics["assertions"] == 1 for trial in report.trials)


def test_malformed_reply_fails_one_task_and_independent_grade_rejects_wrong_answers(
    benchmark: NativeBenchmark, tmp_path: Path,
) -> None:
    report = run_benchmark(benchmark, reader(tmp_path, MODE="bad-west"), directory=tmp_path / "mixed")
    assert report.total == 2 and report.passed_count == 1 and report.harness_failures == 1
    west = next(task.id for task in benchmark.workload.tasks if task.inputs[0].artifact_id == "west")
    assert next(trial for trial in report.trials if trial.task_id == west).failure_code == "response_json_invalid"
    wrong = run_benchmark(benchmark, reader(tmp_path, MODE="wrong"), directory=tmp_path / "wrong")
    assert wrong.passed_count == wrong.harness_failures == 0
    assert all(not trial.grade.passed for trial in wrong.trials)


def _read_public(request: dict[str, Any], inputs: Any) -> NativeSubmission:
    from openpyxl import load_workbook

    source = request["inputs"][0]
    workbook = load_workbook(BytesIO(inputs[source["artifact_id"]]))
    value = str(workbook["Evidence"]["B2"].value)
    workbook.close()
    return NativeSubmission(answers=(NativeAnswer(assertion_id=request["answers"][0]["assertion_id"], value=value,
        citations=(NativeCitation(artifact_id=source["artifact_id"], locator="sheet:Evidence/cell:B2"),)),))


def test_interrupt_preserves_completed_receipts_and_only_reissues_uncommitted_trial(
    benchmark: NativeBenchmark, tmp_path: Path,
) -> None:
    calls: list[str] = []
    def target(request: dict[str, Any], inputs: Any) -> NativeSubmission:
        calls.append(request["execution_id"])
        if len(calls) == 2:
            raise KeyboardInterrupt()
        return _read_public(request, inputs)
    harness = CallableHarness(target, {"implementation": "reader/v1"})
    out = tmp_path / "partial"
    with pytest.raises(KeyboardInterrupt):
        run_benchmark(benchmark, harness, directory=out)
    assert len(list((out / "receipts").glob("*.json"))) == 1
    report = run_benchmark(benchmark, harness, directory=out, resume=True)
    assert report.passed and len(calls) == 3
    assert calls[1] == calls[2] != calls[0]


def test_resume_refuses_command_configuration_task_and_receipt_tampering_before_calls(
    benchmark: NativeBenchmark, tmp_path: Path,
) -> None:
    calls: list[str] = []
    def target(request: dict[str, Any], inputs: Any) -> NativeSubmission:
        calls.append(request["id"])
        return _read_public(request, inputs)
    harness = CallableHarness(target, {"implementation": "reader/v1"})
    out = tmp_path / "run"
    run_benchmark(benchmark, harness, directory=out)
    with pytest.raises(ValueError, match="configuration changed"):
        run_benchmark(benchmark, harness, directory=out, repeats=2, resume=True)
    changed = NativeBenchmark.from_rendered(benchmark.world, benchmark.rendered,
        benchmark.workload.plan.model_copy(update={"objective": "Read a different measure."}))
    with pytest.raises(ValueError, match="configuration changed"):
        run_benchmark(changed, harness, directory=out, resume=True)
    with pytest.raises(ValueError, match="configuration changed"):
        run_benchmark(benchmark, CallableHarness(target, {"implementation": "reader/v2"}), directory=out, resume=True)
    east = next(task.id for task in benchmark.workload.tasks if task.inputs[0].artifact_id == "east")
    west = next(task.id for task in benchmark.workload.tasks if task.inputs[0].artifact_id == "west")
    path = out / "receipts" / (digest([west, 1]) + ".json")
    saved = json.loads(path.read_text())
    saved["submission"]["answers"][0]["value"] = "777"
    # Even a rewritten checksum cannot make a saved success survive the
    # independent byte grader. Also ensure a missing earlier trial is not run.
    saved["digest"] = digest({key: value for key, value in saved.items() if key != "digest"})
    write_json(path, saved)
    (out / "receipts" / (digest([east, 1]) + ".json")).unlink()
    with pytest.raises(ValueError, match="completed benchmark is missing trial receipts"):
        run_benchmark(benchmark, harness, directory=out, resume=True)
    (out / "run.json").unlink()
    with pytest.raises(ValueError, match="independent byte grading"):
        run_benchmark(benchmark, harness, directory=out, resume=True)
    assert calls == sorted(task.id for task in benchmark.workload.tasks)


def test_command_pins_script_bytes_and_explicit_dependency_files(benchmark: NativeBenchmark, tmp_path: Path) -> None:
    script = tmp_path / "target.py"
    dependency = tmp_path / "policy.txt"
    script.write_text(_READER)
    dependency.write_text("policy v1")
    harness = CommandHarness((sys.executable, str(script)), identity_files=(dependency,))
    run_benchmark(benchmark, harness, directory=tmp_path / "run")
    dependency.write_text("policy v2")
    with pytest.raises(ValueError, match="command identity changed"):
        run_benchmark(benchmark, harness, directory=tmp_path / "run", resume=True)
    dependency.write_text("policy v1")
    script.write_text(_READER + "\n# changed")
    with pytest.raises(ValueError, match="command identity changed"):
        _ = harness.identity


@pytest.mark.parametrize(("program", "code"), [
    ("import time; time.sleep(10)", "command_timeout"),
    ("import sys; sys.stdout.write('x' * 20000)", "command_output_limit"),
    ("raise SystemExit(7)", "command_exit_nonzero"),
    ("print('{}')", "response_identity_invalid"),
])
def test_command_time_output_and_response_bounds(benchmark: NativeBenchmark, tmp_path: Path, program: str, code: str) -> None:
    harness = CommandHarness((sys.executable, "-c", program), timeout_seconds=0.25, max_output_bytes=8192)
    task = benchmark.workload.tasks[0]
    with pytest.raises(HarnessFailure) as caught:
        harness.submit({**public_contract(task), "execution_id": "execution"}, {task.inputs[0].artifact_id: benchmark.inputs[task.inputs[0].artifact_id]})
    assert caught.value.code == code


def test_command_is_native_submit_compatible_and_passes_the_actual_policy_body(
    benchmark: NativeBenchmark, tmp_path: Path,
) -> None:
    harness = reader(tmp_path, MODE="policy")
    agent = ScriptedAgent([], name="policy-target")
    agent.policy = AgentPolicy(system="Verify revenue.", skills={"verify": "Read workbook bytes."})
    task = benchmark.workload.tasks[0]
    request = {**public_contract(task), "execution_id": "one"}
    inputs = {task.inputs[0].artifact_id: benchmark.inputs[task.inputs[0].artifact_id]}
    assert harness(agent, request, inputs).answers[0].value == _read_public(request, inputs).answers[0].value
    broken = CommandHarness((sys.executable, "-c", "print('no')"))
    assert broken(agent, request, inputs) == NativeSubmission()
    protocol = protocol_manifest()
    assert protocol["request_schema"] == REQUEST_SCHEMA and protocol["response_schema"] == RESPONSE_SCHEMA
    assert "OS sandbox" in protocol["boundary"]


def test_callable_improvement_receives_policy_and_malformed_replies_fail(
    benchmark: NativeBenchmark,
) -> None:
    policies: list[str] = []
    def callback(request: dict[str, Any], inputs: Any) -> NativeSubmission:
        policies.append(request["agent"]["policy"]["system"])
        assert set(inputs) == {request["inputs"][0]["artifact_id"]}
        return _read_public(request, inputs)
    harness = CallableHarness(callback, {"implementation": "policy-reader/v1"})
    agent = ScriptedAgent([], name="policy-target")
    agent.policy = AgentPolicy(system="Verify revenue.")
    task = benchmark.workload.tasks[0]
    request = {**public_contract(task), "execution_id": "one"}
    source_id = task.inputs[0].artifact_id
    inputs = {source_id: benchmark.inputs[source_id]}
    assert harness(agent, request, inputs).answers[0].value == _read_public(request, inputs).answers[0].value
    assert policies == ["Verify revenue."]
    invalid = CallableHarness(lambda *_: {"claimed_pass": True}, {"implementation": "invalid/v1"})
    assert invalid(agent, request, inputs) == NativeSubmission()


def test_run_refuses_symlinked_receipts_and_never_overwrites_an_existing_run(
    benchmark: NativeBenchmark, tmp_path: Path,
) -> None:
    harness = CallableHarness(_read_public, {"implementation": "reader/v1"})
    out = tmp_path / "run"
    run_benchmark(benchmark, harness, directory=out)
    with pytest.raises(ValueError, match="run exists"):
        run_benchmark(benchmark, harness, directory=out)
    path = next((out / "receipts").glob("*.json"))
    saved = tmp_path / "saved.json"
    path.rename(saved)
    try:
        path.symlink_to(saved)
    except OSError:
        pytest.skip("symlinks unavailable")
    with pytest.raises(ValueError, match="symlink"):
        run_benchmark(benchmark, harness, directory=out, resume=True)


def test_canonical_privacy_and_oracle_reconstruction_run_before_target_or_cache(tmp_path: Path) -> None:
    from test_native_intake_integrity import _changed_section, _corpus

    pytest.importorskip("docx")
    world, rendered, plan = _corpus("docx", "prose")
    benchmark = NativeBenchmark.from_rendered(world, {"review": rendered}, plan)
    calls: list[str] = []
    def submit(request: dict[str, Any], _: Any) -> NativeSubmission:
        calls.append(request["id"])
        return NativeSubmission()
    harness = CallableHarness(submit, {"implementation": "empty/v1"})
    out = tmp_path / "run"
    run_benchmark(benchmark, harness, directory=out)
    count = len(calls)
    private = replace(benchmark, world=_changed_section(world, hidden=True))
    for resume in (False, True):
        with pytest.raises(ValueError, match="private authored section"):
            run_benchmark(private, harness, directory=out if resume else tmp_path / "new", resume=resume)
    revised_task = benchmark.workload.tasks[0].model_copy(update={"prompt": "Answer a different request."})
    revised = replace(benchmark, workload=benchmark.workload.model_copy(update={"tasks": (
        revised_task, *benchmark.workload.tasks[1:])}))
    with pytest.raises(ValueError, match="oracle differs"):
        run_benchmark(revised, harness, directory=tmp_path / "revised")
    assert len(calls) == count
