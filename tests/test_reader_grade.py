"""The reader-grade corpus (``enterprise/v2``): argued prose, provenance off
the page, designed layout, a presenter's deck.

Pins the five gaps the default build had, each fixed as data rather than
code, and the guarantees that make the fixes safe: every figure still agrees
with the fact store (the claim validator accepts every section, and the
world validates), ``legacy`` and ``enterprise/v1`` are unchanged, and a
reader profile keeps every fact id in the file even though no page prints one.
"""
from __future__ import annotations

import re
import zipfile
from io import BytesIO

import pytest

from worldloom import (
    doctypes,
    presentation,
    prose_quality,
    realism_profiles,
    recipe,
    rhetoric,
)
from worldloom.artifact_text import extract
from worldloom.narrative import (
    ComposedProvider,
    DeterministicProvider,
    handshake,
    prompts,
)
from worldloom.narrative.compiler import _request_for
from worldloom.narrative.requests import NarrativeRequest
from worldloom.render.ooxml import custom_properties
from worldloom.retail import RetailWorld
from worldloom.scenarios import MonthEndClose


def _episode():  # type: ignore[no-untyped-def]
    return RetailWorld(seed=8128).build().run(
        MonthEndClose(period="2026-03", include_operational_incident=True))


@pytest.fixture(scope="module")
def planned():  # type: ignore[no-untyped-def]
    world = _episode()
    return world.extend(recipe=realism_profiles.with_realism(world.recipe, "enterprise")).compile()


@pytest.fixture(scope="module")
def narrated(planned):  # type: ignore[no-untyped-def]
    return planned.narrate(ComposedProvider())


@pytest.fixture(scope="module")
def rendered(narrated):  # type: ignore[no-untyped-def]
    return narrated.render("docx", "pdf", "pptx", "markdown")


def _file(world, artifact_type: str, ext: str) -> tuple[str, bytes]:  # type: ignore[no-untyped-def]
    for ir in world.artifact_irs:
        if world.artifact_intents.by_id(ir.intent_id).artifact_type != artifact_type:
            continue
        for item in world._rendered:
            if item.artifact_id == ir.id and item.path.endswith(f".{ext}") and "/revisions/" not in item.path:
                return item.path, bytes(item.payload)
    raise AssertionError(f"no {ext} for {artifact_type}")


# -- the profile seam -----------------------------------------------------------


def test_new_builds_are_reader_grade_and_older_profiles_keep_their_names() -> None:
    assert realism_profiles.DEFAULT_FOR_NEW_BUILDS == realism_profiles.ENTERPRISE_V2
    assert realism_profiles.named("enterprise") == "enterprise/v2"
    assert realism_profiles.named("enterprise/v1") == "enterprise/v1"
    assert realism_profiles.is_enterprise({"artifact_realism": "enterprise/v1"})
    assert not realism_profiles.reader_grade({"artifact_realism": "enterprise/v1"})
    assert realism_profiles.reader_grade({"artifact_realism": "enterprise/v2"})


def test_a_reader_grade_corpus_that_names_no_profile_is_presented_for_a_reader() -> None:
    assert recipe.presentation_of({"artifact_realism": "enterprise/v2"}) == presentation.READER
    assert recipe.presentation_of({"artifact_realism": "enterprise/v1"}) == presentation.AUDIT
    assert recipe.presentation_of({}) == presentation.AUDIT


def test_a_recipe_that_named_a_profile_before_the_new_knobs_keeps_its_bytes() -> None:
    written = recipe.with_presentation({}, "audit")["presentation"]
    assert set(written) == {"name", "appendix", "provenance", "magnitudes", "table_fit"}
    reader = recipe.with_presentation({}, "reader")
    assert reader["presentation"]["citations"] == "appendix"
    assert recipe.presentation_of(reader) == presentation.READER


def test_the_new_knobs_are_linted_like_the_old_ones() -> None:
    seed = presentation.PresentationSeed(name="house-deck", deck="slides", slide_budget="huge")
    findings = presentation.review(seed)
    assert any("deck is 'slides'" in f for f in findings)
    assert any("slide_budget is 'huge'" in f for f in findings)
    assert presentation.READER.slide_cap == presentation.SLIDE_BUDGET_SLIDES["board"]


# -- 1. rhetoric as data ------------------------------------------------------------


def test_every_shipped_doctype_section_declares_moves_the_catalogue_defines() -> None:
    catalogue = rhetoric.catalogue()
    names = set(catalogue["moves"])
    for artifact_type, sections in catalogue["doctypes"].items():
        for heading, moves in sections.items():
            for move in moves:
                assert rhetoric.spec(move).name in names, (artifact_type, heading, move)
    for role, moves in catalogue["roles"].items():
        assert all(rhetoric.spec(m).name in names for m in moves), role
    for doctype in ("cfo_variance_memo", "executive_summary", "working_note", "incident_rca",
                    "knowledge_article", "confluence_page", "meeting_minutes", "email_thread"):
        assert doctype in catalogue["doctypes"], doctype


def test_a_section_is_asked_for_move_by_move_and_every_fact_has_a_move(planned) -> None:  # type: ignore[no-untyped-def]
    facts = {f.id: f for f in planned.facts}
    for ir in planned.artifact_irs:
        for section in ir.sections:
            if not section.awaiting_prose:
                continue
            request = _request_for(planned, ir, section, facts)
            if not request.allowed_fact_ids:
                continue
            assert request.moves, (ir.id, section.heading)
            introduced = {f for m in request.moves if not m.derived for f in m.fact_ids}
            assert introduced == set(request.allowed_fact_ids), (ir.id, section.heading)
            assert all(m.instruction for m in request.moves)


def test_the_memo_moves_are_the_ones_its_catalogue_declares(planned) -> None:  # type: ignore[no-untyped-def]
    facts = {f.id: f for f in planned.facts}
    memo = next(ir for ir in planned.artifact_irs
                if planned.artifact_intents.by_id(ir.intent_id).artifact_type == "cfo_variance_memo")
    by_heading = {s.heading: _request_for(planned, memo, s, facts) for s in memo.sections if s.awaiting_prose}
    assert [m.name for m in by_heading["Position"].moves] == ["headline", "comparison", "implication"]
    headline = by_heading["Position"].moves[0]
    assert [facts[f].kind for f in headline.fact_ids] == ["financial.revenue.variance"]


def test_a_legacy_request_carries_no_moves_and_digests_as_it_always_did() -> None:
    world = _episode().compile()
    facts = {f.id: f for f in world.facts}
    ir = world.artifact_irs[2]
    section = next(s for s in ir.sections if s.awaiting_prose)
    request = _request_for(world, ir, section, facts)
    assert request.moves == [] and request.display == {} and request.recurrence == {} and request.restated == []
    assert request.digest_fields() == {"fact_digest", "moves", "display", "recurrence", "restated"}
    document = handshake.requests_document(world)
    assert document["prompt_version"] == prompts.SECTION_PROSE.key
    assert not any("moves" in r for r in document["requests"])


def test_the_move_brief_is_prompts_pack_text_and_its_version_moves_with_it(planned) -> None:  # type: ignore[no-untyped-def]
    prompt = prompts.for_world(planned)
    assert prompt.key.startswith("section_moves@1+")
    document = handshake.requests_document(planned)
    assert document["prompt_version"] == prompt.key
    assert any("one paragraph per move" in rule for rule in document["rules"])
    request = next(r for r in handshake.pending(planned) if r.moves)
    facts = {f.id: f for f in planned.facts}
    brief = prompt.render(request, facts)
    assert "First paragraph" in brief and request.moves[0].instruction in brief
    assert "{moves}" not in brief and "{{fact:ID}}" in brief


def test_an_authored_doctype_declares_its_own_moves_and_the_lint_reads_them() -> None:
    spec = doctypes.SectionSpec(heading="Network position", kinds=["financial.revenue."], purpose="State it.",
                                moves=["headline", {"move": "comparison", "kinds": ["financial.revenue."]}])
    assert spec.as_plan().moves[0] == "headline"
    assert "moves" not in doctypes.SectionSpec(heading="X", kinds=["a."], purpose="p").model_dump()
    assert rhetoric.lint_moves(["headline", "comparison"], ["financial."]) == []
    findings = rhetoric.lint_moves(["grandstand", {"move": "driver", "kinds": ["esg."]}], ["financial."])
    assert any("'grandstand' is not a move" in f for f in findings)
    assert any("outside the section's own kinds" in f for f in findings)
    broken = doctypes.DocumentType(key="franchise_statement", authority="approved_report", lifecycle="published",
                                   sections=[doctypes.SectionSpec(heading="Position", kinds=["financial.revenue."],
                                                                  purpose="State it.", moves=["grandstand"])])
    assert any("grandstand" in finding for finding in doctypes.lint([broken]))


# -- 2. the offline narrator --------------------------------------------------------------


def test_the_composing_narrator_is_accepted_everywhere_and_validates(narrated) -> None:  # type: ignore[no-untyped-def]
    calls, _replayed, rejected = narrated._narration
    assert calls > 0 and rejected == 0
    assert narrated.validate().ok
    assert {entry.model_id for entry in narrated._ledger if "/" in entry.call_site} >= {ComposedProvider.id}


def test_the_prose_reads_like_a_person_wrote_it(narrated) -> None:  # type: ignore[no-untyped-def]
    reading = prose_quality.measure(narrated)
    assert prose_quality.failures(reading) == [], reading.as_dict()
    bodies = "\n".join(s.body or "" for ir in narrated.artifact_irs for s in ir.sections if not s.hidden)
    assert not re.search(r"(^|\. )For [a-z0-9]+-[a-z0-9-]+,", bodies)
    assert "inventory-valuation," not in re.sub(r"\{\{fact:[^}]+\}\}", "", bodies)


def test_the_contract_fixture_fails_the_same_thresholds() -> None:
    """The reading is not a formality: the fixture's ledger-dump prose misses it."""
    world = _episode().narrate(DeterministicProvider())
    failures = prose_quality.failures(prose_quality.measure(world))
    assert any(f.startswith("template_opener_rate") for f in failures)
    assert any(f.startswith("slug_leaks") for f in failures)


def test_the_composing_narrator_is_deterministic(planned, narrated) -> None:  # type: ignore[no-untyped-def]
    again = planned.narrate(ComposedProvider())
    assert [ir.model_dump() for ir in again.artifact_irs] == [ir.model_dump() for ir in narrated.artifact_irs]


def test_a_section_is_several_paragraphs_one_per_move(narrated) -> None:  # type: ignore[no-untyped-def]
    memo = next(ir for ir in narrated.artifact_irs
                if narrated.artifact_intents.by_id(ir.intent_id).artifact_type == "cfo_variance_memo")
    position = next(s for s in memo.sections if s.heading == "By business unit")
    assert len([p for p in (position.body or "").split("\n\n") if p.strip()]) >= 3


# -- 3. provenance placement ------------------------------------------------------------------


def test_the_reader_body_carries_no_fact_ids_and_the_file_carries_them_all(rendered) -> None:  # type: ignore[no-untyped-def]
    path, payload = _file(rendered, "cfo_variance_memo", "docx")
    read = extract(path, payload)
    assert read is not None
    body = read.text.split("Appendix", 1)[0]
    assert "Figures cited:" not in read.text and "Key figures" not in body and "Figures cited" not in body
    assert not re.search(r"FACT-\d+", read.units[0].text if read.units else "")
    properties = read.extra["properties"]
    assert set(properties["WorldloomFacts"].split()) >= {"FACT-0055"}
    assert "Sources of figures" in read.text
    sections = [u for u in read.units if re.match(r"^\d+ ", u.title)]
    assert sections and not any(re.search(r"FACT-\d+", u.text) for u in sections)


def test_the_audit_profile_keeps_provenance_inline(narrated) -> None:  # type: ignore[no-untyped-def]
    audited = narrated.extend(recipe=recipe.with_presentation(narrated.recipe, "audit")).render("docx")
    path, payload = _file(audited, "cfo_variance_memo", "docx")
    read = extract(path, payload)
    assert read is not None and "Figures cited" in read.text
    assert "properties" not in read.extra


def test_pdf_and_deck_carry_provenance_in_their_properties(rendered) -> None:  # type: ignore[no-untyped-def]
    path, payload = _file(rendered, "cfo_variance_memo", "pdf")
    read = extract(path, payload)
    assert read is not None and "FACT-0055" in read.extra["properties"]["WorldloomFacts"]
    _, deck = _file(rendered, "executive_summary", "pptx")
    assert custom_properties(deck)["WorldloomProfile"] == "reader"


# -- 4. the presenter deck -----------------------------------------------------------------------


def test_the_deck_is_a_presenters(rendered) -> None:  # type: ignore[no-untyped-def]
    from pptx import Presentation as Deck

    _, payload = _file(rendered, "executive_summary", "pptx")
    deck = Deck(BytesIO(payload))
    slides = list(deck.slides)
    assert len(slides) <= presentation.READER.slide_cap
    notes = [s.notes_slide.notes_text_frame.text for s in slides]
    assert all(notes) and not any("FACT-" in n or "ledger entry" in n for n in notes)
    layouts = [s.slide_layout.name for s in slides]
    appendix = layouts.index("Section Header") if "Section Header" in layouts else len(layouts)
    before = slides[:appendix]
    assert any(s.slide_layout.name == "Two Content" and any(sh.has_chart for sh in s.shapes) for s in before)
    assert not any(any(sh.has_table for sh in s.shapes) for s in before), "tables belong in the appendix"
    titled = [s.shapes.title.text for s in before[2:] if s.shapes.title is not None]
    assert sum(1 for t in titled if re.search(r"\d", t)) >= len(titled) // 2, titled
    assert layouts.count("Title Only") < len(layouts) // 2


# -- 5. layout -------------------------------------------------------------------------------------


def test_the_designed_cover_is_a_cover(rendered) -> None:  # type: ignore[no-untyped-def]
    path, payload = _file(rendered, "cfo_variance_memo", "pdf")
    read = extract(path, payload)
    assert read is not None
    cover = read.units[0].text
    for block in ("CONFIDENTIAL", "Summary", "In this paper", "Distribution", "Document"):
        assert block in cover, block
    control = read.units[2].text if len(read.units) > 2 else ""
    assert "Review record" in control and "Approval and sign-off form" in control


def test_word_keeps_short_tables_whole_and_rows_unbroken(rendered) -> None:  # type: ignore[no-untyped-def]
    _, payload = _file(rendered, "cfo_variance_memo", "docx")
    with zipfile.ZipFile(BytesIO(payload)) as archive:
        body = archive.read("word/document.xml").decode()
    assert "<w:cantSplit/>" in body and "<w:keepNext/>" in body and "<w:widowControl/>" in body


def test_the_request_contract_rejects_a_move_outside_the_allowed_set() -> None:
    from worldloom.narrative.requests import RequestMove

    with pytest.raises(ValueError):
        NarrativeRequest(artifact_id="ART-1", artifact_type="x", section="s", persona_id="", voice="",
                         audience="", author_title="", allowed_fact_ids=["FACT-1"],
                         moves=[RequestMove(name="headline", instruction="i", fact_ids=["FACT-2"])])
