from __future__ import annotations

import base64
import hashlib
from io import BytesIO

import pytest

from worldloom.native_tasks import (
    NativeAnswer,
    NativeAssertion,
    NativeCalculation,
    NativeCitation,
    NativeFile,
    NativeInput,
    NativeOutput,
    NativeSubmission,
    NativeTask,
    grade_native_task,
    public_contract,
)


def workbook(*, a: int = 10, b: int = 20, formula: str = "=A1+B1") -> bytes:
    openpyxl = pytest.importorskip("openpyxl")
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = "Sheet1"
    sheet["A1"], sheet["B1"], sheet["C1"] = a, b, formula
    out = BytesIO()
    book.save(out)
    return out.getvalue()


def ref(cell: str, artifact: str = "source") -> NativeCitation:
    return NativeCitation(artifact_id=artifact, locator=f"sheet:Sheet1/cell:{cell}")


def input_file(payload: bytes) -> NativeInput:
    return NativeInput(artifact_id="source", format="xlsx", path="files/source.xlsx", sha256=hashlib.sha256(payload).hexdigest())


def test_analysis_uses_bytes_and_requires_all_actual_citations() -> None:
    payload = workbook()
    task = NativeTask(id="sum", operation="analyze", inputs=(input_file(payload),), assertions=(
        NativeAssertion(id="total", calculation=NativeCalculation(operation="sum", operands=(ref("A1"), ref("B1")))),
    ))
    answer = NativeAnswer(assertion_id="total", value="30", citations=(ref("A1"), ref("B1")))
    assert grade_native_task(task, {"source": payload}, NativeSubmission(answers=(answer,))).passed
    for bad in (answer.model_copy(update={"value": "31"}), answer.model_copy(update={"citations": (ref("A1"),)})):
        assert not grade_native_task(task, {"source": payload}, NativeSubmission(answers=(bad,))).passed
    assert not grade_native_task(task, {"source": workbook(a=11)}, NativeSubmission(answers=(answer,))).passed
    unbound = task.model_copy(update={"inputs": (input_file(payload).model_copy(update={"sha256": ""}),)})
    assert not grade_native_task(unbound, {"source": payload}, NativeSubmission(answers=(answer,))).passed


def test_read_exact_evidence_and_hidden_truth() -> None:
    payload = workbook()
    task = NativeTask(id="read", prompt="Read A1", operation="read", inputs=(input_file(payload),), assertions=(NativeAssertion(id="value", target=ref("A1"), expected="10"),))
    answer = NativeAnswer(assertion_id="value", value="10", citations=(ref("A1"),))
    assert grade_native_task(task, {"source": payload}, NativeSubmission(answers=(answer,))).passed
    public = public_contract(task)
    assert public["answers"] == [{"assertion_id": "value", "requires_citations": True}]
    assert "expected" not in str(public)
    assert "calculation" not in str(public)
    assert not grade_native_task(task, {"source": payload}, NativeSubmission(answers=(answer, answer))).passed


def test_update_checks_version_change_and_all_untouched_cells_and_formulas() -> None:
    payload = workbook()
    task = NativeTask(id="update", operation="update", inputs=(input_file(payload),), output=NativeOutput(
        artifact_id="updated", format="xlsx", source_artifact_id="source",
        assertions=(NativeAssertion(id="change-a", target=ref("A1", "updated"), expected="15"),),
    ))

    def submit(data: bytes, sha: str | None = None) -> NativeSubmission:
        return NativeSubmission(files=(NativeFile(artifact_id="updated", format="xlsx", content_base64=base64.b64encode(data).decode(), source_sha256=sha or input_file(payload).sha256),))

    assert grade_native_task(task, {"source": payload}, submit(workbook(a=15))).passed
    for data in (payload, workbook(a=15, b=999), workbook(a=15, formula="=A1-B1")):
        assert not grade_native_task(task, {"source": payload}, submit(data)).passed
    assert not grade_native_task(task, {"source": payload}, submit(workbook(a=15), "0" * 64)).passed
    bad = NativeSubmission(files=(NativeFile(artifact_id="updated", format="xlsx", content_base64="not-base64"),))
    assert not grade_native_task(task, {"source": payload}, bad).passed


def test_create_requires_native_bytes_and_expected_content() -> None:
    task = NativeTask(id="create", operation="create", output=NativeOutput(artifact_id="new", format="xlsx", assertions=(NativeAssertion(id="formula", target=ref("C1", "new"), expected="=A1+B1"),)))
    def submit(payload: bytes) -> NativeSubmission:
        return NativeSubmission(files=(NativeFile(artifact_id="new", format="xlsx", content_base64=base64.b64encode(payload).decode()),))
    assert grade_native_task(task, {}, submit(workbook())).passed
    assert not grade_native_task(task, {}, submit(b"fake workbook")).passed
    assert not grade_native_task(task, {}, submit(workbook(formula="=A1-B1"))).passed
    assert not grade_native_task(task, {}, NativeSubmission()).passed


def test_empty_and_ambiguous_contracts_refuse() -> None:
    with pytest.raises(ValueError):
        NativeTask(id="empty", operation="read")
    with pytest.raises(ValueError):
        NativeAssertion(id="empty")
    with pytest.raises(ValueError):
        NativeInput(artifact_id="source", format="xlsx", path="../escape")
    with pytest.raises(ValueError):
        NativeCalculation(operation="ratio", operands=(ref("A1"),))


@pytest.mark.parametrize("format", ["docx", "pptx"])
def test_document_and_slide_updates_preserve_other_text_notes_and_hidden_state(format: str) -> None:
    def document(value: str, *, collateral: bool = False) -> bytes:
        out = BytesIO()
        if format == "docx":
            docx = pytest.importorskip("docx")
            doc = docx.Document()
            doc.add_paragraph(value)
            doc.add_paragraph("changed" if collateral else "Preserve the supplier contract")
            doc.save(out)
        else:
            pptx = pytest.importorskip("pptx")
            deck = pptx.Presentation()
            slide = deck.slides.add_slide(deck.slide_layouts[6])
            shape = slide.shapes.add_textbox(0, 0, 1000000, 1000000)
            shape.text = value
            slide.notes_slide.notes_text_frame.text = "changed" if collateral else "Preserve the source caveat"
            deck.save(out)
        return out.getvalue()

    locator = "paragraph:1" if format == "docx" else "slide:1/shape:1/text"
    source = document("Old operational forecast")
    input = NativeInput(artifact_id="source", format=format, path=f"files/source.{format}", sha256=hashlib.sha256(source).hexdigest())
    task = NativeTask(id="update", operation="update", inputs=(input,), output=NativeOutput(
        artifact_id="updated", format=format, source_artifact_id="source", assertions=(
            NativeAssertion(id="revision", target=NativeCitation(artifact_id="updated", locator=locator), expected="Revised operational forecast"),
        ),
    ))
    def grade(payload: bytes):
        return grade_native_task(task, {"source": source}, NativeSubmission(files=(NativeFile(
            artifact_id="updated", format=format, source_sha256=input.sha256, content_base64=base64.b64encode(payload).decode(),
        ),)))
    assert grade(document("Revised operational forecast")).passed
    assert "update_unaffected_content_changed" in grade(document("Revised operational forecast", collateral=True)).findings
    assert not grade(source).passed
    if format == "pptx":
        from pptx import Presentation
        deck = Presentation(BytesIO(document("Revised operational forecast")))
        deck.slides[0]._element.set("show", "0")
        out = BytesIO()
        deck.save(out)
        assert "update_unaffected_content_changed" in grade(out.getvalue()).findings


@pytest.mark.parametrize("format,locator", [("docx", "paragraph:1"), ("pptx", "slide:1/shape:1/text"), ("xlsx", "sheet:Sheet1/cell:A1")])
def test_reference_proves_create_and_update_and_rejects_impossible_targets(format: str, locator: str) -> None:
    from worldloom.native_reference import qualify_native_task, reference_submission
    create = NativeTask(id="create", operation="create", output=NativeOutput(artifact_id="new", format=format, assertions=(NativeAssertion(id="content", target=NativeCitation(artifact_id="new", locator=locator), expected="Company process evidence"),)))
    proof = qualify_native_task(create, {})
    assert proof.passed, proof.findings
    submitted = reference_submission(create, {})
    assert submitted == reference_submission(create, {})
    payload = base64.b64decode(submitted.files[0].content_base64)
    update = NativeTask(id="update", operation="update", inputs=(NativeInput(artifact_id="source", format=format, path=f"source.{format}", sha256=hashlib.sha256(payload).hexdigest()),), output=NativeOutput(artifact_id="updated", format=format, source_artifact_id="source", assertions=(NativeAssertion(id="change", target=NativeCitation(artifact_id="updated", locator=locator), expected="Revised company process evidence"),)))
    proof = qualify_native_task(update, {"source": payload})
    assert proof.passed, proof.findings
    read = NativeTask(id="impossible", operation="read", inputs=update.inputs, assertions=(NativeAssertion(id="missing", target=NativeCitation(artifact_id="source", locator="does-not-exist")),))
    assert not qualify_native_task(read, {"source": payload}).passed
    bad_truth = read.model_copy(update={"assertions": (NativeAssertion(id="wrong", target=NativeCitation(artifact_id="source", locator=locator), expected="Wrong truth"),)})
    assert not qualify_native_task(bad_truth, {"source": payload}).passed


def test_update_rejects_formula_to_text_and_preserves_numeric_cell_type() -> None:
    from openpyxl import load_workbook

    from worldloom.native_reference import reference_submission

    payload = workbook()
    task = NativeTask(id="typed-update", operation="update", inputs=(input_file(payload),), output=NativeOutput(
        artifact_id="updated", format="xlsx", source_artifact_id="source",
        assertions=(NativeAssertion(id="change-a", target=ref("A1", "updated"), expected="15"),),
    ))
    reference = reference_submission(task, {"source": payload})
    assert grade_native_task(task, {"source": payload}, reference).passed
    generated = base64.b64decode(reference.files[0].content_base64)
    book = load_workbook(BytesIO(generated))
    assert book["Sheet1"]["A1"].data_type == "n"
    decimal_task = task.model_copy(update={"output": task.output.model_copy(update={"assertions": (
        task.output.assertions[0].model_copy(update={"expected": "17.20"}),
    )})})
    decimal_reference = reference_submission(decimal_task, {"source": payload})
    assert grade_native_task(decimal_task, {"source": payload}, decimal_reference).passed
    book["Sheet1"]["C1"].data_type = "s"
    out = BytesIO()
    book.save(out)
    forged = reference.model_copy(update={"files": (reference.files[0].model_copy(update={"content_base64": base64.b64encode(out.getvalue()).decode()}),)})
    grade = grade_native_task(task, {"source": payload}, forged)
    assert not grade.passed
    assert "update_unaffected_content_changed" in grade.findings


def test_create_formula_requires_formula_cell_and_allows_explicit_literal() -> None:
    from openpyxl import load_workbook

    from worldloom.native_reference import reference_submission

    task = NativeTask(id="typed-create", operation="create", output=NativeOutput(artifact_id="new", format="xlsx", assertions=(
        NativeAssertion(id="formula", target=ref("C1", "new"), expected="=A1+B1"),
    )))
    reference = reference_submission(task, {})
    book = load_workbook(BytesIO(base64.b64decode(reference.files[0].content_base64)))
    book["Sheet1"]["C1"].data_type = "s"
    out = BytesIO()
    book.save(out)
    forged = reference.model_copy(update={"files": (reference.files[0].model_copy(update={"content_base64": base64.b64encode(out.getvalue()).decode()}),)})
    assert "output_type_failed:formula" in grade_native_task(task, {}, forged).findings
    literal = task.model_copy(update={"output": task.output.model_copy(update={"assertions": (
        task.output.assertions[0].model_copy(update={"expected_type": "string"}),
    )})})
    assert grade_native_task(literal, {}, reference_submission(literal, {})).passed
