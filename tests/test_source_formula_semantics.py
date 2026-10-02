"""Source facts, declared computations, and executable sheets must agree."""
from __future__ import annotations

from dataclasses import replace
from datetime import UTC, datetime
from decimal import ROUND_HALF_UP, Decimal
from io import BytesIO

import pytest
from openpyxl import load_workbook
from pydantic import ValidationError

from tests.test_render import evaluate
from worldloom import (
    BankingWorld,
    MonthEndClose,
    QuarterlyCapitalReturn,
    RetailWorld,
    columns,
    doctypes,
)
from worldloom.formula_semantics import formula_value
from worldloom.generators.finance import _pct
from worldloom.models import (
    ArtifactIR,
    ArtifactSection,
    Authority,
    CanonicalFact,
    Cell,
    Column,
    Company,
    FormulaKind,
    Quantity,
    Row,
    Table,
)
from worldloom.native_artifacts import inspect_artifact
from worldloom.native_business import plan_business_content
from worldloom.native_corpus import (
    NativeContent,
    NativeCorpusPlan,
    render_native_corpus,
)
from worldloom.native_query_evidence import TableEvidenceIndex, validate_table_binding
from worldloom.render.values import corpus_locale
from worldloom.render.xlsx import render
from worldloom.world import World


@pytest.fixture(scope="module")
def retail() -> World:
    return RetailWorld(seed=8128).build().run(MonthEndClose(period="2026-03")).compile()


@pytest.fixture(scope="module")
def banking() -> World:
    return BankingWorld(seed=8128).build().run(QuarterlyCapitalReturn(period="2026-03")).compile()


@pytest.mark.parametrize("numerator,denominator,expected", [
    (201, 20000, 1.01), (-201, 20000, -1.01), (201, -20000, -1.01),
    (1, 8, 12.5), (0, 1, 0), (1, 0, 0), (0, 0, 0),
])
def test_canonical_percentage_rounding_matches_signed_business_boundaries(
    numerator: int, denominator: int, expected: float,
) -> None:
    assert _pct(numerator, denominator) == expected
    assert formula_value(FormulaKind.RATIO_PCT, (numerator, denominator), decimal_places=2) == expected


def test_percentage_point_change_remains_difference_not_ratio() -> None:
    assert formula_value(FormulaKind.DIFFERENCE, (24.89, 23.75), decimal_places=2) == 1.14
    assert formula_value(FormulaKind.RATIO_PCT, (24.89, 23.75), decimal_places=2) == 104.8
    assert formula_value(FormulaKind.SUM, (1.005, 0), decimal_places=2) == 1.01
    assert formula_value(FormulaKind.REFERENCE, (-1.005,), decimal_places=2) == -1.01


@pytest.mark.parametrize("profit,revenue,expected", [(201, 20000, 1.01), (-201, 20000, -1.01), (1, 0, 0)])
@pytest.mark.parametrize("number_format", ["0.00%", '0.00"%"'])
def test_rounding_boundaries_survive_canonical_ir_and_both_workbook_surfaces(
    profit: int, revenue: int, expected: float, number_format: str,
) -> None:
    facts = tuple(CanonicalFact(id=f"FACT-{index}", kind=kind, subject="CO-1", period="2026-03",
        value=Quantity(amount=value, unit=unit), valid_from=datetime(2026, 3, 1, tzinfo=UTC),
        authority=Authority.SYSTEM_OF_RECORD) for index, (kind, value, unit) in enumerate((
            ("profit", profit, "AUD"), ("revenue", revenue, "AUD"), ("margin", _pct(profit, revenue), "percent"))))
    table = Table(key="margin", title="Margin", columns=[Column(key="profit", label="Profit"),
        Column(key="revenue", label="Revenue"), Column(key="margin", label="Margin", number_format=number_format)],
        rows=[Row(key="group", label="Group", cells={
            "profit": Cell(value=profit, fact_id=facts[0].id),
            "revenue": Cell(value=revenue, fact_id=facts[1].id),
            "margin": Cell(value=facts[2].value.amount, fact_id=facts[2].id, formula=FormulaKind.RATIO_PCT,
                operands=["profit", "revenue"], formula_decimal_places=2)})])
    artifact = ArtifactIR(id="ART-1", intent_id="INTENT-1", title="Margin statement",
                          sections=[ArtifactSection(heading="Margin", table=table)])
    world = World(company=Company(id="CO-1", name="Retail", industry="retail", headquarters="Sydney",
        fiscal_year_start_month=7, employees_total=10), _facts=facts, _artifact_irs=(artifact,))
    workbook = load_workbook(BytesIO(render(artifact)))
    divisor = 100 if number_format == "0.00%" else 1
    assert evaluate(workbook, "Margin", "D4") == pytest.approx(expected / divisor, abs=1e-12)
    result = render_native_corpus(world, NativeCorpusPlan(artifact_id="native", format="xlsx",
        title="Margin statement", surface="business", contents=(NativeContent(source_artifact_id="ART-1", section_index=0),)))
    native = load_workbook(BytesIO(result.payload))
    entry = next(entry for entry in result.manifest.evidence if entry.kind == "formula")
    assert entry.value == expected
    assert native["Margin"]["D2"].value == "=ROUND(IF('Margin'!C2=0,0,'Margin'!B2/'Margin'!C2*100),2)"

    # Updating the source literal and its cited fact cannot authorize a wrong
    # computation; the original operands must still produce the declared value.
    wrong = table.rows[0].cells["margin"].model_copy(update={"value": expected + 0.01})
    corrupted_table = table.model_copy(update={"rows": [table.rows[0].model_copy(update={"cells": {
        **table.rows[0].cells, "margin": wrong}})]})
    corrupted_artifact = artifact.model_copy(update={"sections": [ArtifactSection(heading="Margin", table=corrupted_table)]})
    corrupted = replace(world, _facts=(*facts[:2], facts[2].model_copy(update={"value": Quantity(
        amount=expected + 0.01, unit="percent")})), _artifact_irs=(corrupted_artifact,))
    with pytest.raises(ValueError, match="formula literal disagrees"):
        render_native_corpus(corrupted, NativeCorpusPlan(artifact_id="native", format="xlsx",
            title="Invalid margin", surface="business", contents=(NativeContent(source_artifact_id="ART-1", section_index=0),)))


def test_rounded_cross_sheet_reference_converts_percentage_storage_to_value() -> None:
    source = Table(key="source", title="Source", columns=[Column(key="rate", label="Rate", number_format="0.00%")],
        rows=[Row(key="a", label="Actual", cells={"rate": Cell(value=24.895)})])
    summary = Table(key="summary", title="Summary", columns=[Column(key="margin", label="Margin")],
        rows=[Row(key="a", label="Actual", cells={"margin": Cell(value=24.9, formula=FormulaKind.REFERENCE,
            operands=["source:a:rate"], formula_decimal_places=2)})])
    artifact = ArtifactIR(id="ART-1", intent_id="INTENT-1", title="Margin statement", sections=[
        ArtifactSection(heading="Source", table=source), ArtifactSection(heading="Summary", table=summary)])
    workbook = load_workbook(BytesIO(render(artifact)))
    assert evaluate(workbook, "Summary", "B4") == 24.9


def test_precision_is_explicit_serialized_and_authorable() -> None:
    old = Cell(value=1)
    assert "formula_decimal_places" not in old.model_dump()
    cell = Cell(value=1.01, formula=FormulaKind.RATIO_PCT, operands=["a", "b"], formula_decimal_places=2)
    assert Cell.model_validate_json(cell.model_dump_json()) == cell
    for kwargs in ({"value": 1, "formula_decimal_places": 2},
                   {"formula": "ratio_pct", "operands": ["a", "b"], "formula_decimal_places": -1}):
        with pytest.raises(ValidationError):
            Cell.model_validate(kwargs)
    legacy = doctypes.DerivationSpec(formula=FormulaKind.RATIO_PCT, operands=["a", "b"])
    assert "decimal_places" not in legacy.model_dump()
    authored = doctypes.DerivationSpec(formula=FormulaKind.RATIO_PCT, operands=["a", "b"], decimal_places=3)
    assert authored.as_derivation().decimal_places == 3
    assert columns.cut(columns.PNL, "divisions").precisions() == {"gm_pct_actual": 2}


def test_every_generated_retail_ratio_recomputes_canonical_literal_and_excel(retail: World) -> None:
    rendered = retail.render("xlsx")
    payload = next(item.payload for item in rendered._rendered if item.path.endswith(".xlsx"))
    workbook = load_workbook(BytesIO(payload))
    facts = {fact.id: fact for fact in retail.facts}
    checked = 0
    for artifact in retail.artifact_irs:
        for section in artifact.sections:
            table = section.table
            if table is None:
                continue
            for row_index, row in enumerate(table.rows, 4):
                for column_index, column in enumerate(table.columns, 2):
                    cell = row.cells.get(column.key)
                    if cell is None or cell.formula is not FormulaKind.RATIO_PCT:
                        continue
                    numerator, denominator = (Decimal(str(row.cells[key].value)) for key in cell.operands)
                    expected = (numerator / denominator * 100 if denominator else Decimal(0)).quantize(
                        Decimal("0.01"), rounding=ROUND_HALF_UP)
                    assert cell.formula_decimal_places == 2
                    assert cell.value == float(expected)
                    assert facts[cell.fact_id].value.amount == cell.value
                    if table.title[:31] in workbook.sheetnames:
                        target = workbook[table.title[:31]].cell(row_index, column_index)
                        assert target.value.startswith("=ROUND(")
                        assert target.value.endswith(",2)/100")
                        assert evaluate(workbook, table.title[:31], target.coordinate) == pytest.approx(
                            cell.value / 100, abs=1e-12)
                    checked += 1
    assert checked >= 30


def test_rounded_retail_tables_are_native_evidence_and_formula_bytes_are_bound(retail: World) -> None:
    contents, exclusions = plan_business_content(retail)
    selected = {(item.source_artifact_id, item.section_index) for item in contents}
    for artifact in retail.artifact_irs:
        for index, section in enumerate(artifact.sections):
            if section.table is not None and section.table.key in {"pnl", "category", "divisions"}:
                assert (artifact.id, index) in selected
    assert not any("formula literal disagrees" in item.reason for item in exclusions)
    result = render_native_corpus(retail, NativeCorpusPlan(artifact_id="native", format="xlsx",
        title="Financial tables", contents=contents, surface="business"))
    workbook = load_workbook(BytesIO(result.payload))
    snapshot = inspect_artifact(result.payload, "xlsx")
    units = {unit.locator: unit.text for unit in snapshot.units}
    irs = {artifact.id: artifact for artifact in retail.artifact_irs}
    index = TableEvidenceIndex(result.manifest.evidence, irs, world=retail)
    facts = {fact.id: fact for fact in retail.facts}
    rounded = 0
    for entry in result.manifest.evidence:
        if entry.kind != "formula":
            continue
        validate_table_binding(entry, irs[entry.source_artifact_id], index, units, facts,
                               format="xlsx", locale=corpus_locale(retail))
        _, cell = index.cells[entry.source_artifact_id, entry.table_key, entry.row_key, entry.column_key]
        if cell.formula_decimal_places is not None:
            sheet, coordinate = entry.locator.removeprefix("sheet:").split("/cell:")
            assert workbook[sheet][coordinate].value.startswith("=ROUND(")
            assert workbook[sheet][coordinate].value.endswith(",2)")
            rounded += 1
    assert rounded >= 30


def test_banking_subtotals_name_only_measured_children_and_remain_reconciled(banking: World) -> None:
    contents, exclusions = plan_business_content(banking)
    selected = {(item.source_artifact_id, item.section_index) for item in contents}
    checked = 0
    for artifact in banking.artifact_irs:
        if "Divisional and Branch" not in artifact.title:
            continue
        for index, section in enumerate(artifact.sections):
            if section.table is None or section.hidden:
                continue
            table = section.table
            if not any(cell.formula is FormulaKind.SUM for row in table.rows for cell in row.cells.values()):
                continue
            assert (artifact.id, index) in selected
            rows = {row.key: row for row in table.rows}
            for row in table.rows:
                for key, cell in row.cells.items():
                    if cell.formula is not FormulaKind.SUM:
                        continue
                    operands = [rows[child].cells[key] for child in cell.operands]
                    assert all(operand.fact_id and operand.value is not None for operand in operands)
                    assert sum(operand.value for operand in operands) == cell.value
                    checked += 1
    assert checked >= 10
    assert not any("formula" in item.reason for item in exclusions)


def test_banking_absent_measure_does_not_invent_zero_sum(banking: World) -> None:
    # Resolve the shipped kind by the network table, not an assumed fact name.
    artifact = next(ir for ir in banking.artifact_irs if "Divisional and Branch" in ir.title)
    first = artifact.sections[0].table.rows[0].cells["deposits"]
    kind = banking.facts.by_id(first.fact_id).kind
    remaining = tuple(fact for fact in banking.facts
        if fact.kind != kind or fact.subject == banking.company.id)
    present = {fact.id for fact in remaining}
    intent = banking.artifact_intents.by_id(artifact.intent_id)
    intent = intent.model_copy(update={"required_fact_ids": [identifier for identifier in intent.required_fact_ids
                                                            if identifier in present]})
    removed = replace(banking, _facts=remaining, _artifact_intents=(intent,)).compile()
    artifact = next(ir for ir in removed.artifact_irs if "Divisional and Branch" in ir.title)
    total = artifact.sections[0].table.rows[-1].cells["deposits"]
    assert total.fact_id is not None and total.value is not None
    assert total.formula is None and total.operands == []


def test_banking_false_parent_total_still_rejected(banking: World) -> None:
    artifact = next(ir for ir in banking.artifact_irs if "Divisional and Branch" in ir.title)
    parent = artifact.sections[0].table.rows[-1].cells["deposits"]
    changed = replace(banking, _facts=tuple(fact.model_copy(update={"value": fact.value.model_copy(
        update={"amount": fact.value.amount + 1})}) if fact.id == parent.fact_id else fact for fact in banking.facts)).compile()
    with pytest.raises(ValueError, match="formula literal disagrees"):
        render_native_corpus(changed, NativeCorpusPlan(artifact_id="native", format="xlsx",
            title="Invalid totals", surface="business", contents=(NativeContent(
                source_artifact_id=artifact.id, section_index=0),)))
