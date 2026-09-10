from __future__ import annotations

from io import BytesIO
from zipfile import ZipFile

import pytest
from docx import Document
from openpyxl import Workbook
from openpyxl.comments import Comment
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.util import Inches

from worldloom.native_artifacts import (
    NativeArtifactError,
    inspect_artifact,
    package_digests,
)


def _bytes(document):
    stream = BytesIO()
    document.save(stream)
    return stream.getvalue()


def test_docx_extracts_late_evidence_and_table_and_explicit_pages():
    document = Document()
    for page in range(120):
        if page:
            document.add_page_break()
        document.add_paragraph(f"Process control {page + 1}")
    document.add_paragraph("Approved supplier variance: 17")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Owner"
    table.cell(0, 1).text = "Supply Chain"
    payload = _bytes(document)
    snapshot = inspect_artifact(payload, "docx")
    assert snapshot.metrics["explicit_pages"] == 120
    assert "pages" not in snapshot.metrics  # No claim of rendered pagination.
    assert snapshot.units[-1].text == "Supply Chain"
    assert any(unit.text == "Approved supplier variance: 17" for unit in snapshot.units)
    assert inspect_artifact(payload, "docx") == snapshot
    assert "word/document.xml" in package_digests(payload, "docx")


def test_pptx_reads_hidden_slide_notes_table_and_chart():
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    slide._element.set("show", "0")
    box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1))
    box.text = "Quarter-end supplier reconciliation"
    slide.notes_slide.notes_text_frame.text = "Invoice exceeds receipt by 17"
    table = slide.shapes.add_table(1, 1, Inches(1), Inches(2), Inches(2), Inches(1)).table
    table.cell(0, 0).text = "Supplier A"
    data = CategoryChartData()
    data.categories = ["Received", "Invoiced"]
    data.add_series("Amount", [100, 117])
    slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(1), Inches(3), Inches(4), Inches(2), data)
    snapshot = inspect_artifact(_bytes(presentation), "pptx")
    units = {unit.locator: unit.text for unit in snapshot.units}
    assert units["slide:1/hidden"] == "true"
    assert units["slide:1/notes"] == "Invoice exceeds receipt by 17"
    assert units["slide:1/shape:2/row:1/cell:1"] == "Supplier A"
    assert units["slide:1/shape:3/series:1"] == "Amount: 100.0, 117.0"
    assert snapshot.metrics["speaker_notes"] == snapshot.metrics["hidden_slides"] == snapshot.metrics["native_charts"] == 1


def test_xlsx_preserves_formulas_comments_and_sparse_far_cells():
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Receipts"
    sheet["A1"] = 100
    sheet["A2"] = 117
    sheet["B2"] = "=A2-A1"
    sheet["B2"].comment = Comment("Price variance", "Controller")
    sheet["XFD1048576"] = "Late evidence"
    hidden = workbook.create_sheet("Source")
    hidden.sheet_state = "hidden"
    snapshot = inspect_artifact(_bytes(workbook), "xlsx")
    units = {unit.locator: unit.text for unit in snapshot.units}
    assert units["sheet:Receipts/cell:B2"] == "=A2-A1"
    assert units["sheet:Receipts/cell:B2/comment"] == "Price variance"
    assert units["sheet:Receipts/cell:XFD1048576"] == "Late evidence"
    assert snapshot.metrics["formula_cells"] == snapshot.metrics["comments"] == snapshot.metrics["hidden_sheets"] == 1
    assert len(units) == 11


def test_corrupt_mismatched_and_duplicate_packages_refuse():
    with pytest.raises(NativeArtifactError, match="malformed"):
        inspect_artifact(b"not a package", "docx")
    with pytest.raises(NativeArtifactError, match="not pptx"):
        inspect_artifact(_bytes(Document()), "pptx")
    with pytest.raises(NativeArtifactError, match="unsupported"):
        inspect_artifact(b"", "pdf")
    stream = BytesIO(_bytes(Document()))
    with ZipFile(stream, "a") as package, pytest.warns(UserWarning):
        package.writestr("word/document.xml", b"invalid")
    with pytest.raises(NativeArtifactError, match="duplicate"):
        inspect_artifact(stream.getvalue(), "docx")


def test_same_text_different_formatting_changes_exact_part_hash():
    document = Document()
    paragraph = document.add_paragraph("Preserve approved wording")
    original = _bytes(document)
    paragraph.runs[0].bold = True
    edited = _bytes(document)
    assert inspect_artifact(original, "docx").units == inspect_artifact(edited, "docx").units
    assert package_digests(original, "docx")["word/document.xml"] != package_digests(edited, "docx")["word/document.xml"]
