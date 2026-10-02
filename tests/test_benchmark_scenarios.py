"""Generated business cases earn independence through inspected evidence."""
from __future__ import annotations

from dataclasses import replace
from datetime import UTC, datetime
from io import BytesIO
from zipfile import ZipFile

import pytest
from openpyxl import load_workbook

from worldloom import RetailWorld, World
from worldloom.benchmarks.scenarios import NativeScenarioDemand, build_native_scenarios
from worldloom.corpus_scale import _world_digest
from worldloom.evalrun.qualification import audit_splits, evidence_components
from worldloom.models import Authority, CanonicalFact, FormulaKind
from worldloom.native_artifacts import inspect_artifact
from worldloom.native_eval_bridge import native_task_cases
from worldloom.native_query_planning import NativeWorkloadPlan, plan_native_workload


def test_eighteen_generated_cases_fund_analysis_in_every_native_format() -> None:
    source = RetailWorld(seed=8128).build()
    unrelated = CanonicalFact(id="FACT-UNRELATED-POLICY", kind="policy.retention", subject=source.company.id,
        text_value="Archive the source evidence", valid_from=datetime(2026, 1, 1, tzinfo=UTC),
        authority=Authority.APPROVED_REPORT)
    source = replace(source, _facts=(*source._facts, unrelated))
    built = build_native_scenarios(source, NativeScenarioDemand())
    built.world.validate().raise_if_failed()
    rendered = built.render()
    assert len(rendered) == 54
    assert len({episode.template_variant for episode in built.episodes}) == 3
    assert len({episode.process for episode in built.episodes}) == 3
    plan = NativeWorkloadPlan(use_case_id="case-analysis", objective="Reconcile the case and review its variance.",
        operations=("analyze",), discovery_scope="artifact", max_tasks=324)
    workload = plan_native_workload(built.world, rendered, plan)
    assert workload.reference_qualified == len(workload.tasks) == 324
    assert not workload.findings
    assert set(workload.capability_coverage) >= {"actual_budget_variance", "actual_budget_ratio", "authored_sum_reconciliation"}
    cases = native_task_cases(workload.tasks, rendered, world=built.world, namespace="independent-case-test")
    components = evidence_components(cases)
    assert len(set(components.values())) == 18
    assert all(unrelated.id not in case.row["expected_fact_ids"] for case in cases)
    for format in ("docx", "pptx", "xlsx"):
        tasks = {task.id for task in workload.tasks if {item.format for item in task.inputs} == {format}}
        selected = tuple(case for case in cases if case.id in tasks)
        assert len(selected) == 108
        assert len(set(evidence_components(selected).values())) == 18
    first = next(iter(components.values()))
    assert audit_splits(tuple(case for case in cases if components[case.id] == first),
        tuple(case for case in cases if components[case.id] != first)).isolated


def test_registered_recipe_replays_the_world_and_actual_office_bytes(tmp_path) -> None:
    base = RetailWorld(seed=31).build()
    demand = NativeScenarioDemand(episodes=3, start_period="2026-02", batch_id="replay-cases")
    built = build_native_scenarios(base, demand)
    built.verify_source_replay()
    assert built.world.compile()._artifact_irs == built.world._artifact_irs
    built.world.export(tmp_path / "source")
    loaded = World.load(tmp_path / "source")
    assert _world_digest(loaded) == built.source_digest
    again = build_native_scenarios(RetailWorld(seed=31).build(), demand)
    left, right = built.render(), again.render()
    assert left.keys() == right.keys()
    for identifier in left:
        assert left[identifier].payload == right[identifier].payload
        assert left[identifier].manifest == right[identifier].manifest
        result = left[identifier]
        snapshot = inspect_artifact(result.payload, result.manifest.format)
        assert not any(fact.id in unit.text for fact in built.world.facts for unit in snapshot.units)
        assert not any(event_id in unit.text for episode in built.episodes for event_id in episode.event_ids
                       for unit in snapshot.units)
        private_ids = [fact.id.encode() for fact in built.world.facts] + [identifier.encode()
            for episode in built.episodes for identifier in episode.event_ids]
        with ZipFile(BytesIO(result.payload)) as package:
            assert not any(identifier in data for data in (package.read(name) for name in package.namelist())
                           for identifier in private_ids)
        if result.manifest.format == "xlsx":
            workbook = load_workbook(BytesIO(result.payload))
            formulas = [cell.value for sheet in workbook for row in sheet for cell in row
                        if cell.data_type == "f"]
            assert any("SUM(" in formula for formula in formulas)
            assert any("IF(" in formula for formula in formulas)
            assert any("comparison" in section.table.key for section in built.world.artifact_irs.by_id(
                result.manifest.evidence[0].source_artifact_id).sections if section.table)


def test_every_operation_and_format_has_case_level_support() -> None:
    built = build_native_scenarios(RetailWorld(seed=8128).build(), NativeScenarioDemand(episodes=3))
    rendered = built.render()
    workload = plan_native_workload(built.world, rendered, NativeWorkloadPlan(use_case_id="case-workflows",
        objective="Read the source, reconcile the outcome, update the decision and create the review.",
        discovery_scope="artifact", max_tasks=288))
    assert workload.reference_qualified == len(workload.tasks)
    assert not workload.findings
    cases = native_task_cases(workload.tasks, rendered, world=built.world, namespace="workflow-case-test")
    for operation in ("read", "analyze", "update", "create"):
        for format in ("docx", "pptx", "xlsx"):
            tasks = {task.id for task in workload.tasks if task.operation == operation
                     and {item.format for item in task.inputs} == {format}}
            selected = tuple(case for case in cases if case.id in tasks)
            assert len(set(evidence_components(selected).values())) == 3, (operation, format)
    assert workload.capability_coverage["native_preservation"] >= 9


def test_process_arithmetic_has_distinct_units_and_case_local_ancestry() -> None:
    built = build_native_scenarios(RetailWorld(seed=19).build(), NativeScenarioDemand(episodes=3))
    known: set[str] = set()
    for episode in built.episodes:
        assert not known.intersection(episode.fact_ids)
        known.update(episode.fact_ids)
        facts = [built.world.facts.by_id(identifier) for identifier in episode.fact_ids]
        final = next(fact for fact in facts if fact.kind.endswith(".status") and fact.authority is Authority.APPROVED_REPORT)
        review = built.world.facts.by_id(final.supersedes)
        initial = built.world.facts.by_id(review.supersedes)
        assert initial.authority is Authority.INITIAL_HYPOTHESIS
        assert final.authority is Authority.APPROVED_REPORT
        assert initial.valid_to == review.valid_from < final.valid_from == review.valid_to
        assert all(set(fact.derived_from) <= set(episode.fact_ids) for fact in facts)
        assert {fact.subject for fact in facts} == {episode.event_ids[0]}
        assert built.world.entity_names()[episode.event_ids[0]].startswith({
            "supplier_reconciliation": "Supplier invoice", "customer_settlement": "Customer refund",
            "inventory_replenishment": "Inventory replenishment"}[episode.process])
        totals = {fact.kind.rsplit(".", 1)[-1]: fact.value.amount for fact in facts
                  if ".total." in fact.kind and fact.value}
        if episode.process == "supplier_reconciliation":
            assert totals["eligible"] == pytest.approx(totals["actual"] - totals["unreceipted"])
        elif episode.process == "customer_settlement":
            assert totals["actual"] == pytest.approx(totals["rejected"] + totals["paid"] + totals["pending"])
        else:
            assert totals["shortage"] == totals["actual"] - totals["stock"] - totals["inbound"]
        assert {fact.value.unit for fact in facts if fact.value} == (
            {"units"} if episode.process == "inventory_replenishment" else {built.world.company.currency})
        tables = built.world.artifact_irs.by_id(episode.source_artifact_id).tables()
        assert {cell.formula for table in tables for row in table.rows for cell in row.cells.values()} >= {
            FormulaKind.SUM, FormulaKind.DIFFERENCE, FormulaKind.RATIO_PCT, FormulaKind.REFERENCE}


def test_duplicate_batch_mutated_source_and_unrecorded_changes_refuse() -> None:
    built = build_native_scenarios(RetailWorld(seed=41).build(), NativeScenarioDemand(episodes=3))
    with pytest.raises(ValueError, match="batch already exists"):
        build_native_scenarios(built.world, built.demand)
    fact = built.world.facts.by_id(built.episodes[0].fact_ids[0])
    changed = fact.model_copy(update={"text_value": "Unrecorded change"})
    world = replace(built.world, _facts=tuple(changed if item.id == fact.id else item for item in built.world.facts))
    with pytest.raises(ValueError, match="source changed"):
        replace(built, world=world).render()
    with pytest.raises(ValueError, match="does not replay"):
        replace(built, world=world, source_digest=_world_digest(world)).verify_source_replay()


def test_fresh_batches_have_distinct_case_subjects_in_the_same_period() -> None:
    base = RetailWorld(seed=8128).build()
    before = base.entity_names()
    first = build_native_scenarios(base, NativeScenarioDemand(episodes=3, batch_id="first-review"))
    second = build_native_scenarios(first.world, NativeScenarioDemand(episodes=3, batch_id="followup-review"))
    second.world.validate().raise_if_failed()
    second.verify_source_replay()
    first_subjects = {first.world.facts.by_id(identifier).subject for episode in first.episodes for identifier in episode.fact_ids}
    second_subjects = {second.world.facts.by_id(identifier).subject for episode in second.episodes for identifier in episode.fact_ids}
    assert not first_subjects.intersection(second_subjects)
    assert {episode.period for episode in first.episodes} == {episode.period for episode in second.episodes}
    assert all(second.world.entity_names()[identifier] == name for identifier, name in before.items())


@pytest.mark.parametrize("arguments", [{"episodes": True}, {"episodes": 257}, {"rows_per_episode": 9},
    {"episodes": 2}, {"start_period": "2026-13"}, {"processes": ()}, {"batch_id": "../escape"}])
def test_case_demand_is_bounded_and_explicit(arguments) -> None:
    with pytest.raises(ValueError):
        NativeScenarioDemand(**arguments)
