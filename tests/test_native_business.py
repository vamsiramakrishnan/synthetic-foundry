from __future__ import annotations

import hashlib
from dataclasses import replace
from datetime import UTC, datetime
from io import BytesIO
from typing import Literal
from zipfile import ZipFile

import pytest

from worldloom.models import (
    ArtifactIR,
    ArtifactSection,
    Authority,
    BusinessUnit,
    CanonicalFact,
    Cell,
    Column,
    Company,
    FormulaKind,
    Quantity,
    Row,
    Table,
)
from worldloom.native_artifacts import inspect_artifact
from worldloom.native_corpus import (
    NativeContent,
    NativeContentProvenance,
    NativeCorpusPlan,
    native_corpus_coverage,
    plan_native_corpus,
    render_native_corpus,
)
from worldloom.world import World


def _world(*, prose: bool = True) -> World:
    facts = tuple(CanonicalFact(id=f"FACT-{index:04}", kind="financial.revenue." + measure,
        subject=subject, period="2026-01", value=Quantity(amount=amount, unit="AUD"),
        valid_from=datetime(2026, 1, 1, tzinfo=UTC), authority=Authority.SYSTEM_OF_RECORD)
        for index, (subject, measure, amount) in enumerate((
            ("BU-NORTH", "actual", 125.5), ("BU-NORTH", "budget", 100),
            ("BU-SOUTH", "actual", 75), ("BU-SOUTH", "budget", 80))))
    columns = [Column(key=key, label=label, number_format='0.00%' if key == "variance_pct" else '#,##0.00')
        for key, label in (("actual", "Actual revenue"), ("budget", "Budget revenue"),
                           ("variance", "Variance"), ("variance_pct", "Variance percent"))]
    rows = []
    for row_key, label, first, second in (("north", "North division", facts[0], facts[1]),
                                         ("south", "South division", facts[2], facts[3])):
        actual, budget = first.value.amount, second.value.amount
        rows.append(Row(key=row_key, label=label, cells={
            "actual": Cell(value=actual, fact_id=first.id),
            "budget": Cell(value=budget, fact_id=second.id),
            "variance": Cell(value=actual - budget, formula=FormulaKind.DIFFERENCE, operands=["actual", "budget"]),
            "variance_pct": Cell(value=(actual - budget) / budget * 100,
                                  formula=FormulaKind.RATIO_PCT, operands=["variance", "budget"]),
        }))
    rows.append(Row(key="total", label="Group total", emphasis=True, cells={
        "actual": Cell(value=200.5, formula=FormulaKind.SUM, operands=["north", "south"]),
        "budget": Cell(value=180, formula=FormulaKind.SUM, operands=["north", "south"]),
        "variance": Cell(value=20.5, formula=FormulaKind.DIFFERENCE, operands=["actual", "budget"]),
        "variance_pct": Cell(value=20.5 / 180 * 100, formula=FormulaKind.RATIO_PCT, operands=["variance", "budget"]),
    }))
    summary = Table(key="summary", title="Management review", columns=[Column(key="revenue", label="Revenue")],
        rows=[Row(key="group", label="Group actual revenue", cells={"revenue": Cell(value=200.5,
            formula=FormulaKind.REFERENCE, operands=["revenue:total:actual"])})])
    sections = [ArtifactSection(heading="Division revenue", table=Table(key="revenue", title="Revenue [AUD]",
        columns=columns, rows=rows), body="North actual revenue: {{fact:FACT-0000}}." if prose else None),
        ArtifactSection(heading="Management review", table=summary)]
    return World(company=Company(id="CO-1", name="Northstar Retail", industry="retail", headquarters="Sydney",
        fiscal_year_start_month=7, employees_total=100), _facts=facts,
        _business_units=(BusinessUnit(id="BU-NORTH", name="North division", company_id="CO-1", leader_id="PERSON-NORTH", kind="division"),
                         BusinessUnit(id="BU-SOUTH", name="South division", company_id="CO-1", leader_id="PERSON-SOUTH", kind="division")),
        _artifact_irs=(ArtifactIR(id="ART-SOURCE", intent_id="INTENT-1", title="Revenue review", sections=sections),))


def _plan(world: World, format: Literal["docx", "pptx", "xlsx"]) -> NativeCorpusPlan:
    return plan_native_corpus(world, artifact_id="ART-NATIVE", format=format, title="Revenue review",
        minimum_units=2, minimum_distinct_facts=4, surface="business")


def _explicit_plan(world: World, format: Literal["docx", "pptx", "xlsx"]) -> NativeCorpusPlan:
    return NativeCorpusPlan(artifact_id="ART-NATIVE", format=format, title="Revenue review", surface="business",
        minimum_units=2, minimum_distinct_facts=4, contents=tuple(NativeContent(source_artifact_id="ART-SOURCE", section_index=i)
        for i in range(len(world._artifact_irs[0].sections))))


@pytest.mark.parametrize("format", ["docx", "pptx", "xlsx"])
def test_business_tables_preserve_real_source_values_and_dependency_evidence(format: Literal["docx", "pptx", "xlsx"]) -> None:
    world = _world()
    plan = _plan(world, format)
    result = render_native_corpus(world, plan)
    assert result.payload == render_native_corpus(world, plan).payload
    snapshot = inspect_artifact(result.payload, format)
    units = {unit.locator: unit.text for unit in snapshot.units}
    assert all(fact.id not in unit.text for fact in world.facts for unit in snapshot.units)
    assert result.manifest.distinct_fact_count == 4
    assert result.manifest.content_units == 2
    assert result.manifest.native_metrics == snapshot.metrics
    assert snapshot.metrics["tables"] == 2 if format != "xlsx" else snapshot.metrics["sheets"] == 3
    for entry in result.manifest.evidence:
        assert entry.text_sha256 == hashlib.sha256(units[entry.locator].encode()).hexdigest()
        assert all(locator in units for locator in entry.metadata_locators.values())
        assert all(locator in units for locator in entry.dependency_locators)
    reference = next(entry for entry in result.manifest.evidence if entry.table_key == "summary")
    assert reference.kind == "formula" and reference.value == 200.5
    assert reference.fact_ids == ("FACT-0000", "FACT-0002")
    assert len(reference.dependency_locators) == 1
    operand = next(entry for entry in result.manifest.evidence if entry.locator == reference.dependency_locators[0])
    assert operand.table_key == "revenue" and operand.row_key == "total"
    assert len(operand.dependency_locators) == 2
    if format == "xlsx":
        from openpyxl import load_workbook
        workbook = load_workbook(BytesIO(result.payload))
        assert "Facts" not in workbook.sheetnames
        assert workbook["Revenue  AUD"]["B2"].value == 125.5
        assert workbook["Revenue  AUD"]["B4"].value == "=SUM('Revenue  AUD'!B2,'Revenue  AUD'!B3)"
        assert workbook["Management review"]["B2"].value == "='Revenue  AUD'!B4"
        assert workbook["Revenue  AUD"]["E2"].value == "=IF('Revenue  AUD'!C2=0,0,'Revenue  AUD'!D2/'Revenue  AUD'!C2*100)"
        assert workbook["Revenue  AUD"]["E2"].number_format == '0.00"%"'
        assert snapshot.metrics["formula_cells"] == 9


def test_table_only_native_sources_count_grounding_without_a_writer_or_renderer() -> None:
    world = _world(prose=False)
    plan = _plan(world, "xlsx")
    coverage = native_corpus_coverage(world, plan)
    assert coverage.content_units == 2 and coverage.source_artifact_count == 1
    assert coverage.fact_ids == tuple(f.id for f in world.facts)
    result = render_native_corpus(world, plan)
    assert not any(entry.kind == "prose" for entry in result.manifest.evidence)
    assert result.manifest.distinct_fact_count == 4


def test_business_register_uses_typed_values_real_subjects_and_reporting_periods() -> None:
    world = _world()
    sections = [ArtifactSection(heading=f"{fact.subject} {fact.kind}", body=f"Reported revenue: {{{{fact:{fact.id}}}}}")
                for fact in world.facts]
    source = world._artifact_irs[0].model_copy(update={"sections": sections})
    world = replace(world, _artifact_irs=(source,))
    plan = _plan(world, "xlsx")
    result = render_native_corpus(world, plan)
    from openpyxl import load_workbook
    workbook = load_workbook(BytesIO(result.payload))
    assert workbook.sheetnames == ["Briefing", "Financial records"]
    row = tuple(cell.value for cell in workbook["Financial records"][2])
    assert row[:5] == ("Financial Revenue Actual", "North division", "2026-01", 125.5, "AUD")
    entry = next(entry for entry in result.manifest.evidence if entry.locator == "sheet:Financial records/cell:D2")
    assert entry.fact_ids == ("FACT-0000",) and entry.value == 125.5
    units = {unit.locator: unit.text for unit in inspect_artifact(result.payload, "xlsx").units}
    assert units[entry.metadata_locators["subject"]] == "North division"
    assert units[entry.metadata_locators["period"]] == "2026-01"
    with ZipFile(BytesIO(result.payload)) as package:
        assert all(fact.id.encode() not in data for fact in world.facts for data in (package.read(n) for n in package.namelist()))


@pytest.mark.parametrize("defect,message", [("literal", "disagrees with canonical"),
    ("dependency", "dependency is not selected"), ("cycle", "dependency cycle"),
    ("formula_literal", "formula literal disagrees"), ("invented_number", "lacks a canonical fact"),
    ("duplicate_column", "duplicate row or column"), ("arity", "exactly two operands")])
def test_native_business_tables_reject_false_grounding(defect: str, message: str) -> None:
    world = _world()
    source = world._artifact_irs[0]
    table = source.sections[0].table
    assert table is not None
    rows, columns = list(table.rows), list(table.columns)
    cells = dict(rows[0].cells)
    if defect == "literal":
        cells["actual"] = cells["actual"].model_copy(update={"value": 999.0})
    elif defect == "dependency":
        cells["variance"] = cells["variance"].model_copy(update={"operands": ["actual", "missing"]})
    elif defect == "cycle":
        cells["variance"] = Cell(value=25.5, formula=FormulaKind.REFERENCE, operands=["revenue:north:variance"])
    elif defect == "formula_literal":
        cells["variance"] = cells["variance"].model_copy(update={"value": 999.0})
    elif defect == "invented_number":
        cells["actual"] = Cell(value=125.5)
    elif defect == "duplicate_column":
        columns.append(columns[0])
    else:
        cells["variance"] = cells["variance"].model_copy(update={"operands": ["actual"]})
    rows[0] = rows[0].model_copy(update={"cells": cells})
    bad_table = table.model_copy(update={"rows": rows, "columns": columns})
    bad_source = source.model_copy(update={"sections": [source.sections[0].model_copy(update={"table": bad_table}), *source.sections[1:]]})
    bad_world = replace(world, _artifact_irs=(bad_source,))
    with pytest.raises(ValueError, match=message):
        native_corpus_coverage(bad_world, _explicit_plan(bad_world, "xlsx"))


def test_selected_native_sources_cannot_borrow_unselected_dependency_evidence() -> None:
    world = _world()
    plan = NativeCorpusPlan(artifact_id="ART-X", format="xlsx", title="Summary", surface="business",
        contents=(NativeContent(source_artifact_id="ART-SOURCE", section_index=1),))
    with pytest.raises(ValueError, match="dependency is not selected"):
        render_native_corpus(world, plan)


def test_business_formula_operands_cannot_mix_units() -> None:
    world = _world()
    wrong_unit = world._facts[1].model_copy(update={"value": Quantity(amount=100, unit="JPY")})
    bad_world = replace(world, _facts=(world._facts[0], wrong_unit, *world._facts[2:]))
    with pytest.raises(ValueError, match="operands use different units"):
        native_corpus_coverage(bad_world, _explicit_plan(bad_world, "xlsx"))


def test_renamed_native_tables_do_not_purchase_another_grounded_unit() -> None:
    world = _world(prose=False)
    source = world._artifact_irs[0]
    original = source.sections[0]
    assert original.table is not None
    renamed = original.model_copy(update={"heading": "Another view", "table": original.table.model_copy(update={
        "key": "another", "title": "Renamed revenue"})})
    world = replace(world, _artifact_irs=(source.model_copy(update={"sections": [*source.sections, renamed]}),))
    plan = _plan(world, "xlsx")
    assert len(plan.contents) == 2
    assert plan.excluded_contents[0].code == "duplicate_grounded_content"
    with pytest.raises(ValueError, match="duplicate grounded content"):
        native_corpus_coverage(world, _explicit_plan(world, "xlsx"))


def test_business_auto_planner_audits_private_and_invalid_sections_without_repairing_them() -> None:
    world = _world()
    source = world._artifact_irs[0]
    assert source.sections[0].table is not None
    private = source.sections[0].model_copy(update={"heading": "Private lineage", "hidden": True,
        "table": source.sections[0].table.model_copy(update={"key": "private-revenue"})})
    malformed = ArtifactSection(heading="Unguarded receipt", body="Amount: 999.99")
    world = replace(world, _artifact_irs=(source.model_copy(update={"sections": [*source.sections, private, malformed]}),))
    plan = _plan(world, "xlsx")
    assert [content.section_index for content in plan.contents] == [0, 1]
    assert [(entry.section_index, entry.code) for entry in plan.excluded_contents] == [(2, "private_native_section"), (3, "unwritten_native_section")]
    assert "excluded_contents" in plan.model_dump(mode="json")
    assert plan == _plan(world, "xlsx")


def test_business_auto_planner_selects_a_complete_dependency_closure_in_source_order() -> None:
    world = _world(prose=False)
    source = world._artifact_irs[0]
    world = replace(world, _artifact_irs=(source.model_copy(update={"sections": list(reversed(source.sections))}),))
    plan = _plan(world, "xlsx")
    assert [content.section_index for content in plan.contents] == [0, 1]
    assert not plan.excluded_contents
    assert render_native_corpus(world, plan).manifest.distinct_fact_count == 4


def test_business_spreadsheet_keeps_equals_prefixed_source_text_literal() -> None:
    world = _world()
    fact = world._facts[0].model_copy(update={"value": None, "text_value": "=RECONCILE_RECEIPT"})
    section = ArtifactSection(heading="Receiving note", body="={{fact:FACT-0000}}")
    world = replace(world, _facts=(fact,), _artifact_irs=(world._artifact_irs[0].model_copy(update={"sections": [section]}),))
    result = render_native_corpus(world, plan_native_corpus(world, artifact_id="ART-X", format="xlsx", title="Receiving",
        minimum_units=1, surface="business"))
    from openpyxl import load_workbook
    workbook = load_workbook(BytesIO(result.payload))
    assert workbook["Briefing"]["B2"].data_type == "s"
    assert workbook["Financial records"]["D2"].data_type == "s"


def test_business_identity_leakage_is_refused_and_historical_wire_shape_is_preserved() -> None:
    content = NativeContent(source_artifact_id="ART-SOURCE", section_index=0)
    plan = NativeCorpusPlan(artifact_id="ART-X", format="docx", title="Receiving", contents=(content,))
    assert "surface" not in plan.model_dump(mode="json")
    entry = NativeContentProvenance(source_artifact_id="ART-SOURCE", section_index=0, locator="paragraph:2",
        fact_ids=("FACT-0000",), text_sha256="0" * 64)
    assert set(entry.model_dump()) == {"source_artifact_id", "section_index", "locator", "fact_ids", "text_sha256", "value", "unit"}
    world = _world()
    section = world._artifact_irs[0].sections[0].model_copy(update={"heading": "FACT-0000"})
    source = world._artifact_irs[0].model_copy(update={"sections": [section, *world._artifact_irs[0].sections[1:]]})
    with pytest.raises(ValueError, match="identities leaked"):
        render_native_corpus(replace(world, _artifact_irs=(source,)), _plan(world, "docx"))
