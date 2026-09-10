from __future__ import annotations

import os
import shlex
import subprocess
import sys
from datetime import UTC, datetime

import pytest

from worldloom.models import (
    ArtifactIR,
    ArtifactSection,
    Authority,
    CanonicalFact,
    Company,
)
from worldloom.native_corpus import NativeContent, NativeCorpusPlan
from worldloom.native_tasks import (
    NativeAssertion,
    NativeCitation,
    NativeInput,
    NativeTask,
)
from worldloom.studio import RunOptions, Studio, preset
from worldloom.studio.native import execute
from worldloom.world import World


def _setup(tmp_path, monkeypatch):
    fact = CanonicalFact(id="FACT-1", kind="retail.receipt", subject="purchase-order:one",
                         text_value="Receiving confirmed", valid_from=datetime(2026, 1, 1, tzinfo=UTC),
                         authority=Authority.SYSTEM_OF_RECORD)
    world = World(company=Company(id="CO-1", name="Northstar", industry="retail", headquarters="Sydney",
                                  fiscal_year_start_month=7, employees_total=100), _facts=(fact,),
                  _artifact_irs=(ArtifactIR(id="ART-SOURCE", intent_id="INTENT-1", title="Receiving",
                      sections=[ArtifactSection(heading="Receipt", body="{{fact:FACT-1}}", fact_ids=[fact.id])]),))
    studio = Studio(tmp_path / "studio")
    monkeypatch.setattr(studio, "snapshot", lambda spec: (world, tmp_path / "snapshot"))
    base = preset()
    plan = NativeCorpusPlan(artifact_id="ART-BOOK", format="xlsx", title="Receiving",
                            contents=(NativeContent(source_artifact_id="ART-SOURCE", section_index=0),))
    task = NativeTask(id="read-one", use_case_id=base.use_cases[0].id, operation="read",
                      prompt="Read the receipt disposition in Evidence!B2 and cite it.",
                      inputs=(NativeInput(artifact_id="ART-BOOK", format="xlsx", path="book.xlsx", sha256=""),),
                      assertions=(NativeAssertion(id="disposition", target=NativeCitation(
                          artifact_id="ART-BOOK", locator="sheet:Evidence/cell:B2")),))
    spec = base.model_copy(update={"native_corpus": (plan,), "native_tasks": (task, task.model_copy(update={"id": "read-two"}))})
    project = studio.store.create(spec)
    job = studio.store.enqueue(project["id"], project["revision"], RunOptions(operation="native"))
    return studio, job


def test_native_studio_actual_read_replay_and_tamper(tmp_path, monkeypatch):
    studio, job = _setup(tmp_path, monkeypatch)
    agent = tmp_path / "reader.py"
    agent.write_text('''import json, sys
from openpyxl import load_workbook
p = json.load(sys.stdin)
assert "oracles" not in p and "expected" not in json.dumps(p["task"])
w = load_workbook(p["input_files"]["ART-BOOK"])
print(json.dumps({"request_id":p["request_id"],"submission":{"answers":[{"assertion_id":"disposition","value":w["Evidence"]["B2"].value,"citations":[{"artifact_id":"ART-BOOK","locator":"sheet:Evidence/cell:B2"}]}]}}))
''')
    argv = [sys.executable, str(agent)]
    command = subprocess.list2cmdline(argv) if os.name == "nt" else shlex.join(argv)
    result = studio.execute(job["id"], harness_command=command, timeout=30)
    assert result["passed_trials"] == result["observed_trials"] == 2
    assert result["evidence_components"] == 1
    assert result["calibrated"] is False
    import worldloom.studio.checkpoints as checkpoints
    monkeypatch.setattr(checkpoints, "run_exec", lambda *a, **k: pytest.fail("replay called target"))
    assert execute(studio, job, harness_command=command, timeout=30) == result
    root = studio.path("native", job["id"])
    (root / "oracles.json").write_text("[]\n")
    with pytest.raises(ValueError, match="changed after checkpoint"):
        execute(studio, job, harness_command=command, timeout=30)


def test_native_studio_export_does_not_claim_target_observation(tmp_path, monkeypatch):
    studio, job = _setup(tmp_path, monkeypatch)
    result = execute(studio, job, harness_command=None, timeout=30)
    assert result["status"] == "prepared"
    assert result["observed_trials"] == 0
    root = studio.path("native", job["id"])
    payload, format = studio.native_artifact(job["project"], job["id"], "ART-BOOK")
    assert format == "xlsx" and payload == next((root / "inputs").iterdir()).read_bytes()
    assert payload.startswith(b"PK")
    with pytest.raises(ValueError, match="another project"):
        studio.native_artifact("another-project", job["id"], "ART-BOOK")
    with pytest.raises(ValueError, match="unknown native artifact"):
        studio.native_artifact(job["project"], job["id"], "unknown-artifact")
    next((root / "inputs").iterdir()).unlink()
    with pytest.raises(ValueError, match="changed after checkpoint"):
        execute(studio, job, harness_command=None, timeout=30)
    with pytest.raises(ValueError, match="changed after checkpoint"):
        studio.native_artifact(job["project"], job["id"], "ART-BOOK")


def test_native_output_file_boundary_reads_bytes_and_refuses_escape(tmp_path):
    import base64

    from worldloom.studio.native import NativeReply, _submission

    out = tmp_path / "out"
    out.mkdir()
    (out / "result.xlsx").write_bytes(b"actual submitted bytes")
    reply = NativeReply.model_validate({"request_id": "one", "output_files": [
        {"artifact_id": "ART-OUT", "format": "xlsx", "path": "result.xlsx"}]})
    result = _submission(reply, out)
    assert base64.b64decode(result.files[0].content_base64) == b"actual submitted bytes"
    with pytest.raises(ValueError, match="inside the task"):
        NativeReply.model_validate({"request_id": "one", "output_files": [
            {"artifact_id": "ART-OUT", "format": "xlsx", "path": "../result.xlsx"}]})
    (out / "result.xlsx").unlink()
    (tmp_path / "secret.xlsx").write_bytes(b"outside")
    (out / "result.xlsx").symlink_to(tmp_path / "secret.xlsx")
    with pytest.raises(ValueError, match="inside the task"):
        _submission(reply, out)


def test_native_reference_refusal_precedes_any_target(tmp_path, monkeypatch):
    import worldloom.studio.checkpoints as checkpoints
    from worldloom.studio import ProjectSpec

    studio, job = _setup(tmp_path, monkeypatch)
    original = ProjectSpec.model_validate(studio.store.get(job["project"], job["revision"])["spec"])
    task = original.native_tasks[0]
    assertion = task.assertions[0].model_copy(update={"expected": "contradicts source bytes"})
    revised = original.model_copy(update={"native_tasks": (task.model_copy(update={"assertions": (assertion,)}),)})
    project = studio.store.create(revised)
    invalid = studio.store.enqueue(project["id"], project["revision"], RunOptions(operation="native"))
    monkeypatch.setattr(checkpoints, "run_exec", lambda *a, **k: pytest.fail("unqualified target called"))
    with pytest.raises(ValueError, match="reference qualification failed"):
        execute(studio, invalid, harness_command="unused-command", timeout=30)


def test_native_actual_update_file_through_studio_dispatch(tmp_path, monkeypatch):
    from worldloom.native_tasks import NativeOutput
    from worldloom.studio import ProjectSpec

    studio, job = _setup(tmp_path, monkeypatch)
    base = ProjectSpec.model_validate(studio.store.get(job["project"], job["revision"])["spec"])
    task = base.native_tasks[0]
    update = NativeTask(id="update", operation="update", use_case_id=task.use_case_id,
                        prompt="Change Evidence!B2 to Reconciled, preserving every other cell.",
                        inputs=task.inputs, output=NativeOutput(artifact_id="ART-UPDATED", format="xlsx",
                            source_artifact_id="ART-BOOK", assertions=(NativeAssertion(id="updated", target=NativeCitation(
                                artifact_id="ART-UPDATED", locator="sheet:Evidence/cell:B2"), expected="Reconciled"),)))
    project = studio.store.create(base.model_copy(update={"native_tasks": (update,)}))
    job = studio.store.enqueue(project["id"], project["revision"], RunOptions(operation="native"))
    agent = tmp_path / "update.py"
    agent.write_text('''import json, sys
from pathlib import Path
from openpyxl import load_workbook
p = json.load(sys.stdin)
w = load_workbook(p["input_files"]["ART-BOOK"])
w["Evidence"]["B2"] = "Reconciled"
w.save(Path(p["output_directory"]) / "updated.xlsx")
print(json.dumps({"request_id":p["request_id"],"output_files":[{"artifact_id":"ART-UPDATED","format":"xlsx","path":"updated.xlsx","source_sha256":p["task"]["inputs"][0]["sha256"]}]}))
''')
    argv = [sys.executable, str(agent)]
    command = subprocess.list2cmdline(argv) if os.name == "nt" else shlex.join(argv)
    result = studio.execute(job["id"], harness_command=command, timeout=30)
    assert result["passed_trials"] == result["observed_trials"] == 1
    output = next((studio.path("native", job["id"]) / "outputs").rglob("updated.xlsx"))
    assert output.is_file()
    import worldloom.studio.checkpoints as checkpoints
    monkeypatch.setattr(checkpoints, "run_exec", lambda *a, **k: pytest.fail("replay called target"))
    assert studio.execute(job["id"], harness_command=command, timeout=30) == result
    output.write_bytes(b"modified result")
    with pytest.raises(ValueError, match="changed after checkpoint"):
        studio.execute(job["id"], harness_command=command, timeout=30)
