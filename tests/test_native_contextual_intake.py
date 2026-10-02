"""Contextual labels are bound to complete canonical sections, not a cell."""
from __future__ import annotations

from dataclasses import replace
from datetime import UTC, datetime
from io import BytesIO
from typing import Literal

import pytest

from worldloom.models import (
    ArtifactIR,
    ArtifactSection,
    Authority,
    CanonicalFact,
    Cell,
    Column,
    Company,
    FormulaKind,
    Quantity,
    Row,
    Table,
)
from worldloom.native_artifacts import inspect_artifact
from worldloom.native_corpus import (
    NativeContent,
    NativeCorpusPlan,
    NativeCorpusResult,
    render_native_corpus,
)
from worldloom.native_eval_bridge import native_task_cases
from worldloom.native_query_planning import NativeWorkloadPlan, plan_native_workload
from worldloom.native_reference import qualify_native_task
from worldloom.native_tasks import NativeTask
from worldloom.world import World

NativeFormat = Literal["docx", "pptx", "xlsx"]


def _source(format: NativeFormat, *, mixed: bool = False) -> tuple[World, NativeCorpusResult, NativeWorkloadPlan]:
    facts = tuple(CanonicalFact(id="FACT-" + name, kind="financial.revenue.actual", subject="store:north",
        period=period, value=Quantity(amount=amount, unit="AUD"), valid_from=datetime(2026, 1, 1, tzinfo=UTC),
        authority=Authority.SYSTEM_OF_RECORD) for name, period, amount in (("EARLY", "2026-01", 100), ("LATER", "2026-02", 125)))
    sections = []
    for index, fact in enumerate(facts):
        columns = [Column(key="actual", label="Actual")]
        cells = {"actual": Cell(value=fact.value.amount, fact_id=fact.id)}
        if mixed and index == 1:
            columns.append(Column(key="prior", label="Prior"))
            cells["prior"] = Cell(value=100, formula=FormulaKind.REFERENCE, operands=["revenue-0:north:actual"])
        sections.append(ArtifactSection(heading="Commitment", body="Committed revenue: {{fact:" + fact.id + "}}.",
            table=Table(key=f"revenue-{index}", title="Revenue", columns=columns,
                rows=[Row(key="north", label="North store", cells=cells)])))
    world = World(seed=8128, company=Company(id="CO-1", name="Northstar Retail", industry="retail",
        headquarters="Sydney", fiscal_year_start_month=7, employees_total=100), _facts=facts,
        _artifact_irs=(ArtifactIR(id="ART-SOURCE", intent_id="INTENT-1", title="Revenue close", sections=sections),))
    plan = NativeCorpusPlan(artifact_id="review", format=format, title="Revenue close", surface="business",
        contextual_headings=True, contents=tuple(NativeContent(source_artifact_id="ART-SOURCE", section_index=index,
            placement="notes" if format == "pptx" else "body") for index in range(2)))
    result = render_native_corpus(world, plan)
    workload = NativeWorkloadPlan(use_case_id="revenue-review", objective="Review reporting evidence.",
        formats=(format,), max_tasks=32, discovery_scope="artifact")
    return world, result, workload


def _labels(result: NativeCorpusResult, replacements: dict[str, str]) -> NativeCorpusResult:
    """Modify actual labels, then recompute every transport integrity pin."""
    format = result.manifest.format
    stream = BytesIO()
    if format == "docx":
        from docx import Document
        document = Document(BytesIO(result.payload))
        for paragraph in document.paragraphs:
            if paragraph.text in replacements:
                paragraph.text = replacements[paragraph.text]
        for table in document.tables:
            for row in table.rows:
                for cell in row.cells:
                    if cell.text in replacements:
                        cell.text = replacements[cell.text]
        document.save(stream)
    elif format == "pptx":
        from pptx import Presentation
        presentation = Presentation(BytesIO(result.payload))
        for slide in presentation.slides:
            for shape in slide.shapes:
                if shape.has_text_frame and shape.text in replacements:
                    shape.text_frame.text = replacements[shape.text]
                if shape.has_table:
                    for row in shape.table.rows:
                        for cell in row.cells:
                            if cell.text in replacements:
                                cell.text = replacements[cell.text]
        presentation.save(stream)
    else:
        from openpyxl import load_workbook
        workbook = load_workbook(BytesIO(result.payload), data_only=False)
        for sheet in workbook:
            for row in sheet:
                for cell in row:
                    if isinstance(cell.value, str) and cell.value in replacements:
                        cell.value = replacements[cell.value]
        workbook.save(stream)
        workbook.close()
    payload = stream.getvalue()
    snapshot = inspect_artifact(payload, format)
    assert snapshot.sha256 != result.manifest.sha256
    return replace(result, payload=payload, manifest=result.manifest.model_copy(update={
        "sha256": snapshot.sha256, "file_size_bytes": len(payload), "native_metrics": snapshot.metrics}))


def _rebind(tasks: tuple[NativeTask, ...], result: NativeCorpusResult) -> tuple[NativeTask, ...]:
    return tuple(task.model_copy(update={"inputs": tuple(item.model_copy(update={"sha256": result.manifest.sha256})
        for item in task.inputs)}) for task in tasks)


@pytest.mark.parametrize("format", ["docx", "pptx", "xlsx"])
def test_contextual_intake_uses_the_actual_heading_and_title_for_discovery(format: NativeFormat) -> None:
    world, result, plan = _source(format)
    workload = plan_native_workload(world, {"review": result}, plan)
    assert set(workload.operation_counts) == {"read", "analyze", "update", "create"}
    assert not any(finding.code.startswith("ambiguous_") for finding in workload.findings)
    prompts = "\n".join(task.prompt for task in workload.tasks)
    for label in ("Commitment", "Revenue"):
        for period in ("2026-01", "2026-02"):
            assert label + " — reporting period " + period in prompts
    assert all(qualify_native_task(task, {"review": result.payload}).passed for task in workload.tasks)
    cases = native_task_cases(workload.tasks, {"review": result}, namespace="northstar", world=world)
    assert len(cases) == len(workload.tasks)


@pytest.mark.parametrize("format", ["docx", "pptx", "xlsx"])
@pytest.mark.parametrize("label", ["Commitment", "Revenue"])
@pytest.mark.parametrize("mutation", ["forged_period", "moved_period"])
def test_transport_valid_bytes_cannot_forge_or_move_context_between_sections(
    format: NativeFormat, label: str, mutation: str,
) -> None:
    world, original, plan = _source(format)
    tasks = plan_native_workload(world, {"review": original}, plan).tasks
    assert tasks
    earlier = label + " — reporting period 2026-01"
    later = label + " — reporting period 2026-02"
    replacements = {earlier: label + " — reporting period 2026-03"} if mutation == "forged_period" else {
        earlier: later, later: earlier}
    forged = _labels(original, replacements)
    reason = "heading is not bound" if label == "Commitment" else "selector labels disagree"
    with pytest.raises(ValueError, match=reason):
        plan_native_workload(world, {"review": forged}, plan)
    with pytest.raises(ValueError, match=reason):
        native_task_cases(_rebind(tasks, forged), {"review": forged}, namespace="northstar", world=world)


@pytest.mark.parametrize("format", ["docx", "pptx", "xlsx"])
@pytest.mark.parametrize("label", ["Commitment", "Revenue"])
def test_context_cannot_select_one_local_cell_period_from_a_mixed_dependency_section(format: NativeFormat, label: str) -> None:
    world, original, plan = _source(format, mixed=True)
    workload = plan_native_workload(world, {"review": original}, plan)
    assert workload.tasks
    units = {unit.text for unit in inspect_artifact(original.payload, format).units}
    assert label in units
    assert label + " — reporting period 2026-02" not in units
    forged = _labels(original, {label: label + " — reporting period 2026-02"})
    reason = "heading is not bound" if label == "Commitment" else "selector labels disagree"
    with pytest.raises(ValueError, match=reason):
        plan_native_workload(world, {"review": forged}, plan)
    with pytest.raises(ValueError, match=reason):
        native_task_cases(_rebind(workload.tasks, forged), {"review": forged}, namespace="northstar", world=world)
