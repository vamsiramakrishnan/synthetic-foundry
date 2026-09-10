from __future__ import annotations

import hashlib
from dataclasses import replace
from datetime import UTC, datetime
from io import BytesIO
from typing import Literal
from zipfile import ZipFile

import pytest

from worldloom.models import (
    ArtifactIR,
    ArtifactSection,
    Authority,
    CanonicalFact,
    Company,
    Quantity,
)
from worldloom.native_artifacts import inspect_artifact
from worldloom.native_corpus import (
    NativeContent,
    NativeCorpusPlan,
    plan_native_corpus,
    render_native_corpus,
)
from worldloom.world import World


def _world(count: int = 200) -> World:
    facts = tuple(CanonicalFact(
        id=f"FACT-{index:04}", kind="retail.receipt.exception", subject=f"purchase-order:{index}",
        text_value=f"Supplier receipt for purchase order {index} requires receiving-team reconciliation.",
        valid_from=datetime(2026, 1, 1, tzinfo=UTC), authority=Authority.SYSTEM_OF_RECORD,
    ) for index in range(count))
    sections = [ArtifactSection(heading="Supplier reconciliation", body=f"Receiving disposition: {{{{fact:{fact.id}}}}}", fact_ids=[fact.id]) for fact in facts]
    return World(company=Company(id="CO-1", name="Northstar Retail", industry="retail", headquarters="Sydney", fiscal_year_start_month=7, employees_total=100), _facts=facts,
                 _artifact_irs=(ArtifactIR(id="ART-1", intent_id="INTENT-1", title="Receiving register", sections=sections),))


@pytest.mark.parametrize("format", ["docx", "pptx", "xlsx"])
def test_large_native_corpus_grounded_and_replayable(format: Literal["docx", "pptx", "xlsx"]) -> None:
    world = _world()
    plan = plan_native_corpus(world, artifact_id="ART-CORPUS", format=format, title="Receiving review", minimum_units=200, minimum_distinct_facts=200)
    result = render_native_corpus(world, plan)
    assert result.payload == render_native_corpus(world, plan).payload
    assert result.manifest.content_units == 200
    assert result.manifest.distinct_fact_count == 200
    assert len({entry.text_sha256 for entry in result.manifest.evidence[:200]}) == 200
    assert result.manifest.evidence[-1].fact_ids == ("FACT-0199",)
    assert not result.manifest.ballast_is_evidence
    extracted = {unit.locator: unit.text for unit in inspect_artifact(result.payload, format).units}
    for entry in result.manifest.evidence:
        assert entry.text_sha256 == hashlib.sha256(extracted[entry.locator].encode()).hexdigest()
    if format == "docx":
        from docx import Document
        document = Document(BytesIO(result.payload))
        assert "purchase order 199" in document.paragraphs[-1].text
        with ZipFile(BytesIO(result.payload)) as package:
            assert package.read("word/document.xml").count(b'w:type="page"') == 199
        assert result.manifest.explicit_page_floor == 200
    elif format == "pptx":
        from pptx import Presentation
        deck = Presentation(BytesIO(result.payload))
        assert len(deck.slides) == 200
        assert "purchase order 199" in deck.slides[-1].shapes[-1].text
    else:
        from openpyxl import load_workbook
        workbook = load_workbook(BytesIO(result.payload))
        assert workbook["Evidence"].max_row == 201
        assert "purchase order 199" in workbook["Evidence"]["B201"].value


def test_native_corpus_requires_real_new_evidence() -> None:
    world = _world(1)
    with pytest.raises(ValueError, match="insufficient grounded content"):
        plan_native_corpus(world, artifact_id="ART-X", format="docx", title="Receiving", minimum_units=200)
    content = NativeContent(source_artifact_id="ART-1", section_index=0)
    plan = NativeCorpusPlan(artifact_id="ART-X", format="docx", title="Receiving", contents=(content, content), minimum_units=2)
    with pytest.raises(ValueError, match="duplicate grounded content"):
        render_native_corpus(world, plan)
    bad = world._artifact_irs[0].model_copy(update={"sections": [ArtifactSection(heading="Missing", body="{{fact:FACT-MISSING}}", fact_ids=["FACT-MISSING"])]})
    with pytest.raises(ValueError, match="unknown canonical facts"):
        render_native_corpus(replace(world, _artifact_irs=(bad,)), plan.model_copy(update={"contents": (content,), "minimum_units": 1}))


def test_native_corpus_notes_and_literal_spreadsheet_text() -> None:
    world = _world(1)
    content = NativeContent(source_artifact_id="ART-1", section_index=0, placement="notes")
    plan = NativeCorpusPlan(artifact_id="ART-X", format="pptx", title="Receiving", contents=(content,))
    from pptx import Presentation
    result = render_native_corpus(world, plan)
    deck = Presentation(BytesIO(result.payload))
    assert "purchase order 0" in deck.slides[0].notes_slide.notes_text_frame.text
    assert result.manifest.evidence[0].locator == "slide:1/notes"
    with pytest.raises(ValueError, match="speaker notes require pptx"):
        render_native_corpus(world, plan.model_copy(update={"format": "xlsx"}))
    section = world._artifact_irs[0].sections[0].model_copy(update={"body": "={{fact:FACT-0000}}"})
    ir = world._artifact_irs[0].model_copy(update={"sections": [section]})
    world = replace(world, _artifact_irs=(ir,))
    result = render_native_corpus(world, plan.model_copy(update={"format": "xlsx", "contents": (content.model_copy(update={"placement": "body"}),)}))
    from openpyxl import load_workbook
    assert load_workbook(BytesIO(result.payload))["Evidence"]["B2"].data_type == "s"


def test_paraphrases_do_not_inflate_fact_coverage() -> None:
    world = _world(1)
    original = world._artifact_irs[0]
    extra = original.sections[0].model_copy(update={"body": "Disposition review: {{fact:FACT-0000}}"})
    world = replace(world, _artifact_irs=(original.model_copy(update={"sections": [*original.sections, extra]}),))
    with pytest.raises(ValueError, match="insufficient distinct canonical facts"):
        plan_native_corpus(world, artifact_id="ART-X", format="docx", title="Receiving", minimum_units=2, minimum_distinct_facts=2)


def test_workbook_canonical_numeric_cells_support_analysis() -> None:
    world = _world(2)
    first = world._facts[0].model_copy(update={"value": Quantity(amount=125.5, unit="AUD"), "text_value": None})
    second = world._facts[1].model_copy(update={"value": Quantity(amount=50, unit="AUD"), "text_value": None})
    world = replace(world, _facts=(first, second))
    plan = plan_native_corpus(world, artifact_id="ART-X", format="xlsx", title="Invoice reconciliation", minimum_units=2)
    result = render_native_corpus(world, plan)
    from openpyxl import load_workbook
    workbook = load_workbook(BytesIO(result.payload))
    assert workbook["Facts"]["B2"].value + workbook["Facts"]["B3"].value == 175.5
    assert workbook["Facts"]["C2"].value == "AUD"
    numeric = [item for item in result.manifest.evidence if item.value is not None]
    assert [(item.locator, item.fact_ids, item.value) for item in numeric] == [
        ("sheet:Facts/cell:B2", ("FACT-0000",), 125.5),
        ("sheet:Facts/cell:B3", ("FACT-0001",), 50.0),
    ]
    extracted = {unit.locator: unit.text for unit in inspect_artifact(result.payload, "xlsx").units}
    for entry in result.manifest.evidence:
        assert entry.text_sha256 == hashlib.sha256(extracted[entry.locator].encode()).hexdigest()
