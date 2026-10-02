"""Scripted pilot worker over public instructions and staged DOCX bytes.

This deliberately supports one bounded task family, not a general agent. The
worker never imports benchmark plans, reference answers or native graders.
"""
from __future__ import annotations

import base64
import hashlib
import json
import re
import sys
from decimal import Decimal
from io import BytesIO
from pathlib import Path

from docx import Document

sys.path.insert(0, str(Path(__file__).resolve().parents[2] / "src"))
from worldloom.render.ooxml import normalise


def solve(request, inputs):
    assert len(request['inputs']) == 1 and request['inputs'][0]['format'] == 'docx'
    item = request['inputs'][0]
    artifact = item['artifact_id']
    data = inputs[artifact]
    assert hashlib.sha256(data).hexdigest() == item['sha256']
    assert 'assertions' not in request and 'expected' not in request
    document = Document(BytesIO(data))
    prompt = request['prompt']

    def paragraph(heading):
        matches = [(i, p) for i, p in enumerate(document.paragraphs[:-1]) if p.text == heading]
        assert len(matches) == 1
        index = matches[0][0] + 1
        return document.paragraphs[index].text, f'paragraph:{index + 1}'

    def table_value(selector):
        title, row_label, column_label = selector
        tables = [(i, t) for i, t in enumerate(document.tables) if t.cell(0, 0).text == title]
        assert len(tables) == 1
        index, table = tables[0]
        row = next(i for i, r in enumerate(table.rows) if r.cells[0].text == row_label)
        column = next(i for i, c in enumerate(table.rows[0].cells) if c.text == column_label)
        return table.cell(row, column).text, f'table:{index + 1}/row:{row + 1}/cell:{column + 1}'

    def answer(identifier, value, locators):
        return {'assertion_id': identifier, 'value': value, 'citations': [
            {'artifact_id': artifact, 'locator': locator} for locator in locators]}

    headings = re.findall(r"section headed '([^']+)'", prompt)
    if request['operation'] == 'analyze':
        constituents = prompt.split('Locate these constituent values: ', 1)[1]
        selectors = re.findall(r"table '([^']+)', row '([^']+)', column '([^']+)'", constituents)
        found = [table_value(selector) for selector in selectors]
        answers = [answer('total', str(sum((Decimal(value) for value, _ in found), Decimal(0))),
            [locator for _, locator in found])]
    else:
        found = [paragraph(heading) for heading in headings]
        answers = [answer(shape['assertion_id'], value, [locator])
            for shape, (value, locator) in zip(request['answers'], found, strict=True)]
    files = []
    output = request['output']
    if output:
        assert output['format'] == 'docx'
        source_hash = None
        if request['operation'] == 'update':
            index = int(found[0][1].split(':')[1]) - 1
            value = found[0][0] + '\nRelated evidence: ' + found[1][0]
            target = document.paragraphs[index]
            if target.runs:
                target.runs[0].text = value
                for run in target.runs[1:]:
                    run.text = ''
            else:
                target.add_run(value)
            source_hash = item['sha256']
        elif request['operation'] == 'create':
            document = Document()
            for heading, (value, _) in zip(headings, found, strict=True):
                document.add_paragraph(heading)
                document.add_paragraph(value)
        stream = BytesIO()
        document.save(stream)
        content = normalise(stream.getvalue(), created="1980-01-01T00:00:00Z")
        files = [{'artifact_id': output['artifact_id'], 'format': 'docx',
            'content_base64': base64.b64encode(content).decode('ascii'), 'source_sha256': source_hash}]
    return {'answers': answers, 'files': files}


payload = json.load(sys.stdin)
assert payload['schema'] == 'worldloom.native-harness-request/v1'
request = payload['task']
root = Path(payload['input_root'])
assert root == Path.cwd()
assert not list(root.rglob('manifest.json')) and not list(root.rglob('oracle.json'))
inputs = {item['artifact_id']: (root / item['path']).read_bytes() for item in request['inputs']}
submission = solve(request, inputs)
print(json.dumps({'schema': 'worldloom.native-harness-response/v1', 'task_id': request['id'],
    'execution_id': request['execution_id'], 'submission': submission}))
