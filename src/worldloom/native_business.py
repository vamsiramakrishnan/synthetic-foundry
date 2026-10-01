"""Business native surfaces, with oracle identities confined to the manifest.

The artifact IR owns tables and computations. This renderer carries them into
Office rather than replacing them with a global fact ledger. Every dependency
must exist in the selected source, every number must be grounded, and native
locators are verified against the actual bytes before provenance is returned.
"""
from __future__ import annotations

import hashlib
import math
import re
from collections.abc import Iterable
from dataclasses import dataclass, replace
from io import BytesIO
from typing import TYPE_CHECKING, Any, Literal

import networkx as nx

from .models import Cell, FormulaKind
from .narrative import references
from .native_artifacts import inspect_artifact
from .native_corpus import (
    NativeContent,
    NativeContentExclusion,
    NativeContentProvenance,
    NativeCorpusManifest,
    NativeCorpusPlan,
    NativeCorpusResult,
    _Content,
)
from .presentation import of as presentation_of
from .render.ooxml import normalise
from .render.values import corpus_locale

if TYPE_CHECKING:
    from .models import CanonicalFact, Table
    from .world import World

_Address = tuple[str, str, str, str]
_TABLE_CELL_BUDGET = 500_000
_CANONICAL_ID = re.compile(r"\b[A-Z][A-Z0-9]*-[A-Z0-9-]+\b")


@dataclass(frozen=True)
class _Tables:
    cells: dict[_Address, Cell]
    owners: dict[_Address, _Content]
    dependencies: dict[_Address, tuple[_Address, ...]]
    facts: dict[_Address, tuple[str, ...]]
    units: dict[_Address, str | None]


@dataclass(frozen=True)
class _Pending:
    content: _Content
    locator: str
    fact_ids: tuple[str, ...]
    expected_text: str
    value: float | str | None = None
    unit: str | None = None
    kind: Literal["prose", "value", "formula"] = "prose"
    address: _Address | None = None
    metadata: dict[str, str] | None = None


def business_label(value: str) -> str:
    """A stable display label; canonical measure vocabulary remains private."""
    return " ".join(value.replace(".", " ").replace("_", " ").split()).title()


def _reject_identity_leaks(facts: dict[str, CanonicalFact], texts: Iterable[str]) -> None:
    leaked = sorted({match[0] for text in texts for match in _CANONICAL_ID.finditer(text) if match[0] in facts})
    if leaked:
        raise ValueError(f"canonical fact identities leaked into business surface: {leaked[:5]}")


def _same_value(first: float | str | None, second: float | str | None) -> bool:
    if isinstance(first, (int, float)) and isinstance(second, (int, float)):
        return math.isfinite(first) and math.isfinite(second) and math.isclose(first, second, rel_tol=1e-9, abs_tol=1e-6)
    return first == second


def _content_identity(body: str, table: Table | None) -> str:
    identity = " ".join(body.split()).casefold()
    if table is not None:
        identity += repr(tuple(tuple((cell.value, cell.fact_id, cell.formula)
            if (cell := row.cells.get(column.key)) is not None else None
            for column in table.columns) for row in table.rows))
    return identity


def plan_business_content(world: World, *, source_artifact_ids: tuple[str, ...] | None = None
                          ) -> tuple[tuple[NativeContent, ...], tuple[NativeContentExclusion, ...]]:
    """Select public sections with complete, valid dependency closures.

    Auto-selection can exclude an unusable candidate. Explicit plans still
    refuse it. The audit keeps source defects visible after selecting a clean
    subset, without changing any value or computation to make it usable.
    """
    requested = set(source_artifact_ids or ())
    selected: dict[tuple[str, int], _Content] = {}
    seen: set[str] = set()
    excluded: list[NativeContentExclusion] = []
    for ir in sorted(world.artifact_irs, key=lambda artifact: artifact.id):
        if source_artifact_ids is not None and ir.id not in requested:
            continue
        tables: dict[str, list[int]] = {}
        for index, section in enumerate(ir.sections):
            if section.table is not None:
                tables.setdefault(section.table.key, []).append(index)
        for index, section in enumerate(ir.sections):
            if (ir.id, index) in selected:
                continue
            code: Literal["private_native_section", "unwritten_native_section", "duplicate_grounded_content", "invalid_native_dependency", "invalid_native_evidence"] = "invalid_native_evidence"
            if section.hidden:
                code, reason = "private_native_section", "Source section is private."
            elif section.table is None and not references.referenced(section.body or ""):
                code, reason = "unwritten_native_section", "No grounded authored prose or resolved table."
            else:
                closure: set[int] = set()
                remaining = [index]
                try:
                    while remaining:
                        position = remaining.pop()
                        if position in closure:
                            continue
                        dependency = ir.sections[position]
                        if dependency.hidden:
                            code = "invalid_native_dependency"
                            raise ValueError("native table depends on a private source section")
                        closure.add(position)
                        table = dependency.table
                        if table is None:
                            continue
                        if len(tables.get(table.key, ())) != 1:
                            code = "invalid_native_dependency"
                            raise ValueError("native table dependency key is ambiguous")
                        for row in table.rows:
                            for cell in row.cells.values():
                                if cell.formula not in (FormulaKind.REFERENCE, FormulaKind.SUM):
                                    continue
                                for operand in cell.operands:
                                    parts = operand.split(":")
                                    if len(parts) != 3:
                                        continue
                                    targets = tables.get(parts[0], ())
                                    if len(targets) != 1:
                                        code = "invalid_native_dependency"
                                        raise ValueError("native table names an unknown or ambiguous dependency")
                                    remaining.append(targets[0])
                    candidate = NativeCorpusPlan(artifact_id="native-candidate", format="xlsx", title=ir.title,
                        surface="business", contents=tuple(NativeContent(source_artifact_id=ir.id, section_index=p) for p in sorted(closure)))
                    prepared = prepare_business_content(world, candidate)
                    additions = [content for content in prepared if (ir.id, content.source.section_index) not in selected]
                    if any(_content_identity(content.body, content.table) in seen for content in additions):
                        code = "duplicate_grounded_content"
                        raise ValueError("headings, renaming or repetition do not increase grounded coverage")
                except ValueError as error:
                    reason = str(error)
                    if "duplicate grounded content" in reason:
                        code = "duplicate_grounded_content"
                else:
                    for content in additions:
                        selected[ir.id, content.source.section_index] = content
                        seen.add(_content_identity(content.body, content.table))
                    continue
            excluded.append(NativeContentExclusion(source_artifact_id=ir.id, section_index=index, code=code, reason=reason))
    return (tuple(selected[key].source for key in sorted(selected)),
            tuple(sorted(excluded, key=lambda entry: (entry.source_artifact_id, entry.section_index))))


def _table_graph(world: World, contents: tuple[_Content, ...]) -> _Tables:
    facts = {fact.id: fact for fact in world.facts}
    cells: dict[_Address, Cell] = {}
    owners: dict[_Address, _Content] = {}
    selected: set[tuple[str, str]] = set()
    for content in contents:
        table = content.table
        if table is None:
            continue
        key = (content.source.source_artifact_id, table.key)
        if key in selected:
            raise ValueError(f"duplicate selected native table: {key}")
        selected.add(key)
        columns = [column.key for column in table.columns]
        rows = [row.key for row in table.rows]
        if len(set(columns)) != len(columns) or len(set(rows)) != len(rows):
            raise ValueError("native table has duplicate row or column keys")
        if len(rows) > 1_048_573 or len(columns) > 16_383:
            raise ValueError("native table exceeds Excel dimensions")
        for row in table.rows:
            if set(row.cells) - set(columns):
                raise ValueError("native table cell names an undeclared column")
            for column in table.columns:
                cell = row.cells.get(column.key)
                if cell is None:
                    continue
                address = (*key, row.key, column.key)
                cells[address], owners[address] = cell, content
                if len(cells) > _TABLE_CELL_BUDGET:
                    raise ValueError("native table exceeds bounded cell budget")
                if cell.fact_id is not None:
                    fact = facts.get(cell.fact_id)
                    if fact is None:
                        raise ValueError(f"unknown canonical facts: {cell.fact_id}")
                    value = fact.value.amount if fact.value is not None else fact.text_value
                    if cell.value != value:
                        raise ValueError(f"native table value disagrees with canonical fact: {cell.fact_id}")
                elif isinstance(cell.value, (int, float)) and cell.formula is None:
                    raise ValueError("native table number lacks a canonical fact or declared computation")
                if isinstance(cell.value, (int, float)) and not math.isfinite(cell.value):
                    raise ValueError("native table value must be finite")
    dependencies: dict[_Address, tuple[_Address, ...]] = {}
    graph = nx.DiGraph()
    graph.add_nodes_from(sorted(cells))
    for address, cell in sorted(cells.items()):
        source_key, table_key, row_key, column_key = address
        operands: list[_Address] = []
        if cell.formula is FormulaKind.SUM:
            for specification in cell.operands:
                parts = specification.split(":")
                if len(parts) == 1:
                    operands.append((source_key, table_key, specification, column_key))
                elif len(parts) == 3:
                    operands.append((source_key, parts[0], parts[1], parts[2]))
                else:
                    raise ValueError("native SUM operand must name a row or table:row:column")
        elif cell.formula in (FormulaKind.DIFFERENCE, FormulaKind.RATIO_PCT):
            if len(cell.operands) != 2:
                raise ValueError("native arithmetic formula requires exactly two operands")
            operands = [(source_key, table_key, row_key, specification) for specification in cell.operands]
        elif cell.formula is FormulaKind.REFERENCE:
            if len(cell.operands) != 1 or len(parts := cell.operands[0].split(":")) != 3:
                raise ValueError("native reference requires exactly one table:row:column operand")
            operands = [(source_key, parts[0], parts[1], parts[2])]
        for dependency in operands:
            if dependency not in cells:
                raise ValueError(f"native table dependency is not selected or has no value: {dependency}")
            graph.add_edge(dependency, address)
        dependencies[address] = tuple(operands)
    try:
        order = tuple(nx.lexicographical_topological_sort(graph))
    except nx.NetworkXUnfeasible as error:
        raise ValueError("native table formula dependency cycle") from error
    bound: dict[_Address, tuple[str, ...]] = {}
    units: dict[_Address, str | None] = {}
    for address in order:
        cell = cells[address]
        ids = {cell.fact_id} if cell.fact_id else set()
        dependency_addresses = dependencies[address]
        for dependency in dependency_addresses:
            ids.update(bound[dependency])
        bound[address] = tuple(sorted(ids))
        operand_units = {units[dependency] for dependency in dependency_addresses}
        if cell.formula in (FormulaKind.SUM, FormulaKind.DIFFERENCE, FormulaKind.RATIO_PCT) and len(operand_units) != 1:
            raise ValueError("native arithmetic formula operands use different units")
        units[address] = (facts[cell.fact_id].value.unit if cell.fact_id and facts[cell.fact_id].value is not None
                          else ("percent" if cell.formula is FormulaKind.RATIO_PCT
                                else sorted(operand_units, key=lambda unit: unit or "")[0] if len(operand_units) == 1 else None))
        if cell.formula is None:
            continue
        values = [cells[dependency].value for dependency in dependency_addresses]
        if cell.formula is FormulaKind.REFERENCE:
            expected = values[0]
        else:
            if any(not isinstance(value, (int, float)) for value in values):
                raise ValueError("native arithmetic formula operand must be numeric")
            numbers = [float(value) for value in values if isinstance(value, (int, float))]
            if cell.formula is FormulaKind.SUM:
                expected = math.fsum(numbers)
            elif cell.formula is FormulaKind.DIFFERENCE:
                expected = numbers[0] - numbers[1]
            else:
                expected = numbers[0] / numbers[1] * 100 if numbers[1] else 0.0
        if not _same_value(cell.value, expected):
            raise ValueError(f"native table formula literal disagrees with its operands: {address}")
        if not bound[address]:
            raise ValueError("native table computation has no canonical evidence")
    return _Tables(cells, owners, dependencies, bound, units)


def prepare_business_content(world: World, plan: NativeCorpusPlan) -> tuple[_Content, ...]:
    """Accept real prose or resolved tables, without manufacturing evidence."""
    if len(plan.contents) < plan.minimum_units:
        raise ValueError(f"insufficient grounded content: need {plan.minimum_units}, have {len(plan.contents)}")
    if len(plan.contents) > 10_000:
        raise ValueError("native corpus exceeds bounded content budget")
    artifacts = {ir.id: ir for ir in world.artifact_irs}
    facts = {fact.id: fact for fact in world.facts}
    seen: set[str] = set()
    contents = []
    for source in plan.contents:
        ir = artifacts.get(source.source_artifact_id)
        if ir is None or source.section_index >= len(ir.sections):
            raise ValueError(f"unknown authored section: {source.source_artifact_id}:{source.section_index}")
        section = ir.sections[source.section_index]
        if section.hidden:
            # Explicit selection must preserve the same privacy boundary as
            # auto-selection, including tables and presenter notes.
            raise ValueError(f"private authored section: {source.source_artifact_id}:{source.section_index}")
        if source.placement == "notes" and plan.format != "pptx":
            raise ValueError("speaker notes require pptx")
        body = section.body or ""
        ids = set(references.referenced(body))
        carried = set(section.fact_ids) | ids
        if section.table:
            carried.update(cell.fact_id for row in section.table.rows for cell in row.cells.values() if cell.fact_id)
            ids.update(cell.fact_id for row in section.table.rows for cell in row.cells.values() if cell.fact_id)
        unknown = sorted((carried - facts.keys()) | set(references.unresolved(body, facts)))
        if unknown:
            raise ValueError(f"unknown canonical facts: {unknown}")
        if references.bare_numbers(body):
            raise ValueError("native corpus prose must reference canonical figures")
        if body.strip() and not references.referenced(body):
            raise ValueError("native corpus evidence must occur in authored prose")
        if not ids and section.table is None:
            raise ValueError("native corpus section has no canonical evidence")
        body = references.substitute(body, facts, locale=corpus_locale(world), presentation=presentation_of(world))
        visible = [section.heading, body]
        if section.table is not None:
            visible.extend([section.table.title, *(column.label for column in section.table.columns)])
            for row in section.table.rows:
                visible.append(row.label)
                visible.extend(cell.value for cell in row.cells.values() if isinstance(cell.value, str))
        _reject_identity_leaks(facts, visible)
        # A second rendering of one paragraph cannot purchase another evidence
        # unit. A genuinely different table in the same section can.
        identity = " ".join(body.split()).casefold()
        if section.table is not None:
            # Row/column keys and titles are addressing metadata. Renaming
            # them around the same values must not buy another evidence unit.
            identity += repr(tuple(tuple((cell.value, cell.fact_id, cell.formula)
                if (cell := row.cells.get(column.key)) is not None else None
                for column in section.table.columns) for row in section.table.rows))
        if identity in seen:
            raise ValueError("duplicate grounded content: headings or renaming do not increase coverage")
        seen.add(identity)
        contents.append(_Content(source, section.heading, body, tuple(sorted(ids)), section.table))
    result = tuple(contents)
    tables = _table_graph(world, result)
    table_facts: dict[tuple[str, int], set[str]] = {}
    for address, owner in sorted(tables.owners.items()):
        table_facts.setdefault((owner.source.source_artifact_id, owner.source.section_index), set()).update(tables.facts[address])
    result = tuple(replace(content, fact_ids=tuple(sorted(set(content.fact_ids) |
        table_facts.get((content.source.source_artifact_id, content.source.section_index), set())))) for content in result)
    if any(not content.fact_ids for content in result):
        raise ValueError("native corpus section has no canonical evidence")
    if len({fid for content in result for fid in content.fact_ids}) < plan.minimum_distinct_facts:
        raise ValueError("insufficient distinct canonical facts for native business corpus")
    return result


def _text(value: float | str | None) -> str:
    if value is None:
        return ""
    if isinstance(value, (int, float)) and float(value).is_integer():
        return str(int(value))
    return str(value)


def _register_row(fact: CanonicalFact, names: dict[str, str]) -> tuple[str | float, ...]:
    return (business_label(fact.kind), names.get(fact.subject, fact.subject), fact.period or "",
            fact.value.amount if fact.value is not None else fact.text_value or "",
            fact.value.unit if fact.value is not None else "text", fact.valid_from.isoformat())


_REGISTER_HEADERS = ("Measure", "Subject", "Reporting period", "Value", "Unit", "Effective from")


def _table_evidence(tables: _Tables, address: _Address, locator: str, text: str,
                    metadata: dict[str, str]) -> _Pending | None:
    cell = tables.cells[address]
    if not tables.facts[address]:
        return None
    return _Pending(tables.owners[address], locator, tables.facts[address], text, cell.value,
                    tables.units[address], "formula" if cell.formula else "value", address, metadata)


def _docx(world: World, plan: NativeCorpusPlan, contents: tuple[_Content, ...], tables: _Tables,
          stream: BytesIO, pending: list[_Pending], addresses: dict[_Address, str]) -> None:
    from docx import Document
    document = Document()
    document.core_properties.title = plan.title
    document.sections[0].header.paragraphs[0].text = world.company.name
    document.sections[0].footer.paragraphs[0].text = plan.title
    facts, names = {f.id: f for f in world.facts}, world.entity_names()
    for index, content in enumerate(contents):
        if index:
            document.add_page_break()
        document.add_heading(content.heading, level=1)
        if content.body:
            document.add_paragraph(content.body)
            pending.append(_Pending(content, f"paragraph:{len(document.paragraphs)}",
                           tuple(sorted(references.referenced(world.artifact_irs.by_id(content.source.source_artifact_id).sections[content.source.section_index].body or ""))), content.body))
        if content.table is not None:
            source_table = content.table
            table = document.add_table(rows=1, cols=len(source_table.columns) + 1)
            table.style = "Table Grid"
            table.rows[0].cells[0].text = source_table.title
            for c, column in enumerate(source_table.columns, 1):
                table.rows[0].cells[c].text = column.label
            t = len(document.tables)
            for r, row in enumerate(source_table.rows, 2):
                cells = table.add_row().cells
                cells[0].text = row.label
                for c, column in enumerate(source_table.columns, 2):
                    address = (content.source.source_artifact_id, source_table.key, row.key, column.key)
                    cell = tables.cells.get(address)
                    if cell is None:
                        continue
                    text = _text(cell.value)
                    cells[c - 1].text = text
                    locator = f"table:{t}/row:{r}/cell:{c}"
                    addresses[address] = locator
                    entry = _table_evidence(tables, address, locator, text,
                        {"table_title": f"table:{t}/row:1/cell:1", "row_label": f"table:{t}/row:{r}/cell:1",
                         "column_label": f"table:{t}/row:1/cell:{c}"})
                    if entry is not None:
                        pending.append(entry)
        else:
            table = document.add_table(rows=1, cols=len(_REGISTER_HEADERS))
            table.style = "Table Grid"
            for cell, header in zip(table.rows[0].cells, _REGISTER_HEADERS, strict=True):
                cell.text = header
            t = len(document.tables)
            for r, fact_id in enumerate(content.fact_ids, 2):
                fact = facts[fact_id]
                values = _register_row(fact, names)
                for cell, value in zip(table.add_row().cells, values, strict=True):
                    cell.text = _text(value)
                locator = f"table:{t}/row:{r}/cell:4"
                pending.append(_Pending(content, locator, (fact_id,), _text(values[3]), values[3],
                    fact.value.unit if fact.value else None, "value", metadata={key: f"table:{t}/row:{r}/cell:{c}"
                    for key, c in (("measure", 1), ("subject", 2), ("period", 3), ("unit", 5))}))
    document.save(stream)


def _slide(presentation: Any, heading: str) -> Any:
    from pptx.util import Inches, Pt
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    box = slide.shapes.add_textbox(Inches(.6), Inches(.4), Inches(12), Inches(.8))
    box.text_frame.text = heading
    box.text_frame.paragraphs[0].font.size = Pt(24)
    return slide


def _pptx(world: World, plan: NativeCorpusPlan, contents: tuple[_Content, ...], tables: _Tables,
          stream: BytesIO, pending: list[_Pending], addresses: dict[_Address, str]) -> None:
    from pptx import Presentation
    from pptx.util import Inches, Pt
    presentation = Presentation()
    presentation.slide_width, presentation.slide_height = Inches(13.333), Inches(7.5)
    presentation.core_properties.title = plan.title
    facts, names = {f.id: f for f in world.facts}, world.entity_names()
    for content in contents:
        slide = _slide(presentation, content.heading)
        number = len(presentation.slides)
        if content.body:
            if content.source.placement == "notes":
                slide.notes_slide.notes_text_frame.text = content.body
                locator = f"slide:{number}/notes"
            else:
                if len(content.body) > 2400:
                    raise ValueError("slide body exceeds readable budget; split authored sections or use notes")
                box = slide.shapes.add_textbox(Inches(.6), Inches(1.4), Inches(12), Inches(5.4))
                box.text_frame.text = content.body
                box.text_frame.word_wrap = True
                for paragraph in box.text_frame.paragraphs:
                    paragraph.font.size = Pt(18)
                locator = f"slide:{number}/shape:2/text"
            body_ids = references.referenced(world.artifact_irs.by_id(content.source.source_artifact_id).sections[content.source.section_index].body or "")
            pending.append(_Pending(content, locator, tuple(sorted(body_ids)), content.body))
        if content.table is None:
            values = [_register_row(facts[fid], names) for fid in content.fact_ids]
            for offset in range(0, len(values), 8):
                slide = _slide(presentation, content.heading + " — source records")
                prefix = f"slide:{len(presentation.slides)}/shape:2"
                table = slide.shapes.add_table(min(8, len(values) - offset) + 1, len(_REGISTER_HEADERS),
                    Inches(.6), Inches(1.5), Inches(12), Inches(5.2)).table
                for c, header in enumerate(_REGISTER_HEADERS):
                    table.cell(0, c).text = header
                for r, (fact_id, record_values) in enumerate(zip(content.fact_ids[offset:offset + 8], values[offset:offset + 8], strict=True), 2):
                    for c, value in enumerate(record_values):
                        table.cell(r - 1, c).text = _text(value)
                    fact = facts[fact_id]
                    pending.append(_Pending(content, f"{prefix}/row:{r}/cell:4", (fact_id,), _text(record_values[3]), record_values[3],
                        fact.value.unit if fact.value else None, "value", metadata={key: f"{prefix}/row:{r}/cell:{c}"
                        for key, c in (("measure", 1), ("subject", 2), ("period", 3), ("unit", 5))}))
        else:
            source_table = content.table
            for row_offset in range(0, len(source_table.rows), 8):
                for column_offset in range(0, len(source_table.columns), 5):
                    rows = source_table.rows[row_offset:row_offset + 8]
                    columns = source_table.columns[column_offset:column_offset + 5]
                    slide = _slide(presentation, source_table.title)
                    prefix = f"slide:{len(presentation.slides)}/shape:2"
                    table = slide.shapes.add_table(len(rows) + 1, len(columns) + 1, Inches(.6), Inches(1.5), Inches(12), Inches(5.2)).table
                    table.cell(0, 0).text = source_table.title
                    for c, column in enumerate(columns, 1):
                        table.cell(0, c).text = column.label
                    for r, row in enumerate(rows, 2):
                        table.cell(r - 1, 0).text = row.label
                        for c, column in enumerate(columns, 2):
                            address = (content.source.source_artifact_id, source_table.key, row.key, column.key)
                            cell = tables.cells.get(address)
                            if cell is None:
                                continue
                            text = _text(cell.value)
                            table.cell(r - 1, c - 1).text = text
                            locator = f"{prefix}/row:{r}/cell:{c}"
                            addresses[address] = locator
                            entry = _table_evidence(tables, address, locator, text,
                                {"table_title": f"{prefix}/row:1/cell:1", "row_label": f"{prefix}/row:{r}/cell:1",
                                 "column_label": f"{prefix}/row:1/cell:{c}"})
                            if entry is not None:
                                pending.append(entry)
        if len(presentation.slides) > 10_000:
            raise ValueError("native deck exceeds bounded slide budget")
    presentation.save(stream)


def _sheet_name(title: str, used: set[str]) -> str:
    stem = re.sub(r"[\\/*?:\[\]']", " ", title).strip()[:31] or "Business records"
    name, suffix = stem, 1
    while name.casefold() in used:
        suffix += 1
        ending = f" ({suffix})"
        name = stem[:31 - len(ending)] + ending
    used.add(name.casefold())
    return name


def _excel_formula(address: _Address, tables: _Tables, addresses: dict[_Address, str]) -> str:
    from openpyxl.utils import quote_sheetname
    def cell(operand: _Address) -> str:
        locator = addresses[operand]
        sheet, coordinate = locator.removeprefix("sheet:").split("/cell:")
        return f"{quote_sheetname(sheet)}!{coordinate}"
    values = [cell(operand) for operand in tables.dependencies[address]]
    kind = tables.cells[address].formula
    if kind is FormulaKind.SUM:
        return "=SUM(" + ",".join(values) + ")"
    if kind is FormulaKind.DIFFERENCE:
        return f"={values[0]}-{values[1]}"
    if kind is FormulaKind.RATIO_PCT:
        return f"=IF({values[1]}=0,0,{values[0]}/{values[1]}*100)"
    if kind is FormulaKind.REFERENCE:
        return "=" + values[0]
    raise ValueError("native cell has no declared formula")


def _write_cell(sheet: Any, row: int, column: int, value: float | str | None) -> Any:
    if isinstance(value, str) and len(value) > 32767:
        raise ValueError("native value exceeds Excel cell limit; split the source section")
    cell = sheet.cell(row, column, value)
    if isinstance(value, str):
        # Source text may start with '='. A source sentence is never executable.
        cell.data_type = "s"
    return cell


def _xlsx(world: World, plan: NativeCorpusPlan, contents: tuple[_Content, ...], tables: _Tables,
          stream: BytesIO, pending: list[_Pending], addresses: dict[_Address, str]) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter
    workbook = Workbook()
    workbook.properties.title = plan.title
    workbook.properties.creator = world.company.name
    briefing = workbook.active
    briefing.title = "Briefing"
    briefing.append(["Section", "Evidence"])
    for content in contents:
        if not content.body:
            continue
        row = briefing.max_row + 1
        _write_cell(briefing, row, 1, content.heading)
        _write_cell(briefing, row, 2, content.body)
        body_ids = references.referenced(world.artifact_irs.by_id(content.source.source_artifact_id).sections[content.source.section_index].body or "")
        pending.append(_Pending(content, f"sheet:Briefing/cell:B{row}", tuple(sorted(body_ids)), content.body))
    names = {"briefing"}
    sheets: dict[tuple[str, str], Any] = {}
    for content in contents:
        table = content.table
        if table is None:
            continue
        sheet = workbook.create_sheet(_sheet_name(table.title, names))
        sheets[content.source.source_artifact_id, table.key] = sheet
        _write_cell(sheet, 1, 1, table.title)
        for c, column in enumerate(table.columns, 2):
            _write_cell(sheet, 1, c, column.label)
        for r, row in enumerate(table.rows, 2):
            _write_cell(sheet, r, 1, row.label)
            for c, column in enumerate(table.columns, 2):
                address = (content.source.source_artifact_id, table.key, row.key, column.key)
                if address in tables.cells:
                    addresses[address] = f"sheet:{sheet.title}/cell:{get_column_letter(c)}{r}"
        sheet.freeze_panes = "B2"
        sheet.auto_filter.ref = f"A1:{get_column_letter(len(table.columns) + 1)}{len(table.rows) + 1}"
    for address, cell in sorted(tables.cells.items()):
        content = tables.owners[address]
        table = content.table
        assert table is not None
        sheet = sheets[address[0], address[1]]
        locator = addresses[address]
        coordinate = locator.split("/cell:")[1]
        target = sheet[coordinate]
        if cell.formula:
            target.value = _excel_formula(address, tables, addresses)
        else:
            _write_cell(sheet, target.row, target.column, cell.value)
        owner_table = tables.owners[address].table
        selected_column = owner_table.column(address[3]) if owner_table else None
        if selected_column is not None and selected_column.number_format:
            # Canonical percentages are percentage points, not fractions. Keep
            # that stored value exact and make the percent sign a display suffix.
            target.number_format = selected_column.number_format.replace("%", '"%"')
        expected = str(target.value) if target.value is not None else ""
        entry = _table_evidence(tables, address, locator, expected,
            {"table_title": f"sheet:{sheet.title}/cell:A1", "row_label": f"sheet:{sheet.title}/cell:A{target.row}",
             "column_label": f"sheet:{sheet.title}/cell:{get_column_letter(target.column)}1"})
        if entry is not None:
            pending.append(entry)
    facts, labels = {f.id: f for f in world.facts}, world.entity_names()
    covered = {cell.fact_id for cell in tables.cells.values() if cell.fact_id}
    owners: dict[str, _Content] = {}
    for content in contents:
        for fact_id in content.fact_ids:
            if fact_id not in covered:
                owners.setdefault(fact_id, content)
    groups: dict[str, list[str]] = {}
    for fact_id in sorted(owners, key=lambda fid: (facts[fid].kind, facts[fid].subject, facts[fid].period or "", fid)):
        groups.setdefault(facts[fact_id].kind.split(".")[0], []).append(fact_id)
    for domain, ids in sorted(groups.items()):
        sheet = workbook.create_sheet(_sheet_name(business_label(domain) + " records", names))
        sheet.append(_REGISTER_HEADERS)
        for r, fact_id in enumerate(ids, 2):
            fact = facts[fact_id]
            values = _register_row(fact, labels)
            for c, value in enumerate(values, 1):
                _write_cell(sheet, r, c, value)
            locator = f"sheet:{sheet.title}/cell:D{r}"
            pending.append(_Pending(owners[fact_id], locator, (fact_id,), str(values[3]), values[3],
                fact.value.unit if fact.value else None, "value", metadata={key: f"sheet:{sheet.title}/cell:{c}{r}"
                for key, c in (("measure", "A"), ("subject", "B"), ("period", "C"), ("unit", "E"))}))
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = f"A1:F{len(ids) + 1}"
    for sheet in workbook.worksheets:
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="24445C")
        for index in range(1, sheet.max_column + 1):
            sheet.column_dimensions[get_column_letter(index)].width = 32 if index == 1 else 24
    briefing.column_dimensions["B"].width = 100
    briefing.freeze_panes = "A2"
    workbook.save(stream)


def render_business_corpus(world: World, plan: NativeCorpusPlan) -> NativeCorpusResult:
    contents = prepare_business_content(world, plan)
    tables = _table_graph(world, contents)
    stream = BytesIO()
    pending: list[_Pending] = []
    addresses: dict[_Address, str] = {}
    {"docx": _docx, "pptx": _pptx, "xlsx": _xlsx}[plan.format](world, plan, contents, tables, stream, pending, addresses)
    used = {fid for content in contents for fid in content.fact_ids}
    facts = {fact.id: fact for fact in world.facts}
    created = max(facts[fid].valid_from for fid in used).isoformat()
    payload = normalise(stream.getvalue(), created=created)
    snapshot = inspect_artifact(payload, plan.format)
    extracted = {unit.locator: unit.text for unit in snapshot.units}
    _reject_identity_leaks(facts, (unit.text for unit in snapshot.units))
    evidence = []
    for entry in pending:
        text = extracted.get(entry.locator)
        if text != entry.expected_text:
            # Excel spells an integral float as an integer. That conversion is
            # allowed only when the exact source numeric value survived.
            if plan.format != "xlsx" or entry.kind != "value" or not isinstance(entry.value, (int, float)) or text is None or float(text) != entry.value:
                raise ValueError(f"native content did not survive rendering at {entry.locator}")
        assert text is not None
        dependency_locators = tuple(addresses[a] for a in tables.dependencies[entry.address]) if entry.address else ()
        address = entry.address
        evidence.append(NativeContentProvenance(
            source_artifact_id=entry.content.source.source_artifact_id, section_index=entry.content.source.section_index,
            locator=entry.locator, fact_ids=entry.fact_ids, text_sha256=hashlib.sha256(text.encode()).hexdigest(),
            value=entry.value, unit=entry.unit, kind=entry.kind,
            table_key=address[1] if address else None, row_key=address[2] if address else None,
            column_key=address[3] if address else None, dependency_locators=dependency_locators,
            metadata_locators=entry.metadata or {},
        ))
    if len({e.locator for e in evidence}) != len(evidence):
        raise ValueError("native business evidence locators are ambiguous")
    return NativeCorpusResult(payload, NativeCorpusManifest(
        artifact_id=plan.artifact_id, format=plan.format, sha256=snapshot.sha256, file_size_bytes=len(payload),
        content_units=len(contents), distinct_fact_count=len(used),
        source_artifact_count=len({c.source.source_artifact_id for c in contents}),
        explicit_page_floor=snapshot.metrics["explicit_pages"] if plan.format == "docx" else 0,
        evidence=tuple(evidence), native_metrics=snapshot.metrics,
    ))


__all__ = ["business_label", "prepare_business_content", "render_business_corpus"]
