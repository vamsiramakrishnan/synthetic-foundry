"""Append recorded visual bytes without expanding qualified native evidence."""
from __future__ import annotations

from io import BytesIO
from typing import TYPE_CHECKING
from xml.etree import ElementTree
from zipfile import ZipFile

from ..native_artifacts import inspect_artifact
from ..native_corpus import NativeCorpusResult
from ..native_query_planning import NativeWorkloadPlan, plan_native_workload
from ..native_tasks import MAX_FILE_BYTES
from ..render.ooxml import normalise
from .models import VisualError, sha256, validate_source
from .store import VisualAsset, validate_asset

if TYPE_CHECKING:
    from ..world import World


def _fit(width: int, height: int, max_width: int, max_height: int) -> tuple[int, int]:
    scale = min(max_width / width, max_height / height)
    return max(1, int(width * scale)), max(1, int(height * scale))


def attach_visual(native: NativeCorpusResult, asset: VisualAsset, *, world: World) -> NativeCorpusResult:
    """Append an unqualified infographic to a DOCX, PPTX or XLSX corpus.

    The refreshed manifest retains exactly the original evidence and fact
    counts. A visual candidate is not evidence merely because it can be opened.
    Re-plan tasks after attachment: old tasks correctly reject the new digest.
    Preserve the visual store alongside the resulting native corpus for replay.
    """
    validate_source(world, asset.record.request.spec)
    validate_asset(asset)
    if (len(native.payload) > MAX_FILE_BYTES or native.manifest.file_size_bytes != len(native.payload)
            or native.manifest.sha256 != sha256(native.payload)):
        raise VisualError("native corpus bytes do not match their manifest or exceed its budget")
    before = inspect_artifact(native.payload, native.manifest.format)
    units = {unit.locator: unit.text for unit in before.units}
    if any(sha256(units.get(item.locator, "").encode()) != item.text_sha256 for item in native.manifest.evidence):
        raise VisualError("native evidence does not match source bytes")
    source_fact_ids = {fact_id for item in native.manifest.evidence for fact_id in item.fact_ids}
    if any(fact.fact_id not in source_fact_ids for fact in asset.record.request.spec.facts):
        raise VisualError("visual facts must already be grounded in the native source")
    # Share the source inventory's canonical provenance checks rather than
    # adding a weaker, parallel image-specific definition of grounding.
    plan_native_workload(world, {native.manifest.artifact_id: native}, NativeWorkloadPlan(
        use_case_id="visual-attachment", objective="Validate source evidence for visual attachment.",
        formats=(native.manifest.format,), operations=("read",), max_tasks=1,
    ))
    output = BytesIO()
    width, height = asset.record.width, asset.record.height
    if native.manifest.format == "docx":
        from docx import Document
        from docx.shared import Inches as DocxInches
        document = Document(BytesIO(native.payload))
        fitted = _fit(width, height, int(DocxInches(6)), int(DocxInches(8)))
        document.add_picture(BytesIO(asset.payload), width=fitted[0], height=fitted[1])
        document.save(output)
    elif native.manifest.format == "pptx":
        from pptx import Presentation
        from pptx.util import Inches as PptxInches
        deck = Presentation(BytesIO(native.payload))
        if deck.slide_width is None or deck.slide_height is None:
            raise VisualError("native presentation has no slide dimensions")
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        slide_width, slide_height = int(deck.slide_width), int(deck.slide_height)
        margin = int(PptxInches(0.25))
        fitted = _fit(width, height, slide_width - 2 * margin, slide_height - 2 * margin)
        slide.shapes.add_picture(BytesIO(asset.payload), (slide_width - fitted[0]) // 2,
                                 (slide_height - fitted[1]) // 2, width=fitted[0], height=fitted[1])
        deck.save(output)
    else:
        from openpyxl import load_workbook
        from openpyxl.drawing.image import Image
        workbook = load_workbook(BytesIO(native.payload), data_only=False)
        ordinal = 1
        while f"Visual {ordinal}" in workbook.sheetnames:
            ordinal += 1
        sheet = workbook.create_sheet(f"Visual {ordinal}")
        picture = Image(BytesIO(asset.payload))
        picture.width, picture.height = _fit(width, height, 1200, 1200)
        sheet.add_image(picture, "A1")
        workbook.save(output)
    with ZipFile(BytesIO(native.payload)) as package:
        stamp = "1980-01-01T00:00:00Z"
        if "docProps/core.xml" in package.namelist():
            root = ElementTree.fromstring(package.read("docProps/core.xml"))
            stamp = root.findtext("{http://purl.org/dc/terms/}created") or stamp
    payload = normalise(output.getvalue(), created=stamp)
    if len(payload) > MAX_FILE_BYTES:
        raise VisualError("attached native corpus exceeds the file byte budget")
    after = inspect_artifact(payload, native.manifest.format)
    updated = {unit.locator: unit.text for unit in after.units}
    # Some Office writers rewrite unsupported parts. Refuse an attachment if
    # doing so moved or changed any existing grounded evidence.
    if any(sha256(updated.get(item.locator, "").encode()) != item.text_sha256 for item in native.manifest.evidence):
        raise VisualError("visual attachment changed existing native evidence")
    metrics = {**native.manifest.native_metrics, **after.metrics}
    metrics["unqualified_visual_count"] = metrics.get("unqualified_visual_count", 0) + 1
    manifest = native.manifest.model_copy(update={
        "sha256": sha256(payload), "file_size_bytes": len(payload), "native_metrics": metrics,
    })
    return NativeCorpusResult(payload, manifest)
