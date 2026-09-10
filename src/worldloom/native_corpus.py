"""Grounded native corpora assembled from the company's authored artifact IR.

Size is a construction requirement, never permission to manufacture filler.
Every content unit reuses one distinct authored section with canonical evidence.
The explicit DOCX breaks establish a page floor; actual pagination depends on
Word's layout engine and is deliberately not reported as a measured page count.
"""
from __future__ import annotations

import hashlib
from dataclasses import dataclass
from io import BytesIO
from typing import TYPE_CHECKING, Literal

from pydantic import Field

from .models import Model
from .narrative import references
from .native_artifacts import inspect_artifact
from .presentation import of as presentation_of
from .render.ooxml import normalise
from .render.values import corpus_locale

if TYPE_CHECKING:
    from .world import World


class NativeContent(Model):
    source_artifact_id: str
    section_index: int = Field(ge=0)
    placement: Literal["body", "notes"] = "body"


class NativeCorpusPlan(Model):
    artifact_id: str = Field(min_length=1)
    format: Literal["docx", "pptx", "xlsx"]
    title: str = Field(min_length=1)
    minimum_units: int = Field(default=1, ge=1, le=10000)
    minimum_distinct_facts: int = Field(default=1, ge=1)
    contents: tuple[NativeContent, ...]


class NativeContentProvenance(Model):
    source_artifact_id: str
    section_index: int
    locator: str
    fact_ids: tuple[str, ...]
    text_sha256: str
    value: float | str | None = None
    unit: str | None = None


class NativeCorpusManifest(Model):
    artifact_id: str
    format: Literal["docx", "pptx", "xlsx"]
    sha256: str
    file_size_bytes: int
    content_units: int
    distinct_fact_count: int
    source_artifact_count: int
    explicit_page_floor: int = 0
    evidence: tuple[NativeContentProvenance, ...]
    ballast_is_evidence: bool = False


@dataclass(frozen=True)
class NativeCorpusResult:
    payload: bytes
    manifest: NativeCorpusManifest


@dataclass(frozen=True)
class _Content:
    source: NativeContent
    heading: str
    body: str
    fact_ids: tuple[str, ...]


def _contents(world: World, plan: NativeCorpusPlan) -> tuple[_Content, ...]:
    if len(plan.contents) < plan.minimum_units:
        raise ValueError(f"insufficient grounded content: need {plan.minimum_units}, have {len(plan.contents)}")
    if len(plan.contents) > 10000:
        raise ValueError("native corpus exceeds bounded content budget")
    artifacts = {ir.id: ir for ir in world.artifact_irs}
    facts = {fact.id: fact for fact in world.facts}
    seen: set[str] = set()
    result: list[_Content] = []
    for content in plan.contents:
        ir = artifacts.get(content.source_artifact_id)
        if ir is None or content.section_index >= len(ir.sections):
            raise ValueError(f"unknown authored section: {content.source_artifact_id}:{content.section_index}")
        section = ir.sections[content.section_index]
        if not section.body or not section.body.strip():
            raise ValueError("native corpus requires authored prose; narrate the source section first")
        if content.placement == "notes" and plan.format != "pptx":
            raise ValueError("speaker notes require pptx")
        fact_ids = tuple(sorted(set(section.fact_ids) | set(references.referenced(section.body))))
        if not fact_ids:
            raise ValueError("native corpus section has no canonical evidence")
        unknown = sorted(set(fact_ids) - set(facts))
        malformed = references.unresolved(section.body, facts)
        if unknown or malformed:
            raise ValueError(f"unknown canonical facts: {sorted(set(unknown + malformed))}")
        if references.bare_numbers(section.body):
            raise ValueError("native corpus prose must reference canonical figures")
        # Citation metadata alone does not prove a fact occurs in the bytes.
        # Require an explicit reference for every provenance fact reported.
        rendered_ids = tuple(sorted(set(references.referenced(section.body))))
        if not rendered_ids:
            raise ValueError("native corpus evidence must occur in authored prose")
        body = references.substitute(section.body, facts, locale=corpus_locale(world), presentation=presentation_of(world))
        identity = " ".join(body.split()).casefold()
        if identity in seen:
            raise ValueError("duplicate grounded content: headings or renaming do not increase coverage")
        seen.add(identity)
        result.append(_Content(content, section.heading, body, rendered_ids))
    distinct_facts = {fact_id for content in result for fact_id in content.fact_ids}
    if len(distinct_facts) < plan.minimum_distinct_facts:
        raise ValueError(f"insufficient distinct canonical facts: need {plan.minimum_distinct_facts}, have {len(distinct_facts)}")
    return tuple(result)


def plan_native_corpus(
    world: World,
    *,
    artifact_id: str,
    format: Literal["docx", "pptx", "xlsx"],
    title: str,
    minimum_units: int,
    minimum_distinct_facts: int = 1,
    source_artifact_ids: tuple[str, ...] | None = None,
) -> NativeCorpusPlan:
    """Select authored sections in stable source order, refusing missing scope.

    The planner never calls a writer. A shortage must feed back to the existing
    episode and narration pipeline to produce additional business evidence.
    """
    requested = set(source_artifact_ids or ())
    existing = {ir.id for ir in world.artifact_irs}
    if requested - existing:
        raise ValueError(f"unknown source artifacts: {sorted(requested - existing)}")
    contents = tuple(
        NativeContent(source_artifact_id=ir.id, section_index=index)
        for ir in sorted(world.artifact_irs, key=lambda item: item.id)
        if source_artifact_ids is None or ir.id in requested
        for index, section in enumerate(ir.sections)
        if section.body and references.referenced(section.body)
    )
    plan = NativeCorpusPlan(artifact_id=artifact_id, format=format, title=title, minimum_units=minimum_units, minimum_distinct_facts=minimum_distinct_facts, contents=contents)
    _contents(world, plan)
    return plan


def render_native_corpus(world: World, plan: NativeCorpusPlan) -> NativeCorpusResult:
    """Render native bytes and hidden provenance from the same canonical world."""
    contents = _contents(world, plan)
    stream = BytesIO()
    locators: list[str] = []
    facts = {fact.id: fact for fact in world.facts}
    used = {fact_id for content in contents for fact_id in content.fact_ids}
    fact_locators: dict[str, str] = {}
    if plan.format == "docx":
        from docx import Document
        document = Document()
        document.core_properties.title = plan.title
        paragraph = 0
        for index, content in enumerate(contents):
            if index:
                document.add_page_break()
                paragraph += 1
            document.add_heading(content.heading, level=1)
            document.add_paragraph(content.body)
            paragraph += 2
            locators.append(f"paragraph:{paragraph}")
        document.save(stream)
    elif plan.format == "pptx":
        from pptx import Presentation
        from pptx.util import Inches, Pt
        presentation = Presentation()
        presentation.slide_width = Inches(13.333)
        presentation.slide_height = Inches(7.5)
        presentation.core_properties.title = plan.title
        for index, content in enumerate(contents, 1):
            slide = presentation.slides.add_slide(presentation.slide_layouts[6])
            heading = slide.shapes.add_textbox(Inches(.6), Inches(.4), Inches(12), Inches(.8))
            heading.text_frame.text = content.heading
            heading.text_frame.paragraphs[0].font.size = Pt(24)
            if content.source.placement == "notes":
                slide.notes_slide.notes_text_frame.text = content.body
                locators.append(f"slide:{index}/notes")
            else:
                if len(content.body) > 2400:
                    raise ValueError("slide body exceeds readable budget; split authored sections or use notes")
                box = slide.shapes.add_textbox(Inches(.7), Inches(1.5), Inches(11.9), Inches(5.3))
                box.text_frame.word_wrap = True
                box.text_frame.text = content.body
                box.text_frame.paragraphs[0].font.size = Pt(18)
                locators.append(f"slide:{index}/shape:2/text")
        presentation.save(stream)
    else:
        from openpyxl import Workbook
        workbook = Workbook()
        workbook.properties.title = plan.title
        sheet = workbook.active
        assert sheet is not None
        sheet.title = "Evidence"
        sheet.append(["Section", "Authored evidence"])
        for index, content in enumerate(contents, 2):
            if len(content.body) > 32767:
                raise ValueError("authored section exceeds Excel cell limit; split the source section")
            sheet.append([content.heading, content.body])
            # Company text beginning with '=' must remain text, not a formula.
            for cell in sheet[index]:
                cell.data_type = "s"
            locators.append(f"sheet:Evidence/cell:B{index}")
        sheet.freeze_panes = "A2"
        sheet.column_dimensions["A"].width = 36
        sheet.column_dimensions["B"].width = 100
        fact_sheet = workbook.create_sheet("Facts")
        fact_sheet.append(["Fact ID", "Value", "Unit", "Measure", "Subject"])
        for index, fact_id in enumerate(sorted(used), 2):
            fact = facts[fact_id]
            value = fact.value.amount if fact.value is not None else fact.text_value
            unit = fact.value.unit if fact.value is not None else "text"
            if isinstance(value, str) and len(value) > 32767:
                raise ValueError("canonical fact text exceeds Excel cell limit")
            fact_sheet.append([fact.id, value, unit, fact.kind, fact.subject])
            for cell in fact_sheet[index]:
                if isinstance(cell.value, str):
                    cell.data_type = "s"
            fact_locators[fact_id] = f"sheet:Facts/cell:B{index}"
        fact_sheet.freeze_panes = "A2"
        fact_sheet.column_dimensions["A"].width = 24
        fact_sheet.column_dimensions["B"].width = 32
        fact_sheet.column_dimensions["C"].width = 16
        fact_sheet.column_dimensions["D"].width = 36
        fact_sheet.column_dimensions["E"].width = 36
        workbook.save(stream)
    # Use the ledger's deterministic valid time, never the host's clock.
    created = max(fact.valid_from for fact in world.facts if fact.id in used).isoformat()
    payload = normalise(stream.getvalue(), created=created)
    extracted = {unit.locator: unit.text for unit in inspect_artifact(payload, plan.format).units}
    for content, locator in zip(contents, locators, strict=True):
        if extracted.get(locator) != content.body:
            raise ValueError(f"native content did not survive rendering at {locator}")
    evidence = tuple(
        NativeContentProvenance(
            source_artifact_id=content.source.source_artifact_id,
            section_index=content.source.section_index,
            locator=locator,
            fact_ids=content.fact_ids,
            text_sha256=hashlib.sha256(content.body.encode()).hexdigest(),
        )
        for content, locator in zip(contents, locators, strict=True)
    )
    sources: dict[str, NativeContent] = {}
    for content in contents:
        for fact_id in content.fact_ids:
            sources.setdefault(fact_id, content.source)
    fact_evidence: list[NativeContentProvenance] = []
    for fact_id, locator in sorted(fact_locators.items()):
        fact = facts[fact_id]
        text = extracted.get(locator)
        value = fact.value.amount if fact.value is not None else fact.text_value
        if text is None:
            raise ValueError(f"canonical fact is missing from native bytes at {locator}")
        if fact.value is not None:
            if float(text) != fact.value.amount:
                raise ValueError(f"canonical number did not survive native rendering at {locator}")
        elif text != value:
            raise ValueError(f"canonical text did not survive native rendering at {locator}")
        source = sources[fact_id]
        fact_evidence.append(NativeContentProvenance(
            source_artifact_id=source.source_artifact_id, section_index=source.section_index,
            locator=locator, fact_ids=(fact_id,), text_sha256=hashlib.sha256(text.encode()).hexdigest(),
            value=value, unit=fact.value.unit if fact.value is not None else None,
        ))
    return NativeCorpusResult(payload, NativeCorpusManifest(
        artifact_id=plan.artifact_id, format=plan.format,
        sha256=hashlib.sha256(payload).hexdigest(), file_size_bytes=len(payload),
        content_units=len(contents), distinct_fact_count=len(used),
        source_artifact_count=len({content.source.source_artifact_id for content in contents}),
        explicit_page_floor=len(contents) if plan.format == "docx" else 0,
        evidence=evidence + tuple(fact_evidence),
    ))


__all__ = ["NativeContent", "NativeCorpusPlan", "NativeContentProvenance", "NativeCorpusManifest", "NativeCorpusResult", "plan_native_corpus", "render_native_corpus"]
