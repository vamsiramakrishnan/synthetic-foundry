"""The controlled Word document: what a board paper actually is.

Cover page, document control (reference, version, owner, approvers,
classification), revision history, approvals and review record, a contents
field, numbered sections and subsections, captioned tables and native charts,
appendices, and running headers and footers carrying the classification and
``Page X of Y`` as fields Word computes. The layout helpers are the legacy
renderer's own (`render.docx`), so a table or chart looks the same in both
profiles; only the document around them changes.

A *revision* renders through the same function: a draft is marked as one and
states the facts that were true at its date, a reviewed version carries the
reviewers' comments as native Word comments, and an amendment shows each
restated figure as a tracked deletion and insertion by its author.
"""

from __future__ import annotations

from itertools import count
from typing import TYPE_CHECKING, Any

from ... import longform
from .. import Rendered, ooxml
from .. import docx as legacy

if TYPE_CHECKING:  # pragma: no cover
    from ...compiler.style import StyleGenome
    from ...models import ArtifactIR
    from . import Context

MEDIA_TYPE = legacy.MEDIA_TYPE

_TS_TITLE, _TS_HEADING, _TS_SUBHEADING, _TS_BODY, _TS_CAPTION = range(5)


def _pt(value: float):  # type: ignore[no-untyped-def]
    from docx.shared import Pt

    return Pt(value)


def _muted(run, g: StyleGenome, size_band: int = _TS_CAPTION) -> None:  # type: ignore[no-untyped-def]
    run.font.size = _pt(g.type_scale[size_band])
    run.font.color.rgb = legacy._rgb(legacy._MUTED)


def _heading(document, text: str, level: int, g: StyleGenome):  # type: ignore[no-untyped-def]
    heading = document.add_heading(text, level=level)
    band = _TS_HEADING if level <= 1 else _TS_SUBHEADING
    legacy._style_heading(
        heading, size_pt=g.type_scale[band], colour_hex=g.colour_roles["body_text"],
        alignment="left", space_before_pt=legacy._space_pt(g, 0 if level <= 1 else 1),
    )
    return heading


#: One styled cell: ``(text, bold, right_aligned, fill_hex_or_empty, colour_hex)``.
_Cell = tuple[str, bool, bool, str, str]


def _write_rows(document, rows: list[list[_Cell]], g: StyleGenome, size_band: int) -> None:  # type: ignore[no-untyped-def]
    """Append a table built as one XML string.

    The legacy renderer styles every run through python-docx's object API,
    which costs a tree search per property; a schedule of a hundred and sixty
    stores rendered three times over a revision chain spent eight seconds
    there. The XML written here is what those calls produce (a shaded
    ``w:tcPr``, a ``w:jc`` for numbers, ``w:b``/``w:color``/``w:sz`` on the
    run), built once and parsed once.
    """
    from xml.sax.saxutils import escape

    from docx.oxml import parse_xml
    from docx.oxml.ns import nsdecls

    if not rows:
        return
    columns = max(len(row) for row in rows)
    grid = document.add_table(rows=0, cols=columns)
    legacy._apply_table_borders(grid, g.gridline_policy, g.rule_weight)
    legacy._apply_cell_padding(grid, legacy._cell_padding_pt(g))
    size = round(g.type_scale[size_band] * 2)
    xml_rows = []
    for row in rows:
        cells = []
        for text, bold, right, fill, colour in row:
            shade = f'<w:shd w:val="clear" w:color="auto" w:fill="{fill}"/>' if fill else ""
            align = '<w:pPr><w:jc w:val="right"/></w:pPr>' if right else ""
            weight = "<w:b/>" if bold else ""
            run = (f'<w:r><w:rPr>{weight}<w:color w:val="{colour}"/><w:sz w:val="{size}"/></w:rPr>'
                   f'<w:t xml:space="preserve">{escape(text)}</w:t></w:r>') if text else ""
            cells.append(f'<w:tc><w:tcPr><w:tcW w:w="0" w:type="auto"/>{shade}</w:tcPr><w:p>{align}{run}</w:p></w:tc>')
        xml_rows.append("<w:tr>" + "".join(cells) + "</w:tr>")
    fragment = parse_xml(f'<w:tbl {nsdecls("w")}>' + "".join(xml_rows) + "</w:tbl>")
    for tr in list(fragment):
        grid._tbl.append(tr)


def _grid(document, header: list[str], rows: list[list[str]], g: StyleGenome) -> None:  # type: ignore[no-untyped-def]
    """A plain text table in the house style: shaded header, genome rules."""
    roles = g.colour_roles
    spec: list[list[_Cell]] = [[(text, True, False, roles["header_fill"], roles["header_text"]) for text in header]]
    spec.extend([(text, False, False, "", roles["body_text"]) for text in values] for values in rows)
    _write_rows(document, spec, g, _TS_CAPTION)


def _ir_table(document, table, g: StyleGenome, locale) -> None:  # type: ignore[no-untyped-def]
    """An IR table, cell for cell as `render.docx._table` draws it."""
    from ..values import format_value

    roles = g.colour_roles
    spec: list[list[_Cell]] = [[(table.title, True, False, roles["header_fill"], roles["header_text"])] + [
        (column.label, True, True, roles["header_fill"], roles["header_text"]) for column in table.columns]]
    for row in table.rows:
        fill = roles["subtotal_fill"] if row.emphasis else ""
        ink = roles["subtotal_text"] if row.emphasis else roles["body_text"]
        cells: list[_Cell] = [(row.label, row.emphasis, False, fill, ink)]
        for column in table.columns:
            cell = row.cells.get(column.key)
            value = cell.value if cell else None
            text = format_value(value, column.number_format, locale=locale) if cell else ""
            negative = isinstance(value, (int, float)) and value < 0
            colour = roles["negative_text"] if negative and not row.emphasis else ink
            cells.append((text, row.emphasis, isinstance(value, (int, float)), fill, colour))
        spec.append(cells)
    _write_rows(document, spec, g, _TS_BODY)
    if table.note:
        note = document.add_paragraph()
        run = note.add_run(table.note)
        run.italic = True
        _muted(run, g)


def _running_heads(document, doc: longform.LongDocument, g: StyleGenome) -> None:  # type: ignore[no-untyped-def]
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    section = document.sections[0]
    section.different_first_page_header_footer = True
    head = section.header.paragraphs[0]
    head.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    marker = "DRAFT | " if doc.draft else ""
    run = head.add_run(f"{marker}{doc.company} | {doc.title} | {doc.classification} | ")
    _muted(run, g)
    legacy._field(head, ' STYLEREF "Heading 1" ')
    for run in head.runs:
        _muted(run, g)

    foot = section.footer.paragraphs[0]
    foot.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = foot.add_run(f"{doc.reference} | Version {doc.revision.version} | {doc.revision.status} | "
                       f"{doc.classification} | Page ")
    _muted(run, g)
    legacy._field(foot, " PAGE ")
    _muted(foot.add_run(" of "), g)
    legacy._field(foot, " NUMPAGES ")
    for run in foot.runs:
        _muted(run, g)


def _cover(document, doc: longform.LongDocument, g: StyleGenome) -> None:  # type: ignore[no-untyped-def]
    company = document.add_paragraph()
    run = company.add_run(doc.company.upper())
    run.bold = True
    run.font.size = _pt(g.type_scale[_TS_SUBHEADING])
    genre = document.add_paragraph()
    _muted(genre.add_run(doc.genre), g, _TS_SUBHEADING)
    for _ in range(3):
        document.add_paragraph()
    title = document.add_heading(doc.title, level=0)
    legacy._style_heading(title, size_pt=g.type_scale[_TS_TITLE], colour_hex=g.colour_roles["body_text"],
                          alignment="left")
    if doc.subtitle:
        subtitle = document.add_paragraph()
        _muted(subtitle.add_run(doc.subtitle), g, _TS_SUBHEADING)
    if doc.draft:
        draft = document.add_paragraph()
        mark = draft.add_run("DRAFT FOR REVIEW")
        mark.bold = True
        mark.font.size = _pt(g.type_scale[_TS_HEADING])
        mark.font.color.rgb = legacy._rgb(g.colour_roles["negative_text"])
    for _ in range(2):
        document.add_paragraph()
    rows = [
        ["Reference", doc.reference],
        ["Version", f"{doc.revision.version} ({doc.revision.status})"],
        ["Date", doc.revision.at.strftime("%d %B %Y")],
        ["Classification", doc.classification],
        ["Prepared by", f"{doc.author.name}, {doc.author.title}".strip(", ")],
    ]
    if doc.approver is not None:
        rows.append(["Approved by", f"{doc.approver.name}, {doc.approver.title}"])
    _grid(document, ["Document", ""], rows, g)
    notice = document.add_paragraph()
    _muted(notice.add_run(doc.metadata.get("note", "Synthetic corpus generated by Worldloom. Not a real company.")), g)
    document.add_page_break()


def _control(document, doc: longform.LongDocument, g: StyleGenome) -> None:  # type: ignore[no-untyped-def]
    _heading(document, "Document control", 1, g)
    _grid(document, ["Field", "Value"], [[label, value] for label, value in longform.control_rows(doc)], g)
    _heading(document, "Revision history", 2, g)
    _grid(document, ["Version", "Date", "Author", "Status", "Description of change"],
          [list(row) for row in longform.history_rows(doc)], g)
    _heading(document, "Approvals", 2, g)
    _grid(document, ["Role", "Name", "Title", "Decision", "Date"],
          [list(row) for row in longform.approval_rows(doc)], g)
    if doc.comments:
        _heading(document, "Review record", 2, g)
        _grid(document, ["Section", "Reviewer", "Comment", "Resolution"],
              [[c.part, c.author.name, c.text, c.resolution] for c in doc.comments], g)
    if doc.long:
        document.add_page_break()


def _contents(document, doc: longform.LongDocument, g: StyleGenome) -> None:  # type: ignore[no-untyped-def]
    """A TOC field whose cached result is the numbered outline.

    A field, so Word rebuilds it (with page numbers) on open; a cached result,
    so a reader, an extractor or a connector that never opens Word still sees
    the outline rather than an empty paragraph.
    """
    from docx.oxml import parse_xml
    from docx.oxml.ns import nsdecls

    _heading(document, "Contents", 1, g)
    entries = longform.contents(doc)
    first = document.add_paragraph()
    first._element.append(parse_xml(f'<w:r {nsdecls("w")}><w:fldChar w:fldCharType="begin"/></w:r>'))
    first._element.append(parse_xml(
        f'<w:r {nsdecls("w")}><w:instrText xml:space="preserve"> TOC \\o "1-2" \\h \\z \\u </w:instrText></w:r>'))
    first._element.append(parse_xml(f'<w:r {nsdecls("w")}><w:fldChar w:fldCharType="separate"/></w:r>'))
    paragraphs = [first]
    for index, (level, label) in enumerate(entries):
        paragraph = first if index == 0 else document.add_paragraph()
        paragraph.paragraph_format.left_indent = _pt(0 if level == 1 else 14)
        run = paragraph.add_run(label)
        run.font.size = _pt(g.type_scale[_TS_BODY])
        run.bold = level == 1
        if index:
            paragraphs.append(paragraph)
    paragraphs[-1]._element.append(parse_xml(f'<w:r {nsdecls("w")}><w:fldChar w:fldCharType="end"/></w:r>'))
    document.add_page_break()


def _tracked(paragraph, segments, revision: longform.Revision, g: StyleGenome, ids) -> None:  # type: ignore[no-untyped-def]
    """Write *segments* with each replacement as a tracked change."""
    from xml.sax.saxutils import escape

    from docx.oxml import parse_xml
    from docx.oxml.ns import nsdecls

    stamp = revision.at.strftime("%Y-%m-%dT%H:%M:00Z")
    author = escape(revision.by.name, {'"': "&quot;"})
    size = int(g.type_scale[_TS_BODY] * 2)
    for old, new in segments:
        if new is None:
            run = paragraph.add_run(old)
            run.font.size = _pt(g.type_scale[_TS_BODY])
            continue
        paragraph._p.append(parse_xml(
            f'<w:del {nsdecls("w")} w:id="{next(ids)}" w:author="{author}" w:date="{stamp}">'
            f'<w:r><w:rPr><w:sz w:val="{size}"/></w:rPr><w:delText xml:space="preserve">{escape(old)}</w:delText></w:r></w:del>'))
        paragraph._p.append(parse_xml(
            f'<w:ins {nsdecls("w")} w:id="{next(ids)}" w:author="{author}" w:date="{stamp}">'
            f'<w:r><w:rPr><w:sz w:val="{size}"/></w:rPr><w:t xml:space="preserve">{escape(new)}</w:t></w:r></w:ins>'))


def _blocks(document, blocks, doc: longform.LongDocument, g: StyleGenome, ctx: Context,
            chart_index, ids) -> None:  # type: ignore[no-untyped-def]
    for block in blocks:
        if block.kind == "prose":
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.space_after = _pt(legacy._space_pt(g, 2))
            if block.segments:
                _tracked(paragraph, block.segments, doc.revision, g, ids)
            else:
                run = paragraph.add_run(block.text)
                run.font.size = _pt(g.type_scale[_TS_BODY])
                run.font.color.rgb = legacy._rgb(g.colour_roles["body_text"])
        elif block.kind == "note":
            paragraph = document.add_paragraph()
            run = paragraph.add_run(block.text)
            run.italic = True
            _muted(run, g, _TS_BODY)
        elif block.kind == "table" and block.table is not None:
            caption = document.add_paragraph(style="Caption")
            caption.add_run(f"Table {block.number}: {block.caption}").bold = True
            _ir_table(document, block.table, g, ctx.locale)
            if block.source:
                source = document.add_paragraph()
                _muted(source.add_run(f"Source: {block.source}"), g)
        elif block.kind == "figure" and block.chart is not None and block.table is not None:
            categories, series = legacy._chart_series(block.chart, block.table)
            if not categories or not series:
                continue
            from docx.oxml.parser import parse_xml

            caption = document.add_paragraph(style="Caption")
            caption.add_run(f"Figure {block.number}: {block.caption}").bold = True
            rid = legacy._add_chart_part(document, legacy._chart_xml(block.chart, categories, series))
            drawing = document.add_paragraph()
            drawing._p.append(parse_xml(legacy._chart_drawing_xml(rid, next(chart_index))))
            if block.source:
                source = document.add_paragraph()
                _muted(source.add_run(f"Source: {block.source}"), g)


def render_document(doc: longform.LongDocument, ctx: Context, ir: ArtifactIR) -> bytes:
    """One long document as DOCX bytes."""
    docx_pkg = legacy._require_docx()
    document = docx_pkg.Document()
    g = legacy._genome_for(ir)
    legacy._apply_typeface(document, g)
    legacy._page_setup(document)
    _running_heads(document, doc, g)
    _cover(document, doc, g)
    _control(document, doc, g)
    if doc.long:
        _contents(document, doc, g)

    chart_index = count(1)
    ids = count(9001)
    anchors: dict[str, Any] = {}
    for position, part in enumerate(doc.all_parts()):
        # A report opens each numbered section on a page of its own; a short
        # controlled document runs on, and only its appendices turn the page.
        if position and (doc.long or (part.appendix and part is doc.appendices[0])):
            document.add_page_break()
        heading = _heading(document, part.label, 1, g)
        anchors[part.number] = heading
        _blocks(document, part.blocks, doc, g, ctx, chart_index, ids)
        for child in part.children:
            _heading(document, f"{child.number} {child.heading}", 2, g)
            _blocks(document, child.blocks, doc, g, ctx, chart_index, ids)

    if doc.revision.status == "Reviewed":
        # Native comments on the version the reviewers actually read. The
        # approved version carries the same comments as its review record.
        for comment in doc.comments:
            heading = anchors.get(comment.part)
            if heading is None or not heading.runs:
                continue
            added = document.add_comment(heading.runs, text=comment.text, author=comment.author.name,
                                         initials="".join(w[0] for w in comment.author.name.split() if w))
            added._comment_elm.date = comment.at

    properties = document.core_properties
    properties.title = doc.title
    properties.subject = doc.genre
    properties.author = doc.author.name
    properties.last_modified_by = doc.revision.by.name
    properties.category = doc.genre
    properties.content_status = doc.revision.status
    properties.identifier = doc.reference
    properties.version = doc.revision.version
    properties.keywords = ", ".join(doc.labels)
    properties.comments = doc.metadata.get("note", "Synthetic corpus generated by Worldloom.")
    properties.revision = len(doc.history)

    from io import BytesIO

    buffer = BytesIO()
    document.save(buffer)
    return ooxml.normalise(buffer.getvalue(), created=doc.revision.at.isoformat())


def render_all(ctx: Context) -> list[Rendered]:
    out: list[Rendered] = []
    for ir in ctx.world.artifact_irs:
        intent = ctx.intent(ir)
        if intent.artifact_type not in legacy.HANDLES:
            continue
        out.append(Rendered(artifact_id=ir.id, path=ctx.path(ir, "docx"), media_type=MEDIA_TYPE,
                            payload=render_document(ctx.doc(ir), ctx, ir)))
    return out


def revisions(ctx: Context) -> list[Rendered]:
    """Every non-canonical version of every controlled Word document."""
    out: list[Rendered] = []
    for ir in ctx.world.artifact_irs:
        intent = ctx.intent(ir)
        if intent.artifact_type not in legacy.HANDLES:
            continue
        for revision in ctx.history(ir):
            if revision.as_of is None:
                continue
            out.append(Rendered(artifact_id=ir.id, path=ctx.revision_path(ir, revision, "docx"),
                                media_type=MEDIA_TYPE,
                                payload=render_document(ctx.doc(ir, revision), ctx, ir)))
    return out


def agendas(ctx: Context) -> list[Rendered]:
    """The agenda of each executive committee pack, as the Word file it is."""
    from . import FAMILIES_DIR

    irs = {ir.id: ir for ir in ctx.world.artifact_irs}
    out: list[Rendered] = []
    for family in ctx.families:
        if family.kind != "board":
            continue
        items = [irs[a] for a, _ in family.members if a in irs]
        minutes = next((ir for ir in items if ctx.intent(ir).artifact_type == "meeting_minutes"), None)
        chair_ir = minutes or items[0]
        chair = longform._person(ctx.world, ctx.intent(chair_ir).approver_id) or longform._person(
            ctx.world, ctx.intent(chair_ir).author_id)
        docx_pkg = legacy._require_docx()
        document = docx_pkg.Document()
        g = legacy._genome_for(chair_ir)
        legacy._apply_typeface(document, g)
        legacy._page_setup(document)
        title = document.add_heading(f"{ctx.world.company.name}: {family.title}", level=0)
        legacy._style_heading(title, size_pt=g.type_scale[_TS_TITLE], colour_hex=g.colour_roles["body_text"],
                              alignment="left")
        _heading(document, "Agenda", 1, g)
        if chair is not None:
            document.add_paragraph(f"Chair: {chair.name}, {chair.title}")
        rows = []
        for index, ir in enumerate(items, start=1):
            intent = ctx.intent(ir)
            presenter = longform._person(ctx.world, intent.author_id)
            role = family.roles()[ir.id]
            purpose = "For decision" if "decision" in role.lower() else (
                "For approval" if intent.artifact_type == "meeting_minutes" else "For noting")
            rows.append([str(index), ir.title, presenter.name if presenter else intent.author_id,
                         purpose, ir.id])
        _grid(document, ["Item", "Paper", "Presented by", "Purpose", "Reference"], rows, g)
        _heading(document, "Papers", 1, g)
        for ir in items:
            document.add_paragraph(f"{ir.id}: {ir.title} ({longform.genre_label(ctx.intent(ir).artifact_type)})",
                                   style="List Bullet")
        from io import BytesIO

        buffer = BytesIO()
        document.save(buffer)
        created = max(ctx.doc(ir).created_at for ir in items).isoformat()
        out.append(Rendered(artifact_id="", path=f"{FAMILIES_DIR}/{family.key}/agenda.docx",
                            media_type=MEDIA_TYPE, payload=ooxml.normalise(buffer.getvalue(), created=created)))
    return out
