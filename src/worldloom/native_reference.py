"""Reference native submissions prove contracts are satisfiable before trials.

Only explicit native locators have setters. Unsupported edits refuse rather than
silently generating an easier task or counting a planner defect as model failure.
"""
from __future__ import annotations

import base64
import re
from collections.abc import Mapping
from decimal import Decimal, localcontext
from io import BytesIO
from typing import Any

from .native_artifacts import inspect_artifact
from .native_tasks import (
    MAX_FILE_BYTES,
    NativeAnswer,
    NativeFile,
    NativeGrade,
    NativeSubmission,
    NativeTask,
    _number,
    expected_cell_type,
    grade_native_task,
)
from .render.ooxml import normalise


def _output(task: NativeTask, inputs: Mapping[str, bytes]) -> bytes:
    output = task.output
    assert output is not None
    source = inputs[output.source_artifact_id] if output.source_artifact_id else None
    create = source is None
    document: Any
    if output.format == "docx":
        from docx import Document
        document = Document(BytesIO(source)) if source is not None else Document()
    elif output.format == "xlsx":
        from openpyxl import Workbook, load_workbook
        document = load_workbook(BytesIO(source), data_only=False) if source is not None else Workbook()
    elif output.format == "pptx":
        from pptx import Presentation
        document = Presentation(BytesIO(source)) if source is not None else Presentation()
    else:
        raise ValueError("reference native output format unsupported")
    for assertion in output.assertions:
        assert assertion.target is not None and assertion.expected is not None
        locator, expected = assertion.target.locator, assertion.expected
        if output.format == "docx":
            match = re.fullmatch(r"paragraph:([1-9][0-9]*)", locator)
            if match is None:
                raise ValueError("unsupported DOCX output locator")
            index = int(match[1])
            if index > 10000:
                raise ValueError("DOCX reference paragraph budget exceeded")
            if create:
                while len(document.paragraphs) < index:
                    document.add_paragraph("")
            if index > len(document.paragraphs):
                raise ValueError("update paragraph does not exist")
            paragraph = document.paragraphs[index - 1]
            # Retain the existing paragraph properties and first run formatting.
            if paragraph.runs:
                paragraph.runs[0].text = expected
                for run in paragraph.runs[1:]:
                    run.text = ""
            else:
                paragraph.add_run(expected)
        elif output.format == "xlsx":
            match = re.fullmatch(r"sheet:([^/]+)/cell:([A-Z]{1,3}[1-9][0-9]{0,6})", locator)
            if match is None:
                raise ValueError("unsupported XLSX output locator")
            sheet, cell = match.groups()
            if sheet not in document.sheetnames:
                if not create:
                    raise ValueError("update sheet does not exist")
                document.create_sheet(sheet)
            target_cell = document[sheet][cell]
            required_type = expected_cell_type(assertion, target_cell.data_type if not create else None)
            if required_type == "n":
                target_cell.value = _number(expected)
            elif required_type == "b":
                if expected not in {"True", "False"}:
                    raise ValueError("boolean cell requires True or False")
                target_cell.value = expected == "True"
            elif required_type in {"s", "f"}:
                if required_type == "f" and not expected.startswith("="):
                    raise ValueError("formula cell requires a formula expression")
                target_cell.value = expected
                target_cell.data_type = required_type
            else:
                raise ValueError("unsupported reference cell type")
        else:
            match = re.fullmatch(r"slide:([1-9][0-9]*)/(notes|hidden|shape:([1-9][0-9]*)/text)", locator)
            if match is None:
                raise ValueError("unsupported PPTX output locator")
            index = int(match[1])
            if index > 1000:
                raise ValueError("PPTX reference slide budget exceeded")
            if create:
                while len(document.slides) < index:
                    document.slides.add_slide(document.slide_layouts[6])
            if index > len(document.slides):
                raise ValueError("update slide does not exist")
            slide = document.slides[index - 1]
            if match[2] == "notes":
                slide.notes_slide.notes_text_frame.text = expected
            elif match[2] == "hidden":
                if expected not in {"true", "false"}:
                    raise ValueError("hidden flag must be true or false")
                slide._element.set("show", "0" if expected == "true" else "1")
            else:
                shape_index = int(match[3])
                if shape_index > 1000:
                    raise ValueError("PPTX reference shape budget exceeded")
                if create:
                    while len(slide.shapes) < shape_index:
                        slide.shapes.add_textbox(0, 0, 1000000, 1000000)
                if shape_index > len(slide.shapes) or not slide.shapes[shape_index - 1].has_text_frame:
                    raise ValueError("update text shape does not exist")
                slide.shapes[shape_index - 1].text = expected
    stream = BytesIO()
    document.save(stream)
    return normalise(stream.getvalue(), created="1980-01-01T00:00:00Z")


def reference_submission(task: NativeTask, inputs: Mapping[str, bytes]) -> NativeSubmission:
    units: dict[str, dict[str, str]] = {}
    for item in task.inputs:
        payload = inputs[item.artifact_id]
        if len(payload) > MAX_FILE_BYTES:
            raise ValueError("input exceeds size budget")
        snapshot = inspect_artifact(payload, item.format)
        if snapshot.sha256 != item.sha256:
            raise ValueError("input checksum mismatch")
        units[item.artifact_id] = {unit.locator: unit.text for unit in snapshot.units}
    answers: list[NativeAnswer] = []
    for assertion in task.assertions:
        if assertion.calculation:
            refs = assertion.calculation.operands
            values = [_number(units[ref.artifact_id][ref.locator]) for ref in refs]
            with localcontext() as context:
                context.prec = 50
                if assertion.calculation.operation == "sum":
                    value = sum(values, Decimal(0))
                elif assertion.calculation.operation == "difference":
                    value = values[0] - values[1]
                else:
                    value = values[0] / values[1]
            answer = str(value)
        else:
            assert assertion.target is not None
            refs = (assertion.target,)
            answer = units[assertion.target.artifact_id][assertion.target.locator]
        answers.append(NativeAnswer(assertion_id=assertion.id, value=answer, citations=refs))
    files: tuple[NativeFile, ...] = ()
    if task.output is not None:
        payload = _output(task, inputs)
        source = next((item for item in task.inputs if item.artifact_id == task.output.source_artifact_id), None)
        files = (NativeFile(artifact_id=task.output.artifact_id, format=task.output.format, content_base64=base64.b64encode(payload).decode("ascii"), source_sha256=source.sha256 if source else None),)
    return NativeSubmission(answers=tuple(answers), files=files)


def qualify_native_task(task: NativeTask, inputs: Mapping[str, bytes]) -> NativeGrade:
    try:
        reference = reference_submission(task, inputs)
        return grade_native_task(task, inputs, reference)
    except (ValueError, KeyError, OSError, ArithmeticError, ImportError, TypeError) as error:
        return NativeGrade(passed=False, findings=(f"reference_invalid:{type(error).__name__}",), metrics={"reference": "native_bytes"})
