"""One-company, evidence-backed corpus scale contracts and materialisation.

A million announced rows are not a million delivered rows. Profiles separate
construction floors from measured files. Business mechanisms remain in the
existing synthesis engine; native prose remains in accepted ArtifactIR. This
module joins those seams without asserting that operational money reconciles
to an unrelated macro close, or counting repeated formats as new evidence.

CSV materialisation streams one record at a time and rotates at both a row and
a byte budget. Optional spreadsheets use write-only cells, bounded by the same
shards. Resume verifies the source recipe and reconstructs its projections;
editing a checksum manifest cannot turn fabricated rows into evidence.
"""
from __future__ import annotations

import csv
import hashlib
import io
import json
import math
import os
from collections.abc import Iterable
from dataclasses import asdict, dataclass, field
from decimal import Decimal
from pathlib import Path
from tempfile import TemporaryDirectory
from typing import TYPE_CHECKING, BinaryIO, Literal

from pydantic import Field

from . import corpus
from .ids import content_key
from .models import Model
from .native_corpus import (
    NativeContent,
    NativeContentExclusion,
    NativeCorpusManifest,
    NativeCorpusPlan,
    native_corpus_coverage,
    plan_native_corpus,
    render_native_corpus,
)
from .synthesis.compiler import digest
from .synthesis.engine import Simulator, entity_id
from .synthesis.models import (
    Column,
    Expr,
    Limits,
    Program,
    Relation,
    Row,
    SynthesisError,
    Table,
    expr,
    literal,
    ref,
)
from .synthesis.storage import Manifest, Recipe, export, iter_export, verify_export

if TYPE_CHECKING:
    from .world import World


class NativeScaleTarget(Model):
    format: Literal["docx", "pptx", "xlsx"]
    minimum_units: int = Field(ge=1, le=10000, strict=True)
    minimum_distinct_facts: int = Field(ge=1, strict=True)


class CorpusScaleProfile(Model):
    name: str = Field(pattern=r"^[a-z][a-z0-9_-]{0,63}$")
    minimum_relational_rows: int = Field(default=0, ge=0, strict=True)
    native_targets: tuple[NativeScaleTarget, ...] = ()


SCALE_PROFILES: dict[str, CorpusScaleProfile] = {
    "development": CorpusScaleProfile(name="development", minimum_relational_rows=10_000,
        native_targets=(NativeScaleTarget(format="docx", minimum_units=30, minimum_distinct_facts=30),
                        NativeScaleTarget(format="pptx", minimum_units=20, minimum_distinct_facts=20))),
    "enterprise": CorpusScaleProfile(name="enterprise", minimum_relational_rows=1_000_000,
        native_targets=(NativeScaleTarget(format="docx", minimum_units=200, minimum_distinct_facts=200),
                        NativeScaleTarget(format="pptx", minimum_units=100, minimum_distinct_facts=100))),
    "stress": CorpusScaleProfile(name="stress", minimum_relational_rows=10_000_000,
        native_targets=(NativeScaleTarget(format="docx", minimum_units=500, minimum_distinct_facts=500),
                        NativeScaleTarget(format="pptx", minimum_units=200, minimum_distinct_facts=200))),
}


class ScaleFinding(Model):
    code: str
    target: str
    expected: int
    observed: int
    message: str = ""


class CorpusScaleAdequacy(Model):
    company_id: str
    profile: str
    relational_rows_planned: int
    grounded_authored_units: int
    distinct_authored_facts: int
    grounded_native_units: int = 0
    distinct_native_facts: int = 0
    adequate: bool
    findings: tuple[ScaleFinding, ...]
    excluded_native_contents: tuple[NativeContentExclusion, ...] = ()


class FactReconciliation(Model):
    table: str
    column: str
    fact_id: str
    decimals: int = Field(default=2, ge=0, le=6, strict=True)


@dataclass(frozen=True)
class LedgerAllocationSource:
    simulator: Simulator
    reconciliation: FactReconciliation


def _fact_amount(world: World, binding: FactReconciliation) -> int:
    fact = next((fact for fact in world.facts if fact.id == binding.fact_id), None)
    if fact is None or fact.value is None or not math.isfinite(fact.value.amount):
        raise SynthesisError("reconciliation_fact", f"finite numeric canonical fact required: {binding.fact_id}")
    scaled = Decimal(str(fact.value.amount)) * 10 ** binding.decimals
    if scaled != scaled.to_integral_value():
        raise SynthesisError("reconciliation_precision", f"{fact.id} cannot be represented at declared precision")
    return int(scaled)


def ledger_allocation_source(world: World, *, fact_id: str, rows: int, seed: int,
                             decimals: int = 2, head_population_bps: int = 1000,
                             head_amount_bps: int = 8000, limits: Limits | None = None) -> LedgerAllocationSource:
    """Generate related ledger/transaction rows that sum to a real World fact.

    The default declared allocation gives a tenth of records four fifths of
    the quantity. A seeded modular permutation assigns strata; integer
    quotient/remainder allocation closes the total exactly without retaining
    row weights or sorting millions of fractional remainders. This is a
    transparent synthetic assumption, not a fitted enterprise distribution.
    """
    from .rng import Rng

    if (type(rows) is not int or rows < 1 or type(head_population_bps) is not int
            or not 1 <= head_population_bps < 10_000 or type(head_amount_bps) is not int
            or not 0 <= head_amount_bps <= 10_000):
        raise SynthesisError("allocation_parameters", "positive rows and valid basis-point strata required")
    binding = FactReconciliation(table="transaction", column="amount_scaled", fact_id=fact_id, decimals=decimals)
    amount = _fact_amount(world, binding)
    if amount < 0:
        raise SynthesisError("allocation_parameters", "transaction allocations require a nonnegative source quantity")
    fact = next(fact for fact in world.facts if fact.id == fact_id)
    assert fact.value is not None
    rng = Rng(seed).derive(f"ledger-allocation/v1/{world.company.id}/{fact_id}")
    stride = rng.integer(1, rows)
    while math.gcd(stride, rows) != 1:
        stride = stride % rows + 1
    offset = rng.integer(0, rows - 1)
    head_rows = max(1, min(rows - 1, rows * head_population_bps // 10_000)) if rows > 1 else 1
    head_amount = amount * head_amount_bps // 10_000 if rows > 1 else amount
    rank = expr("mod", expr("add", expr("mul", ref("_entity"), literal(stride)), literal(offset)), literal(rows))

    def allocation(rank_expr: Expr, count: int, total: int) -> Expr:
        return expr("add", literal(total // count), expr("if", expr("lt", rank_expr, literal(total % count)), literal(1), literal(0)))

    head_value = allocation(ref("allocation_rank"), head_rows, head_amount)
    tail_value = allocation(expr("sub", ref("allocation_rank"), literal(head_rows)), max(1, rows - head_rows), amount - head_amount)
    unit = f"{fact.value.unit}*10^-{decimals}"
    program = Program(namespace="ledger_" + digest([world.company.id, fact_id])[:24], tables=(
        Table(name="ledger", count=1, columns=(
            Column(name="ledger_reference", kind="str", expression=literal(f"{fact.kind}: {fact.subject}")),
            Column(name="unit", kind="str", expression=literal(fact.value.unit)),
            Column(name="decimal_places", expression=literal(decimals)),
            Column(name="control_total_scaled", expression=literal(amount), unit=unit))),
        Table(name="transaction", count=rows, temporal=True,
            relations=(Relation(name="ledger", table="ledger"),), columns=(
                Column(name="allocation_rank", expression=rank),
                Column(name="amount_scaled", expression=expr("if", expr("lt", ref("allocation_rank"), literal(head_rows)), head_value, tail_value), unit=unit))),
    ))
    return LedgerAllocationSource(Simulator(program, seed=seed, limits=limits), binding)


class CorpusScalePlan(Model):
    schema_version: Literal["worldloom.corpus.scale/v1"] = "worldloom.corpus.scale/v1"
    company_id: str
    company_name: str
    source_world_digest: str
    profile: CorpusScaleProfile
    recipe: Recipe
    limits: Limits
    native_plans: tuple[NativeCorpusPlan, ...] = ()
    native_surface: Literal["legacy", "business"] = "business"
    reconciliations: tuple[FactReconciliation, ...] = ()
    csv_shard_rows: int = Field(default=100_000, ge=1, le=1_048_575, strict=True)
    csv_shard_bytes: int = Field(default=32 * 1024 * 1024, ge=4096, le=512 * 1024 * 1024, strict=True)
    maximum_files: int = Field(default=10_000, ge=4, le=100_000, strict=True)
    spreadsheets: bool = False


class CorpusFile(Model):
    path: str
    sha256: str
    file_size_bytes: int = Field(ge=0, strict=True)
    kind: Literal["source", "csv", "xlsx", "native"]


class CorpusColumn(Model):
    name: str
    kind: Literal["int", "bool", "str"]
    unit: str | None = None
    minimum: int | None = None
    maximum: int | None = None
    total: int | None = None
    nonzero_count: int = 0
    distinct_observed: int = 0
    distinct_exact: bool = True


class CorpusTable(Model):
    name: str
    rows: int
    entities: int
    foreign_key_links: int
    columns: tuple[CorpusColumn, ...]
    foreign_keys: dict[str, str]
    csv_shards: tuple[str, ...]
    spreadsheet_shards: tuple[str, ...] = ()


class CorpusScaleManifest(Model):
    schema_version: Literal["worldloom.corpus.scale.manifest/v1"] = "worldloom.corpus.scale.manifest/v1"
    company_id: str
    plan_digest: str
    source: Manifest
    relational_rows: int
    foreign_key_links: int
    tables: tuple[CorpusTable, ...]
    native: tuple[NativeCorpusManifest, ...]
    distinct_native_facts: int
    native_content_units: int
    reconciliations: tuple[FactReconciliation, ...] = ()
    # A page-break floor is a construction metric. Pagination remains a
    # renderer measurement and must never be substituted by this module.
    physical_native_pages: int | None = None
    files: tuple[CorpusFile, ...]
    target_files: tuple[str, ...] = ()


def scale_profile(profile: CorpusScaleProfile | str) -> CorpusScaleProfile:
    if isinstance(profile, CorpusScaleProfile):
        result = profile
    else:
        if profile not in SCALE_PROFILES:
            raise SynthesisError("scale_profile", f"choose one of {', '.join(sorted(SCALE_PROFILES))}")
        result = SCALE_PROFILES[profile]
    formats = [target.format for target in result.native_targets]
    if len(formats) != len(set(formats)):
        raise SynthesisError("scale_profile", "native targets must have distinct formats")
    return result


def _world_digest(world: World) -> str:
    # Fact valid times and prose are part of the binding: a name match alone
    # does not license yesterday's native evidence in today's corpus.
    from .presentation import of as presentation_of
    from .render.values import corpus_locale

    return digest({"company": world.company.model_dump(mode="json"),
                   "recipe": world.recipe,
                   "resolved_locale": corpus_locale(world).as_dict(),
                   "resolved_presentation": asdict(presentation_of(world)),
                   "entity_labels": world.entity_names(),
                   "facts": [fact.model_dump(mode="json") for fact in sorted(world.facts, key=lambda f: f.id)],
                   "artifact_ir": [ir.model_dump(mode="json") for ir in sorted(world.artifact_irs, key=lambda a: a.id)]})


def assess_corpus_scale(world: World, simulator: Simulator, *,
                        profile: CorpusScaleProfile | str = "enterprise",
                        native_surface: Literal["legacy", "business"] = "business") -> CorpusScaleAdequacy:
    """Report construction adequacy; no materialised-row claim is made here."""
    from .narrative import references
    from .presentation import of as presentation_of
    from .render.values import corpus_locale

    target = scale_profile(profile)
    facts = {fact.id: fact for fact in world.facts}
    seen: set[str] = set()
    used: set[str] = set()
    contents: list[NativeContent] = []
    for ir in sorted(world.artifact_irs, key=lambda item: item.id):
        for index, section in enumerate(ir.sections):
            if not section.body:
                if native_surface == "business" and section.table is not None and section.table.rows:
                    contents.append(NativeContent(source_artifact_id=ir.id, section_index=index))
                continue
            referenced = references.referenced(section.body)
            if (not referenced or (set(referenced) | set(section.fact_ids)) - facts.keys()
                    or references.unresolved(section.body, facts)
                    or references.bare_numbers(section.body)):
                continue
            body = references.substitute(section.body, facts, locale=corpus_locale(world), presentation=presentation_of(world))
            identity = " ".join(body.split()).casefold()
            if identity in seen:
                continue
            seen.add(identity)
            used.update(referenced)
            contents.append(NativeContent(source_artifact_id=ir.id, section_index=index))
    findings: list[ScaleFinding] = []
    native_units = len(seen)
    native_fact_count = len(used)
    exclusions: tuple[NativeContentExclusion, ...] = ()
    if contents and native_surface == "business":
        try:
            inventory = plan_native_corpus(world, artifact_id="scale-inventory", format="xlsx",
                title="Company inventory", surface="business", minimum_units=1, minimum_distinct_facts=1)
            coverage = native_corpus_coverage(world, inventory)
            native_units = coverage.content_units
            native_fact_count = len(coverage.fact_ids)
            exclusions = inventory.excluded_contents
        except ValueError as error:
            findings.append(ScaleFinding(code="native_content_invalid", target="company_inventory",
                expected=1, observed=0, message=str(error)))
    if simulator.compiled.rows < target.minimum_relational_rows:
        findings.append(ScaleFinding(code="relational_row_shortfall", target="relational",
            expected=target.minimum_relational_rows, observed=simulator.compiled.rows))
    for native in target.native_targets:
        if native_units < native.minimum_units:
            findings.append(ScaleFinding(code="authored_content_shortfall", target=native.format,
                expected=native.minimum_units, observed=native_units))
        if native_fact_count < native.minimum_distinct_facts:
            findings.append(ScaleFinding(code="canonical_fact_shortfall", target=native.format,
                expected=native.minimum_distinct_facts, observed=native_fact_count))
    return CorpusScaleAdequacy(company_id=world.company.id, profile=target.name,
        relational_rows_planned=simulator.compiled.rows, grounded_authored_units=len(seen),
        distinct_authored_facts=len(used), grounded_native_units=native_units,
        distinct_native_facts=native_fact_count, adequate=not findings, findings=tuple(findings),
        excluded_native_contents=exclusions)


def plan_corpus_scale(world: World, simulator: Simulator, *,
                      profile: CorpusScaleProfile | str = "enterprise",
                      native_plans: Iterable[NativeCorpusPlan] | None = None,
                      native_surface: Literal["legacy", "business"] = "business",
                      reconciliations: Iterable[FactReconciliation] = (),
                      csv_shard_rows: int = 100_000,
                      csv_shard_bytes: int = 32 * 1024 * 1024,
                      maximum_files: int = 10_000,
                      spreadsheets: bool = False) -> CorpusScalePlan:
    """Bind measured-scale targets to one world's facts and a declared process.

    The simulator is an explicit operational scope for this company. Its
    integer money is not silently promoted into pre-existing macro facts.
    Pass native plans to control file topology; otherwise one long native
    file is selected for each profile target.
    """
    target = scale_profile(profile)
    if simulator.compiled.rows < target.minimum_relational_rows:
        raise SynthesisError("relational_row_shortfall",
            f"need {target.minimum_relational_rows}, program declares {simulator.compiled.rows}")
    tables = simulator.program.tables
    if len({table.name.casefold() for table in tables}) != len(tables):
        raise SynthesisError("portable_table_names", "table names collide on case-insensitive filesystems")
    for table in tables:
        reserved = {"company_id", "row_id", "entity_id", "tick"} | {f"fk_{r.name}" for r in table.relations}
        if reserved & {column.name for column in table.columns}:
            raise SynthesisError("reserved_projection_column", table.name)
    if native_plans is None:
        chosen = tuple(plan_native_corpus(world,
            artifact_id=content_key("scale-native/v1", world.company.id, target.name, native.format),
            format=native.format, title=f"{world.company.name} — company evidence",
            surface=native_surface,
            minimum_units=native.minimum_units,
            minimum_distinct_facts=native.minimum_distinct_facts)
            for native in target.native_targets)
    else:
        chosen = tuple(native_plans)
    if len({plan.artifact_id for plan in chosen}) != len(chosen):
        raise SynthesisError("native_plan_identity", "native artifact ids must be distinct")
    for native in target.native_targets:
        # A set of one-page files is not a 200-page document. At least one
        # file must carry the requested evidence depth independently.
        if not any(plan.format == native.format and plan.minimum_units >= native.minimum_units
                   and plan.minimum_distinct_facts >= native.minimum_distinct_facts for plan in chosen):
            raise SynthesisError("native_target_shortfall", native.format)
    bindings = tuple(sorted(reconciliations, key=lambda b: (b.table, b.column, b.fact_id)))
    if len({(binding.table, binding.column) for binding in bindings}) != len(bindings):
        raise SynthesisError("reconciliation_binding", "a quantity column has more than one canonical total")
    for binding in bindings:
        _fact_amount(world, binding)
        bound_table = next((table for table in tables if table.name == binding.table), None)
        column = next((column for column in bound_table.columns if column.name == binding.column), None) if bound_table else None
        fact = next(fact for fact in world.facts if fact.id == binding.fact_id)
        assert fact.value is not None
        if column is None or column.kind != "int" or column.unit != f"{fact.value.unit}*10^-{binding.decimals}":
            raise SynthesisError("reconciliation_binding", "integer column and canonical scaled unit required")
    result = CorpusScalePlan(company_id=world.company.id, company_name=world.company.name,
        source_world_digest=_world_digest(world), profile=target,
        recipe=Recipe.model_validate(simulator.recipe()), limits=simulator.compiled.limits,
        native_plans=chosen, native_surface=native_surface, reconciliations=bindings, csv_shard_rows=csv_shard_rows, csv_shard_bytes=csv_shard_bytes,
        maximum_files=maximum_files, spreadsheets=spreadsheets)
    minimum_files = 5 + len(chosen) + sum(
        (table.count * (simulator.program.ticks if table.temporal else 1) + result.csv_shard_rows - 1)
        // result.csv_shard_rows * (2 if spreadsheets else 1) for table in tables)
    if minimum_files > maximum_files:
        raise SynthesisError("scale_file_budget", f"at least {minimum_files} files exceed {maximum_files}")
    return result


def _file(path: Path, root: Path, kind: Literal["source", "csv", "xlsx", "native"]) -> CorpusFile:
    checksum = hashlib.sha256()
    size = 0
    with path.open("rb") as stream:
        while chunk := stream.read(1024 * 1024):
            checksum.update(chunk)
            size += len(chunk)
    return CorpusFile(path=path.relative_to(root).as_posix(), sha256=checksum.hexdigest(),
                      file_size_bytes=size, kind=kind)


def _csv_line(values: Iterable[object]) -> bytes:
    stream = io.StringIO(newline="")
    csv.writer(stream, lineterminator="\n").writerow(values)
    return stream.getvalue().encode("utf-8")


@dataclass
class _Statistics:
    distinct: set[int | bool | str] = field(default_factory=set)
    exact: bool = True
    minimum: int | None = None
    maximum: int | None = None
    total: int | None = None
    nonzero: int = 0

    def accept(self, value: int | bool | str) -> None:
        if len(self.distinct) < 512:
            self.distinct.add(value)
        elif value not in self.distinct:
            self.exact = False
        if type(value) is int:
            self.minimum = value if self.minimum is None else min(self.minimum, value)
            self.maximum = value if self.maximum is None else max(self.maximum, value)
            self.total = (self.total or 0) + value
            self.nonzero += value != 0
        elif type(value) is bool:
            self.nonzero += value


def _header(table: Table) -> tuple[str, ...]:
    return ("company_id", "row_id", "entity_id", "tick",
            *sorted(column.name for column in table.columns),
            *(f"fk_{relation.name}" for relation in sorted(table.relations, key=lambda r: r.name)))


def _project(row: Row, table: Table, plan: CorpusScalePlan,
             dimension_counts: dict[str, int]) -> tuple[object, ...]:
    values = row.values()
    links = {link.relation: link for link in row.links}
    if set(links) != {relation.name for relation in table.relations}:
        raise SynthesisError("projection_foreign_keys", row.id)
    projected: list[object] = [plan.company_id, row.id, row.entity_id, row.tick]
    projected.extend(values[column.name] for column in sorted(table.columns, key=lambda c: c.name))
    for relation in sorted(table.relations, key=lambda r: r.name):
        link = links[relation.name]
        if (link.table != relation.table or not 0 <= link.entity < dimension_counts[relation.table]
                or link.entity_id != entity_id(plan.recipe.program.namespace, plan.recipe.seed,
                                               relation.table, link.entity)):
            raise SynthesisError("projection_foreign_keys", row.id)
        projected.append(link.entity_id)
    return tuple(projected)


def _spreadsheet(csv_path: Path, target: Path, table: Table, stamp: str) -> None:
    from openpyxl import Workbook
    from openpyxl.cell import WriteOnlyCell

    from .render.ooxml import normalise

    workbook = Workbook(write_only=True)
    sheet = workbook.create_sheet("Records")
    sheet.freeze_panes = "A2"
    kinds = {column.name: column.kind for column in table.columns} | {"tick": "int"}
    with csv_path.open("r", encoding="utf-8", newline="") as stream:
        reader = csv.reader(stream)
        headers = next(reader)
        sheet.append(headers)
        for row in reader:
            cells = []
            for name, text in zip(headers, row, strict=True):
                kind = kinds.get(name, "str")
                value: str | int | bool = int(text) if kind == "int" else (text == "True" if kind == "bool" else text)
                # Excel preserves only fifteen significant decimal digits.
                # Keeping larger exact ledger integers as literal text avoids
                # silently changing the amount while the CSV remains exact.
                if kind == "int" and abs(int(text)) > 999_999_999_999_999:
                    value = text
                cell = WriteOnlyCell(sheet, value=value)
                if isinstance(value, str):
                    cell.data_type = "s"  # source text can begin with '=', including the company name
                cells.append(cell)
            sheet.append(cells)
    payload = io.BytesIO()
    workbook.save(payload)
    target.write_bytes(normalise(payload.getvalue(), created=stamp))


def _materialise_tables(plan: CorpusScalePlan, source: Path, root: Path,
                        stamp: str) -> tuple[tuple[CorpusTable, ...], tuple[CorpusFile, ...]]:
    tables = {table.name: table for table in plan.recipe.program.tables}
    dimensions = {table.name: table.count for table in tables.values() if not table.temporal}
    streams = iter_export(source, limits=plan.limits, replay=False)
    reports: list[CorpusTable] = []
    files: list[CorpusFile] = []
    # sorted source order permits grouping without retaining a table's rows.
    from itertools import groupby

    def close_shard(path: Path, output: BinaryIO, table: Table, workbooks: list[str]) -> None:
        output.close()
        files.append(_file(path, root, "csv"))
        if plan.spreadsheets:
            workbook_path = path.with_suffix(".xlsx")
            _spreadsheet(path, workbook_path, table, stamp)
            files.append(_file(workbook_path, root, "xlsx"))
            workbooks.append(workbook_path.relative_to(root).as_posix())

    for table_name, rows in groupby(streams, key=lambda row: row.table):
        table = tables[table_name]
        statistics = {column.name: _Statistics() for column in table.columns}
        shards: list[str] = []
        workbooks: list[str] = []
        count = entities = links = 0
        previous_entity: int | None = None
        shard_count = shard_rows = shard_bytes = 0
        output: BinaryIO | None = None
        path: Path | None = None

        try:
            for row in rows:
                projected = _csv_line(_project(row, table, plan, dimensions))
                header = _csv_line(_header(table))
                if len(projected) + len(header) > plan.csv_shard_bytes:
                    raise SynthesisError("scale_row_bytes", f"{row.id} does not fit one shard")
                if output is not None and (shard_rows >= plan.csv_shard_rows
                                           or shard_bytes + len(projected) > plan.csv_shard_bytes):
                    assert path is not None
                    close_shard(path, output, table, workbooks)
                    output = None
                if output is None:
                    # Include source metadata, plan and the final manifest in
                    # the total budget, before opening another physical file.
                    if len(files) + len(plan.native_plans) + 5 + (2 if plan.spreadsheets else 1) > plan.maximum_files:
                        raise SynthesisError("scale_file_budget", "byte sharding exceeds maximum_files")
                    path = root / "tables" / f"{table.name}-{shard_count:05}.csv"
                    path.parent.mkdir(parents=True, exist_ok=True)
                    output = path.open("wb")
                    output.write(header)
                    shard_count += 1
                    shard_rows = 0
                    shard_bytes = len(header)
                    shards.append(path.relative_to(root).as_posix())
                output.write(projected)
                shard_rows += 1
                shard_bytes += len(projected)
                count += 1
                links += len(row.links)
                if previous_entity != row.entity:
                    entities += 1
                    previous_entity = row.entity
                for name, value in row.values().items():
                    statistics[name].accept(value)
            if output is not None:
                assert path is not None
                close_shard(path, output, table, workbooks)
                output = None
        finally:
            if output is not None:
                output.close()
        columns = tuple(CorpusColumn(name=column.name, kind=column.kind, unit=column.unit,
            minimum=statistics[column.name].minimum, maximum=statistics[column.name].maximum,
            total=statistics[column.name].total, nonzero_count=statistics[column.name].nonzero,
            distinct_observed=len(statistics[column.name].distinct), distinct_exact=statistics[column.name].exact)
            for column in sorted(table.columns, key=lambda c: c.name))
        reports.append(CorpusTable(name=table.name, rows=count, entities=entities,
            foreign_key_links=links, columns=columns,
            foreign_keys={f"fk_{r.name}": f"{r.table}.entity_id" for r in sorted(table.relations, key=lambda r: r.name)},
            csv_shards=tuple(shards), spreadsheet_shards=tuple(workbooks)))
    return tuple(reports), tuple(files)


def _stamp(world: World) -> str:
    return max(fact.valid_from for fact in world.facts).isoformat() if world.facts else "2000-01-01T00:00:00Z"


def _native_path(plan: NativeCorpusPlan) -> Path:
    return Path("native") / f"{digest(plan.artifact_id)[:32]}.{plan.format}"


def _validate_binding(world: World, plan: CorpusScalePlan) -> Simulator:
    if (world.company.id, world.company.name, _world_digest(world)) != (
        plan.company_id, plan.company_name, plan.source_world_digest
    ):
        raise SynthesisError("scale_company_binding", "world, canonical facts or authored evidence changed")
    simulator = Simulator(plan.recipe.program, seed=plan.recipe.seed,
                          interventions=plan.recipe.interventions, limits=plan.limits)
    checked = plan_corpus_scale(world, simulator, profile=plan.profile,
        native_plans=plan.native_plans, native_surface=plan.native_surface, reconciliations=plan.reconciliations, csv_shard_rows=plan.csv_shard_rows,
        csv_shard_bytes=plan.csv_shard_bytes, maximum_files=plan.maximum_files,
        spreadsheets=plan.spreadsheets)
    if checked != plan:
        raise SynthesisError("invalid_scale_plan", "plan differs from its validated construction")
    return simulator


def _check_reconciliations(world: World, plan: CorpusScalePlan, tables: tuple[CorpusTable, ...]) -> None:
    quantities = {(table.name, column.name): column.total for table in tables for column in table.columns}
    for binding in plan.reconciliations:
        observed = quantities.get((binding.table, binding.column))
        expected = _fact_amount(world, binding)
        if observed != expected:
            raise SynthesisError("canonical_total_mismatch", f"{binding.table}.{binding.column}: {observed} != {binding.fact_id} ({expected})")


def export_corpus_scale(world: World, plan: CorpusScalePlan, directory: Path, *,
                        resume: bool = False) -> CorpusScaleManifest:
    """Atomically materialise actual rows/native bytes; never overwrite a corpus."""
    directory = Path(directory)
    if directory.is_symlink():
        raise SynthesisError("invalid_scale_corpus", "corpus root must be a directory, not a symbolic link")
    simulator = _validate_binding(world, plan)
    if directory.exists():
        if not resume:
            raise SynthesisError("destination_exists", str(directory))
        manifest = verify_corpus_scale(world, directory)
        if manifest.plan_digest != digest(plan.model_dump(mode="json")):
            raise SynthesisError("resume_mismatch", "corpus scale plan changed")
        return manifest
    directory.parent.mkdir(parents=True, exist_ok=True)
    lock = directory.parent / f".{directory.name}.corpus-scale.lock"
    try:
        descriptor = os.open(lock, os.O_CREAT | os.O_EXCL | os.O_WRONLY, 0o600)
    except FileExistsError as error:
        raise SynthesisError("destination_locked", str(directory)) from error
    try:
        os.close(descriptor)
        with TemporaryDirectory(prefix=f".{directory.name}.stage-", dir=directory.parent) as temporary:
            stage = Path(temporary)
            source = export(simulator, stage / "source")
            tables, files = _materialise_tables(plan, stage / "source", stage, _stamp(world))
            _check_reconciliations(world, plan, tables)
            all_files = list(files)
            all_files.extend(_file(path, stage, "source") for path in sorted((stage / "source").iterdir()))
            natives = []
            for native_plan in sorted(plan.native_plans, key=lambda p: p.artifact_id):
                result = render_native_corpus(world, native_plan)
                native_path = stage / _native_path(native_plan)
                native_path.parent.mkdir(parents=True, exist_ok=True)
                native_path.write_bytes(result.payload)
                natives.append(result.manifest)
                all_files.append(_file(native_path, stage, "native"))
            if len(all_files) + 2 > plan.maximum_files:
                raise SynthesisError("scale_file_budget", "materialised files exceed maximum_files")
            distinct_facts = {fact_id for native in natives for evidence in native.evidence for fact_id in evidence.fact_ids}
            manifest = CorpusScaleManifest(company_id=plan.company_id,
                plan_digest=digest(plan.model_dump(mode="json")), source=source,
                relational_rows=sum(table.rows for table in tables),
                foreign_key_links=sum(table.foreign_key_links for table in tables), tables=tables,
                native=tuple(natives), distinct_native_facts=len(distinct_facts),
                native_content_units=sum(native.content_units for native in natives),
                reconciliations=plan.reconciliations,
                files=tuple(sorted(all_files, key=lambda f: f.path)),
                target_files=tuple(sorted(file.path for file in all_files if file.kind != "source")))
            corpus.write_json(stage / "plan.json", plan.model_dump(mode="json"))
            corpus.write_json(stage / "manifest.json", manifest.model_dump(mode="json"))
            if directory.exists():
                raise SynthesisError("destination_exists", str(directory))
            stage.rename(directory)
            return manifest
    finally:
        lock.unlink()


def verify_corpus_scale(world: World, directory: Path) -> CorpusScaleManifest:
    """Verify canonical metadata, recipe replay and every projected file's bytes."""
    directory = Path(directory)
    if directory.is_symlink():
        raise SynthesisError("invalid_scale_corpus", "corpus root must be a directory, not a symbolic link")
    for name in ("plan.json", "manifest.json"):
        path = directory / name
        if path.is_symlink() or not path.is_file() or path.stat().st_size > 128 * 1024 * 1024:
            raise SynthesisError("invalid_scale_corpus", f"invalid or excessive {name}")
    try:
        plan = CorpusScalePlan.model_validate(corpus.read_json(directory / "plan.json"))
        manifest = CorpusScaleManifest.model_validate(corpus.read_json(directory / "manifest.json"))
    except ValueError as error:
        raise SynthesisError("invalid_scale_corpus", str(error)) from error
    for name, metadata in (("plan.json", plan), ("manifest.json", manifest)):
        encoded = (json.dumps(metadata.model_dump(mode="json"), indent=2, sort_keys=True) + "\n").encode("utf-8")
        if (directory / name).read_bytes() != encoded:
            raise SynthesisError("scale_metadata_encoding", f"noncanonical {name}")
    simulator = _validate_binding(world, plan)
    if manifest.company_id != plan.company_id or manifest.plan_digest != digest(plan.model_dump(mode="json")):
        raise SynthesisError("scale_manifest_binding", "manifest does not bind this plan and company")
    expected = {"plan.json", "manifest.json"} | {file.path for file in manifest.files}
    paths = tuple(directory.rglob("*"))
    actual = {path.relative_to(directory).as_posix() for path in paths if path.is_file()}
    if (any(path.is_symlink() for path in paths) or actual != expected
            or len(expected) > plan.maximum_files or len(manifest.files) != len(expected) - 2):
        raise SynthesisError("scale_file_set", "missing, duplicate, additional or symbolic-link files")
    for file in manifest.files:
        if _file(directory / file.path, directory, file.kind) != file:
            raise SynthesisError("scale_checksum", file.path)
    source = verify_export(directory / "source", limits=plan.limits)
    if (source != manifest.source or source.recipe_digest != simulator.run_digest
            or source.shard_index != 0 or source.shard_count != 1):
        raise SynthesisError("scale_source_binding", "source manifest changed")
    # Reconstruct projections from independently replayed operational evidence.
    # Checksums alone would accept a CSV modified alongside its manifest.
    with TemporaryDirectory(prefix="worldloom-scale-verify-") as temporary:
        tables, files = _materialise_tables(plan, directory / "source", Path(temporary), _stamp(world))
    _check_reconciliations(world, plan, tables)
    expected_files = tuple(file for file in manifest.files if file.kind in {"csv", "xlsx"})
    if tables != manifest.tables or tuple(sorted(files, key=lambda f: f.path)) != expected_files:
        raise SynthesisError("scale_projection_replay", "structured projections are not generated by the recipe")
    natives = []
    for native_plan in sorted(plan.native_plans, key=lambda p: p.artifact_id):
        result = render_native_corpus(world, native_plan)
        if (directory / _native_path(native_plan)).read_bytes() != result.payload:
            raise SynthesisError("scale_native_replay", native_plan.artifact_id)
        natives.append(result.manifest)
    used = {fact_id for native in natives for evidence in native.evidence for fact_id in evidence.fact_ids}
    if (tuple(natives) != manifest.native or len(used) != manifest.distinct_native_facts
            or sum(native.content_units for native in natives) != manifest.native_content_units
            or sum(table.rows for table in tables) != manifest.relational_rows
            or sum(table.foreign_key_links for table in tables) != manifest.foreign_key_links
            or manifest.physical_native_pages is not None or manifest.reconciliations != plan.reconciliations
            or manifest.target_files != tuple(sorted(file.path for file in manifest.files if file.kind != "source"))):
        raise SynthesisError("scale_measurement", "reported measurements differ from materialised evidence")
    return manifest


__all__ = [
    "NativeScaleTarget", "CorpusScaleProfile", "SCALE_PROFILES", "scale_profile",
    "ScaleFinding", "CorpusScaleAdequacy", "assess_corpus_scale", "CorpusScalePlan", "plan_corpus_scale",
    "FactReconciliation", "LedgerAllocationSource", "ledger_allocation_source",
    "CorpusFile", "CorpusColumn", "CorpusTable", "CorpusScaleManifest",
    "export_corpus_scale", "verify_corpus_scale",
]
